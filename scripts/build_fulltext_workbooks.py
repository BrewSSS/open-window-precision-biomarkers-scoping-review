#!/usr/bin/env python3
"""Build the two full-text screening workbooks (B and C) for the confirmed validation layer (PRE-008).

Input: a CSV of records sent to full text. Once verification is merged this is the human-confirmed V list
(e.g. ai_triage/human_verification_merged.csv filtered to the confirmed V layer); until then it can be tested on
ai_triage/final_triage_4298.csv with --filter final_layer=V. DOI/PMID are joined from ai_stage2/ta_ai_merged.csv
when the input lacks them. An optional retrieval log (record_id, retrieval_status, pdf_path) pre-fills retrieval.

Output (git-ignored except the manifest): <out-dir>/ft_screen_B.xlsx, ft_screen_C.xlsx, ft_workbooks_manifest.json.
Sheets: README, screen (decision columns with dropdowns), lists (hidden). No abstracts and no AI hints are written.

Usage:
  python3 scripts/build_fulltext_workbooks.py \
      --input 04_screening/formal_2026-10-05_v0.9/ai_triage/final_triage_4298.csv --filter final_layer=V
  python3 scripts/build_fulltext_workbooks.py --input <human_verification_merged.csv> --filter <col>=<value> \
      --retrieval-log <retrieval_log.csv>
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
FORMAL = ROOT / "04_screening/formal_2026-10-05_v0.9"
DEFAULT_OUT = FORMAL / "fulltext"
META_CSV = FORMAL / "ai_stage2/ta_ai_merged.csv"
CODES_JSON = ROOT / "04_screening/fulltext_exclusion_codes.json"

DISPOSITIONS = ["INCLUDE_A", "INCLUDE_B", "INCLUDE_A_AND_B", "RETAIN_BACKGROUND", "EXCLUDE",
                "AWAITING_CLASSIFICATION", "NOT_RETRIEVED", "DUPLICATE_REPORT"]
AGE_CHECK = ["adults_confirmed", "mean_minus_2SD_below_18_unresolved", "under_18_included", "not_applicable"]
VALIDATION = ["yes", "no", "unclear"]
RETRIEVAL = ["not_yet_sought", "retrieved", "not_retrieved_after_attempts", "awaiting_translation"]

FIXED = ["record_id", "title", "journal", "year", "doi", "pmid", "retrieval_status", "pdf_path", "abstract_E5_subtypes"]
DECISION = ["your_disposition", "your_primary_code", "your_secondary_notes", "your_age_rule_check",
            "your_validation_element_confirmed", "your_comment"]
COLUMNS = FIXED + DECISION
WIDTHS = {"record_id": 12, "title": 60, "journal": 28, "year": 6, "doi": 26, "pmid": 10, "retrieval_status": 18,
          "pdf_path": 30, "abstract_E5_subtypes": 22, "your_disposition": 24, "your_primary_code": 38,
          "your_secondary_notes": 40, "your_age_rule_check": 34, "your_validation_element_confirmed": 14,
          "your_comment": 50}


def ft_codes() -> list[str]:
    return json.loads(CODES_JSON.read_text())["primary_reason_rule"]["hierarchy"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_rows(path: Path, filters: list[str]) -> list[dict]:
    rows = list(csv.DictReader(path.open(newline="", encoding="utf-8")))
    for f in filters:
        col, _, val = f.partition("=")
        if rows and col not in rows[0]:
            sys.exit(f"filter column {col!r} not in {path}")
        rows = [r for r in rows if r.get(col, "") == val]
    if "abstract" in (rows[0] if rows else {}):
        for r in rows:
            r.pop("abstract", None)  # never carried into the workbook
    ids = [r["record_id"] for r in rows]
    if len(ids) != len(set(ids)):
        sys.exit("duplicate record_id in input")
    return sorted(rows, key=lambda r: r["record_id"])


def readme_lines(reviewer: str, n: int, src: str) -> list[str]:
    codes = ft_codes()
    return [
        f"Full-text screening workbook: reviewer {reviewer}",
        f"Records: {n} (confirmed validation layer, PRE-008). Source list: {src}",
        "",
        "Rules: protocol v3.2 sections C-D; screening_manual.md 3C; fulltext_exclusion_codes.json; boundary exercises in 04_screening/fulltext_boundary_exercises.md.",
        "Work alone. Do not look at the other reviewer's workbook before both are locked. Read methods, results, relevant supplements and linked reports.",
        "",
        "Fill only the 'your_' columns of the screen sheet; do not edit, sort or delete the other columns or rows.",
        "your_disposition (one per report):",
        "  INCLUDE_A / INCLUDE_B / INCLUDE_A_AND_B - eligible; A = identifiable acute bout + post-cessation sample + baseline/comparator; B = same marker, same people, >=2 identified bouts, known timing, baseline/comparator.",
        "  RETAIN_BACKGROUND - original healthy-adult study with an eligible immune link that fails only through a background exposure pattern (habitual/resting cross-section, periodic monitoring with unknown last-bout timing, training snapshot with unknown interval). No primary code; write the failed dimension (normally FT04) in your_secondary_notes.",
        "  EXCLUDE - a scientific reason is established: enter exactly one primary code (the first in hierarchy order); other failed dimensions go to your_secondary_notes.",
        "  AWAITING_CLASSIFICATION - a material fact is unresolved (age rule, health status, timing, separability). Also use it, with a comment, for a preprint-only report ('frontier preprint') or a report awaiting translation.",
        "  NOT_RETRIEVED - full text not obtained after documented attempts (not a scientific exclusion).",
        "  DUPLICATE_REPORT - same report as another record (version/duplicate); name the retained record_id in your_comment. Distinct reports of one cohort are NOT duplicates: screen each.",
        "your_primary_code: only with EXCLUDE. Hierarchy: " + " > ".join(codes),
        "  FT05 = only during-exercise samples. Sampling only after 72 h is NOT FT05 (tag outside-window recovery). No intensity/duration threshold under FT04.",
        "your_age_rule_check: adults_confirmed (stated range/minimum >=18 y, or mean - 2*SD >= 18 y); mean_minus_2SD_below_18_unresolved (-> AWAITING_CLASSIFICATION unless resolved from full text/supplement/linked report); under_18_included (report states <18 y and no separable adult stratum -> FT03); not_applicable.",
        "your_validation_element_confirmed: does the full text show the validation-readiness element (protocol C, v1.2: metric validation, outcome linkage, omics discovery, repeated monitoring)? yes / no / unclear. It is not an eligibility criterion; it does not change your disposition.",
        "abstract_E5_subtypes: the sub-types recorded at abstract level, shown so that you know what to confirm. It is not a disposition hint.",
        "your_comment: page/section of the passage supporting the decision (required for EXCLUDE, RETAIN_BACKGROUND and AWAITING_CLASSIFICATION).",
        "",
        "Red cells flag a code without EXCLUDE, or EXCLUDE without a code. When finished, save, close and tell D; D records the SHA-256 before merging (scripts/merge_fulltext_screening.py).",
    ]


def build(rows: list[dict], reviewer: str, out: Path, src: str) -> None:
    codes = ft_codes()
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    for i, line in enumerate(readme_lines(reviewer, len(rows), src), 1):
        ws0.cell(i, 1, line).alignment = Alignment(wrap_text=False)
    ws0["A1"].font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 160

    lists = wb.create_sheet("lists")
    for j, (name, vals) in enumerate([("disposition", DISPOSITIONS), ("primary_code", codes),
                                      ("age_rule_check", AGE_CHECK), ("validation", VALIDATION)], 1):
        lists.cell(1, j, name)
        for i, v in enumerate(vals, 2):
            lists.cell(i, j, v)
    lists.sheet_state = "hidden"

    ws = wb.create_sheet("screen", 1)
    ws.append(COLUMNS)
    head_fill = PatternFill("solid", fgColor="DDDDDD")
    you_fill = PatternFill("solid", fgColor="FFF2CC")
    for j, c in enumerate(COLUMNS, 1):
        cell = ws.cell(1, j)
        cell.font = Font(bold=True)
        cell.fill = you_fill if c in DECISION else head_fill
        ws.column_dimensions[cell.column_letter].width = WIDTHS[c]
    for r in rows:
        ws.append([r.get(c, "") for c in FIXED] + [None] * len(DECISION))
    n = len(rows)
    last = max(n + 1, 2)
    col = {c: ws.cell(1, j).column_letter for j, c in enumerate(COLUMNS, 1)}

    def dv(letter: str, ref: str, name: str) -> None:
        v = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True,
                           errorTitle="Invalid value", error=f"Choose a value from the {name} list")
        v.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(v)

    dv(col["your_disposition"], f"lists!$A$2:$A${len(DISPOSITIONS) + 1}", "disposition")
    dv(col["your_primary_code"], f"lists!$B$2:$B${len(codes) + 1}", "FT code")
    dv(col["your_age_rule_check"], f"lists!$C$2:$C${len(AGE_CHECK) + 1}", "age-rule")
    dv(col["your_validation_element_confirmed"], f"lists!$D$2:$D${len(VALIDATION) + 1}", "yes/no/unclear")
    red = PatternFill("solid", fgColor="F4CCCC")
    d, p = col["your_disposition"], col["your_primary_code"]
    ws.conditional_formatting.add(
        f"{p}2:{p}{last}",
        FormulaRule(formula=[f'OR(AND(${p}2<>"",${d}2<>"EXCLUDE"),AND(${d}2="EXCLUDE",${p}2=""))'], fill=red))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"
    wb.active = 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def verify(path: Path, n: int) -> None:
    wb = load_workbook(path)
    assert wb.sheetnames == ["README", "screen", "lists"], wb.sheetnames
    ws = wb["screen"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.max_row == n + 1, (ws.max_row, n)
    assert len(ws.data_validations.dataValidation) == 4
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert all(v is None for v in row[len(FIXED):]), "decision cells must be empty"


def main(argv=None) -> dict:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--filter", action="append", default=[], help="col=value (repeatable)")
    ap.add_argument("--retrieval-log", type=Path, help="CSV with record_id, retrieval_status, pdf_path")
    ap.add_argument("--meta", type=Path, default=META_CSV, help="CSV with record_id, doi, pmid")
    ap.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--reviewers", default="B,C")
    a = ap.parse_args(argv)

    rows = read_rows(a.input, a.filter)
    if not rows:
        sys.exit("no rows after filtering")
    meta = {}
    if a.meta and a.meta.exists():
        meta = {m["record_id"]: m for m in csv.DictReader(a.meta.open(newline="", encoding="utf-8"))}
    retr = {}
    if a.retrieval_log:
        retr = {m["record_id"]: m for m in csv.DictReader(a.retrieval_log.open(newline="", encoding="utf-8"))}
    for r in rows:
        m = meta.get(r["record_id"], {})
        for k in ("title", "journal", "year", "doi", "pmid"):
            if not r.get(k):
                r[k] = m.get(k, "")
        r["abstract_E5_subtypes"] = r.get("final_E5_subtypes", r.get("abstract_E5_subtypes", ""))
        lg = retr.get(r["record_id"], {})
        status = lg.get("retrieval_status") or "not_yet_sought"
        if status not in RETRIEVAL:
            sys.exit(f"unknown retrieval_status {status!r} for {r['record_id']}")
        r["retrieval_status"] = status
        r["pdf_path"] = lg.get("pdf_path", "")

    def rel(p: Path) -> str:
        p = p.resolve()
        return str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p)

    out = {}
    for rv in a.reviewers.split(","):
        path = a.out_dir / f"ft_screen_{rv}.xlsx"
        build(rows, rv, path, rel(a.input))
        verify(path, len(rows))
        out[rv] = path
    counts_retr = {}
    for r in rows:
        counts_retr[r["retrieval_status"]] = counts_retr.get(r["retrieval_status"], 0) + 1
    manifest = {
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "script": "scripts/build_fulltext_workbooks.py",
        "input": {"path": rel(a.input), "sha256": sha256(a.input), "filters": a.filter},
        "retrieval_log": ({"path": rel(a.retrieval_log), "sha256": sha256(a.retrieval_log)} if a.retrieval_log else None),
        "n_records": len(rows),
        "record_ids_sha256": hashlib.sha256("\n".join(r["record_id"] for r in rows).encode()).hexdigest(),
        "counts": {"retrieval_status": counts_retr,
                   "missing_doi": sum(1 for r in rows if not r["doi"]),
                   "missing_pmid": sum(1 for r in rows if not r["pmid"])},
        "columns": COLUMNS,
        "dropdowns": {"your_disposition": DISPOSITIONS, "your_primary_code": ft_codes(),
                      "your_age_rule_check": AGE_CHECK, "your_validation_element_confirmed": VALIDATION},
        "workbooks": {rv: {"path": rel(p), "sha256_blank": sha256(p)} for rv, p in out.items()},
        "locks": {rv: None for rv in out},
        "note": "Workbooks are git-ignored (titles only, but reviewer comments may quote full texts). D records the SHA-256 of each locked workbook under 'locks' and commits it before merging.",
    }
    (a.out_dir / "ft_workbooks_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"n_records": len(rows), "workbooks": {k: rel(v) for k, v in out.items()},
                      "retrieval_status": counts_retr}, ensure_ascii=False))
    return manifest


if __name__ == "__main__":
    main()
