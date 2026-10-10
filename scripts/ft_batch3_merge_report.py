#!/usr/bin/env python3
"""Prepare batch-3 D review, combine D resolutions, and report screened full-text counts.

The 2026-10-09 D workbook remains the source of record for its 39 resolutions.
This script never edits it or either reviewer workbook. Generated CSVs contain titles,
codes, and counts only; page-cited reviewer comments remain in ignored workbooks.

Usage:
  python3 scripts/ft_batch3_merge_report.py scope-ids --ids-file ft_todo_batch3.txt --scope-file ft_scope_477_2026-10-09.csv --fetch-manifest fulltext_fetch_manifest.csv --human-final human_verification_final.csv --text-dir V_text --out ft_todo_batch3_in_scope.txt --hold ft_batch3_scope_hold.txt --audit ft_batch3_scope_audit.md
  python3 scripts/ft_batch3_merge_report.py prepare --b B.xlsx --c C.xlsx --ids-file ft_todo_batch3_in_scope.txt --out ft_conflicts_for_D_batch3.xlsx [--template ft_conflicts_for_D.xlsx]
  python3 scripts/ft_batch3_merge_report.py combine --old ft_conflicts_for_D.xlsx --batch3 ft_conflicts_for_D_batch3.xlsx --out ft_D_resolutions_combined.csv
  python3 scripts/ft_batch3_merge_report.py report --b B.xlsx --c C.xlsx --ids-file ft_todo_batch3_in_scope.txt --source-ids-file ft_todo_batch3.txt --held-ids-file ft_batch3_scope_hold.txt --merge-summary ft_merge_summary.json --merged ft_merged.csv --ledger token_ledger_batch3.csv --out ft_batch3_summary_2026-10-10.md --newly-included newly_included_batch3.txt
"""
from __future__ import annotations

import argparse
import copy
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_fulltext_workbooks as bfw  # noqa: E402
import merge_fulltext_screening as merge  # noqa: E402

ADJ_COLUMNS = ["conflict_type", "record_id", "title", "doi", "B_disposition", "B_code",
               "B_comment", "C_disposition", "C_code", "C_comment", "final_disposition",
               "final_primary_code", "D_note", "text_file"]
RES_COLUMNS = ["record_id", "title", "final_disposition", "final_primary_code"]
FIELDS = ["disposition", "primary_code", "age_rule_check", "validation_element_confirmed",
          "secondary_notes", "comment"]


def ids_from(path: Path) -> list[str]:
    ids = [line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(ids) != len(set(ids)):
        raise ValueError(f"duplicate record ID in {path}")
    return ids


def scope_ids(args) -> dict:
    ids = ids_from(args.ids_file)
    with args.scope_file.open(newline="", encoding="utf-8") as fh:
        scope = {r["record_id"] for r in csv.DictReader(fh)}
    with args.fetch_manifest.open(newline="", encoding="utf-8") as fh:
        fetch = {r["record_id"]: r for r in csv.DictReader(fh)}
    with args.human_final.open(newline="", encoding="utf-8") as fh:
        human = {r["record_id"]: r for r in csv.DictReader(fh)}
    included, outside = [], []
    for rid in ids:
        if rid not in fetch or rid not in human or not (args.text_dir / f"{rid}.txt").is_file():
            raise ValueError(f"scope audit lacks fetch, human decision, or text for {rid}")
        if rid in scope:
            if fetch[rid].get("scope") != "confirmed" or human[rid].get("final_layer") not in ("V", "U"):
                raise ValueError(f"inconsistent in-scope decision for {rid}")
            included.append(rid)
        else:
            if fetch[rid].get("scope") != "out_of_scope" or human[rid].get("final_layer") != "M":
                raise ValueError(f"outside-scope status not confirmed for {rid}")
            outside.append(rid)
    held_deciders = Counter(human[rid].get("decided_by") or "unknown" for rid in outside)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.hold.parent.mkdir(parents=True, exist_ok=True)
    args.audit.parent.mkdir(parents=True, exist_ok=True)
    for path, values in ((args.out, included), (args.hold, outside)):
        content = "\n".join(values) + ("\n" if values else "")
        if path.exists() and path.read_text(encoding="utf-8") != content:
            raise ValueError(f"existing scope list differs; refusing to overwrite: {path.name}")
        if not path.exists():
            path.write_text(content, encoding="utf-8")
    lines = ["# Batch-3 full-text scope audit", "",
             f"The source list has {len(ids)} text files: {len(included)} belong to the confirmed "
             f"{len(scope)}-record full-text scope and {len(outside)} have a human final layer M "
             "and are marked out_of_scope in the fetch manifest.",
             f"Held-record human decisions: {json.dumps(dict(sorted(held_deciders.items())), sort_keys=True)}.",
             f"The screened denominator after this batch is the prior 190 plus {len(included)}; "
             f"the remaining confirmed-scope records without paired screening total {len(scope) - 190 - len(included)}.",
             "The excluded-from-batch IDs are retained as retrieved out-of-scope records, "
             "not counted as full-text exclusions or D adjudication conflicts.", "",
             "| Record ID | Human final layer | Fetch scope |", "|---|---|---|"]
    lines.extend(f"| {rid} | M | out_of_scope |" for rid in outside)
    audit_text = "\n".join(lines) + "\n"
    if args.audit.exists() and args.audit.read_text(encoding="utf-8") != audit_text:
        raise ValueError(f"existing scope audit differs; refusing to overwrite: {args.audit.name}")
    if not args.audit.exists():
        args.audit.write_text(audit_text, encoding="utf-8")
    return {"source_ids": len(ids), "in_scope": len(included), "out_of_scope": len(outside),
            "screened_after": 190 + len(included), "scope_unpaired_after": len(scope) - 190 - len(included)}


def conflict_type(b: str, c: str) -> str:
    if b == c:
        return "3_OTHER"
    if (b in merge.INCLUDE) != (c in merge.INCLUDE):
        return "1_INCLUDE_VS_NOT"
    if b in merge.INCLUDE and c in merge.INCLUDE:
        return "2_STREAM_ONLY"
    return "3_OTHER"


def prepare(args) -> dict:
    if args.out.exists():
        raise ValueError(f"refusing to replace existing D workbook: {args.out}")
    template = args.template or args.out.parent / "ft_conflicts_for_D.xlsx"
    if not template.is_file() or template.resolve() == args.out.resolve():
        raise ValueError("an existing original D workbook is required as a separate template")
    B, C = merge.read_wb(args.b), merge.read_wb(args.c)
    ids = ids_from(args.ids_file)
    if set(B) != set(C) or set(ids) - set(B):
        raise ValueError("reviewer workbook record sets do not cover the batch-3 ID list")
    missing = [i for i in ids if not B[i]["your_disposition"] or not C[i]["your_disposition"]]
    if missing:
        raise ValueError(f"both reviewer decisions required before D workbook: {len(missing)} missing; first {missing[:5]}")
    rows = []
    for i in ids:
        b, c = B[i], C[i]
        pb, pc = merge.rule_problems(b), merge.rule_problems(c)
        if merge.decision(b) == merge.decision(c) and not pb and not pc:
            continue
        rows.append([conflict_type(b["your_disposition"], c["your_disposition"]), i,
                     b["title"], b["doi"], b["your_disposition"], b["your_primary_code"],
                     b["your_comment"], c["your_disposition"], c["your_primary_code"],
                     c["your_comment"], None, None, None, f"V_text/{i}.txt"])
    wb = load_workbook(template)
    if "adjudicate" not in wb or "README" not in wb:
        raise ValueError("D workbook template lacks adjudicate or README sheet")
    ws = wb["adjudicate"]
    if [c.value for c in ws[1]] != ADJ_COLUMNS or len(ws.data_validations.dataValidation) != 2:
        raise ValueError("unexpected D workbook template layout or validation count")
    row_styles = [copy.copy(ws.cell(2, col)._style) for col in range(1, len(ADJ_COLUMNS) + 1)]
    ws.delete_rows(2, ws.max_row - 1)
    for row in rows:
        ws.append(row)
        for col, style in enumerate(row_styles, 1):
            ws.cell(ws.max_row, col)._style = copy.copy(style)
    for validation in ws.data_validations.dataValidation:
        old_range = str(validation.sqref)
        if old_range.startswith("K2:K"):
            validation.sqref = f"K2:K{max(2, len(rows) + 1)}"
        elif old_range.startswith("L2:L"):
            validation.sqref = f"L2:L{max(2, len(rows) + 1)}"
        else:
            raise ValueError("unexpected D workbook validation range")
    if "_adjud_D" in wb:
        del wb["_adjud_D"]
    readme = wb["README"]
    readme.delete_rows(1, readme.max_row)
    for line in ["Batch-3 full-text conflicts for D.",
                 "Fill final_disposition, final_primary_code (EXCLUDE only), and D_note in the yellow columns.",
                 "B/C comments may quote full texts; keep this workbook git-ignored.",
                 "The earlier 39 resolutions remain in the original ft_conflicts_for_D.xlsx."]:
        readme.append([line])
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    check = load_workbook(args.out, read_only=True, data_only=True)
    assert check["adjudicate"].max_row == len(rows) + 1
    return {"batch3_records": len(ids), "conflicts": len(rows), "workbook": args.out.name}


def read_d(path: Path) -> dict[str, dict[str, str]]:
    ws = load_workbook(path, read_only=True, data_only=True)["adjudicate"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    if head != ADJ_COLUMNS:
        raise ValueError(f"unexpected adjudication layout: {path}")
    rows = {}
    for values in it:
        r = dict(zip(head, values))
        rid = r["record_id"]
        if not rid:
            continue
        if rid in rows:
            raise ValueError(f"duplicate D row: {rid}")
        d, code = r["final_disposition"] or "", r["final_primary_code"] or ""
        if not d:
            raise ValueError(f"D decision missing in {path}: {rid}")
        if d not in bfw.DISPOSITIONS or (d == "EXCLUDE") != bool(code) or (code and code not in bfw.ft_codes()):
            raise ValueError(f"invalid D decision in {path}: {rid} {d} {code}")
        rows[rid] = {"record_id": rid, "title": r["title"] or "",
                     "final_disposition": d, "final_primary_code": code}
    return rows


def combine(args) -> dict:
    old = read_d(args.old)
    batch = read_d(args.batch3)
    if len(old) != 39:
        raise ValueError(f"expected 39 preserved old D resolutions, found {len(old)}")
    if set(old) & set(batch):
        raise ValueError(f"overlapping old and batch-3 D resolutions: {sorted(set(old) & set(batch))[:5]}")
    if args.out.exists():
        raise ValueError(f"refusing to replace existing combined resolutions: {args.out}")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=RES_COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows([old[i] for i in sorted(old)] + [batch[i] for i in sorted(batch)])
    return {"old_D_resolutions_preserved": len(old), "batch3_D_resolutions": len(batch),
            "combined": len(old) + len(batch), "out": args.out.name}


def override_rates(path: Path, ids: set[str]) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    names = [s for s in wb.sheetnames if s.startswith("_override_")]
    if len(names) != 1:
        raise ValueError(f"expected one override sheet in {path}")
    ws = wb[names[0]]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    latest = {}
    def normalized(value) -> str:
        return "" if value is None or not str(value).strip() else str(value).strip()
    for values in it:
        r = dict(zip(head, values))
        if r.get("record_id") in ids:
            field = str(r.get("column") or "").removeprefix("your_")
            if field in FIELDS:
                latest[(r["record_id"], field)] = (str(r.get("agree_or_override") or "").lower(),
                                                    normalized(r.get("old_value")) != normalized(r.get("new_value")))
    counts = defaultdict(Counter)
    for (_, field), (flag, changed) in latest.items():
        counts[field][flag] += 1
        counts[field]["cell_changed"] += int(changed)
    return {f: {"n": sum(v for k, v in counts[f].items() if k != "cell_changed"),
                "model_reported_overrides": counts[f]["override"],
                "cell_changes": counts[f]["cell_changed"],
                "model_reported_rate": round(counts[f]["override"] / (sum(v for k, v in counts[f].items() if k != "cell_changed")), 4) if counts[f] else None,
                "cell_change_rate": round(counts[f]["cell_changed"] / (sum(v for k, v in counts[f].items() if k != "cell_changed")), 4) if counts[f] else None}
            for f in FIELDS}


def prefill_dispositions(path: Path, ids: set[str]) -> dict[str, int]:
    ws = load_workbook(path, read_only=True, data_only=True)["_prefill"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    latest = {}
    for values in it:
        r = dict(zip(head, values))
        if r.get("record_id") in ids and r.get("column") == "your_disposition":
            latest[r["record_id"]] = str(r.get("value") or "")
    return dict(sorted(Counter(latest.values()).items()))


def token_stats(path: Path, processed_ids: set[str], held_ids: set[str]) -> dict:
    totals = Counter()
    by_stage = defaultdict(Counter)
    record_total = Counter()
    record_calls = Counter()
    record_unknown = Counter()
    category = defaultdict(Counter)
    with path.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rid = r.get("record_id") or ""
            stage = r.get("stage") or "unspecified"
            group = "processed" if rid in processed_ids else "held_out_of_scope" if rid in held_ids else "probe_or_other"
            usage_unknown = any(not str(r.get(key) or "").strip() for key in ("input_tokens", "output_tokens"))
            if usage_unknown:
                totals["calls_without_usage"] += 1
                by_stage[stage]["calls_without_usage"] += 1
                category[group]["calls_without_usage"] += 1
                record_unknown[rid] += 1
            for key in ("input_tokens", "cached_input_tokens", "output_tokens", "reasoning_tokens"):
                raw = str(r.get(key) or "").strip()
                if not raw:
                    continue  # Missing usage is unknown; known fields still form a lower bound.
                n = int(raw)
                totals[key] += n
                by_stage[stage][key] += n
                category[group][key] += n
                if key in ("input_tokens", "output_tokens"):
                    record_total[rid] += n
            totals["calls"] += 1
            by_stage[stage]["calls"] += 1
            category[group]["calls"] += 1
            record_calls[rid] += 1
            if str(r.get("valid", "")).lower() in ("true", "1", "yes"):
                totals["valid_calls"] += 1
    totals["billable_tokens_measured_lower_bound"] = totals["input_tokens"] + totals["output_tokens"]
    for group in category.values():
        group["billable_tokens_measured_lower_bound"] = group["input_tokens"] + group["output_tokens"]
    all_ids = sorted(set(record_calls) | processed_ids | held_ids)
    return {"totals": dict(totals), "by_stage": {s: dict(c) for s, c in sorted(by_stage.items())},
            "by_scope_group": {g: dict(c) for g, c in sorted(category.items())},
            "per_record_tokens": {i: {"measured_lower_bound": record_total[i], "calls": record_calls[i],
                                      "calls_without_usage": record_unknown[i]} for i in all_ids},
            "n_records_with_calls": sum(record_calls[i] > 0 for i in processed_ids | held_ids),
            "mean_tokens_per_source_record_lower_bound":
            round(totals["billable_tokens_measured_lower_bound"] / len(processed_ids | held_ids), 1)
            if processed_ids or held_ids else None,
            "mean_tokens_per_processed_record_lower_bound":
            round(category["processed"]["billable_tokens_measured_lower_bound"] / len(processed_ids), 1)
            if processed_ids else None,
            "estimated_weekly_quota_pct_lower_bound":
            round(totals["billable_tokens_measured_lower_bound"] / 1_010_000, 3)}


def report(args) -> dict:
    ids = set(ids_from(args.ids_file))
    source_ids = set(ids_from(args.source_ids_file))
    held_ids = set(ids_from(args.held_ids_file))
    if ids & held_ids or ids | held_ids != source_ids:
        raise ValueError("processed and held IDs must partition the original source list")
    B, C = merge.read_wb(args.b), merge.read_wb(args.c)
    if ids - set(B) or ids - set(C):
        raise ValueError("batch-3 IDs missing from reviewer workbooks")
    if any(not B[i]["your_disposition"] or not C[i]["your_disposition"] for i in ids):
        raise ValueError("batch-3 paired decisions incomplete")
    batch_disp_pairs = [(B[i]["your_disposition"], C[i]["your_disposition"]) for i in sorted(ids)]
    batch_dec_pairs = [(merge.decision(B[i]), merge.decision(C[i])) for i in sorted(ids)]
    batch_agreement = {"disposition": merge.agree(batch_disp_pairs),
                       "disposition_with_ft_code": merge.agree(batch_dec_pairs)}
    summary = json.loads(args.merge_summary.read_text(encoding="utf-8"))
    if not summary.get("final_complete") or not summary.get("prisma_final"):
        raise ValueError("merge is not final for paired screened reports; complete D adjudication first")
    if summary["n_complete_pairs"] != args.prior_screened + len(ids):
        raise ValueError(f"expected {args.prior_screened} prior plus {len(ids)} batch-3 screened pairs; "
                         f"merge has {summary['n_complete_pairs']}")
    with args.merged.open(newline="", encoding="utf-8") as fh:
        rows = {r["record_id"]: r for r in csv.DictReader(fh)}
    if ids - set(rows):
        raise ValueError("merged CSV omits batch-3 IDs")
    batch_final = Counter(rows[i]["final_disposition"] for i in ids)
    if "" in batch_final:
        raise ValueError("batch-3 final decisions incomplete")
    new_included = sorted(i for i in ids if rows[i]["final_disposition"] in merge.INCLUDE)
    all_included = sorted(i for i, r in rows.items() if r["final_disposition"] in merge.INCLUDE)
    old_included = set(all_included) - set(new_included)
    conflict_ids = [i for i in sorted(ids) if merge.decision(B[i]) != merge.decision(C[i])
                    or merge.rule_problems(B[i]) or merge.rule_problems(C[i])]
    batch_conflicts = len(conflict_ids)
    rates = {side: override_rates(path, ids) for side, path in (("B", args.b), ("C", args.c))}
    prefill = {side: prefill_dispositions(path, ids) for side, path in (("B", args.b), ("C", args.c))}
    for side in ("B", "C"):
        if sum(prefill[side].values()) != len(ids) or any(rates[side][f]["n"] != len(ids) for f in FIELDS):
            raise ValueError(f"incomplete batch-3 prefill/override log for reviewer {side}")
    final_review = {"B": dict(sorted(Counter(B[i]["your_disposition"] for i in ids).items())),
                    "C": dict(sorted(Counter(C[i]["your_disposition"] for i in ids).items()))}
    resolved_conflicts = dict(sorted(Counter(rows[i]["final_disposition"] for i in conflict_ids).items()))
    match = Counter()
    for i in conflict_ids:
        final_disposition = rows[i]["final_disposition"]
        final = final_disposition + ("|" + rows[i]["final_primary_code"] if final_disposition == "EXCLUDE" else "")
        bmatch = final == merge.decision(B[i])
        cmatch = final == merge.decision(C[i])
        match["both" if bmatch and cmatch else "B" if bmatch else "C" if cmatch else "neither"] += 1
    tokens = token_stats(args.ledger, ids, held_ids)
    stats = batch_agreement
    pf = summary["prisma_final"]
    lines = ["# Full-text screening batch 3", "",
             f"Batch-3 source list: {len(source_ids)} text files; processed within confirmed scope: {len(ids)}; "
             f"held outside scope: {len(held_ids)}; newly included: {len(new_included)}. "
             f"Paired screened across all batches: {summary['n_complete_pairs']}; "
             f"outside paired screening: {summary['n_without_paired_decision']}.",
             f"Scope deviation: the {len(held_ids)} held texts have human final layer M after V-to-M "
             "reclassification and are marked out_of_scope in the fetch manifest. "
             f"The nominal {args.prior_screened + len(source_ids)} available texts include these {len(held_ids)}; "
             f"{args.prior_screened + len(ids)} retrieved texts belong to the confirmed {summary['n_records']}-record scope. "
             "Held records are not full-text exclusions or adjudication conflicts.",
             f"Batch-3 B/C raw agreement: disposition {stats['disposition']['agree']}/{stats['disposition']['n']} "
             f"({stats['disposition']['raw_agreement']:.1%}), Cohen's kappa {stats['disposition']['kappa']}; "
             f"disposition with FT code {stats['disposition_with_ft_code']['agree']}/{stats['disposition_with_ft_code']['n']} "
             f"({stats['disposition_with_ft_code']['raw_agreement']:.1%}), "
             f"kappa {stats['disposition_with_ft_code']['kappa']}. Batch-3 B/C conflicts: {batch_conflicts}.",
             f"D conflicts across all batches: {summary['n_conflicts_for_D']}; "
             f"resolved: {summary['n_resolved_by_D']}; pending: {summary['n_unresolved_screened_conflicts']}.",
             f"Batch-3 prefill dispositions: B {json.dumps(prefill['B'], sort_keys=True)}; "
             f"C {json.dumps(prefill['C'], sort_keys=True)}.",
             f"Batch-3 final reviewer dispositions: B {json.dumps(final_review['B'], sort_keys=True)}; "
             f"C {json.dumps(final_review['C'], sort_keys=True)}.",
             f"Batch-3 D conflict resolutions by final disposition: {json.dumps(resolved_conflicts, sort_keys=True)}; "
             f"matched reviewer disposition and FT code: {json.dumps(dict(sorted(match.items())), sort_keys=True)}.",
             "", "## Final screened counts", "",
             f"Included: {pf['included_reports']} (stream A {pf['stream_A_reports']}, "
             f"stream B {pf['stream_B_reports']}); previously included outside batch 3: {len(old_included)}.",
             "Included record IDs across screened reports: " + ", ".join(all_included) + ".",
             f"Batch-3 final dispositions: {json.dumps(dict(sorted(batch_final.items())), sort_keys=True)}.",
             "", "| Disposition | Count |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in pf["by_disposition"].items()]
    lines += ["", "| Exclusion code | Count |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in pf["excluded_by_ft_code"].items()]
    lines += ["", "## Batch-3 override rates", "",
              "Model-reported agree/override flags are self-reports from the second read. Cell changes compare "
              "old and new workbook values after blank normalization. Comment and age-rule formatting, and "
              "human evidence-warning annotations, can change cells without changing a scientific decision.", "",
              "| Field | B model-reported override / n | B cell changed / n | C model-reported override / n | C cell changed / n |",
              "|---|---:|---:|---:|---:|"]
    for f in FIELDS:
        b, c = rates["B"][f], rates["C"][f]
        lines.append(f"| {f} | {b['model_reported_overrides']}/{b['n']} ({b['model_reported_rate']:.1%}) | "
                     f"{b['cell_changes']}/{b['n']} ({b['cell_change_rate']:.1%}) | "
                     f"{c['model_reported_overrides']}/{c['n']} ({c['model_reported_rate']:.1%}) | "
                     f"{c['cell_changes']}/{c['n']} ({c['cell_change_rate']:.1%}) |")
    t = tokens["totals"]
    lines += ["", "## Model usage", "",
              f"Ledger calls: {t.get('calls', 0)} (valid {t.get('valid_calls', 0)}; "
              f"usage unavailable on {t.get('calls_without_usage', 0)}). "
              f"Measured input {t.get('input_tokens', 0):,}, cached input {t.get('cached_input_tokens', 0):,}, "
              f"output {t.get('output_tokens', 0):,}, of which reasoning {t.get('reasoning_tokens', 0):,}.",
              f"Measured lower bound for billable input plus output: "
              f"{t.get('billable_tokens_measured_lower_bound', 0):,}; "
              f"lower-bound mean {tokens['mean_tokens_per_source_record_lower_bound']:,.1f} "
              f"per original source-list record; processed-record mean "
              f"{tokens['mean_tokens_per_processed_record_lower_bound']:,.1f} "
              f"for the {len(ids)} in-scope records (processed calls only). "
              f"The calibrated weekly-quota share is at least "
              f"{tokens['estimated_weekly_quota_pct_lower_bound']:.3f}% "
              "using 1.01M tokens per 1%; the total share is unknown while usage is missing.",
              "The ledger covers recorded screening calls and probes. Coordinator and script-development "
              "agent usage is not exposed here and is excluded from these figures.",
              f"Token use by scope group (including held calls and probes): "
              f"{json.dumps(tokens['by_scope_group'], sort_keys=True)}.",
              f"Per-stage token totals: {json.dumps(tokens['by_stage'], sort_keys=True)}.",
              "", "| Record ID | Measured input + output tokens | Calls without usage |", "|---|---:|---:|"]
    for i, usage in tokens["per_record_tokens"].items():
        measured = (f"≥{usage['measured_lower_bound']:,} (unknown remainder)"
                    if usage["calls_without_usage"] else
                    f"{usage['measured_lower_bound']:,}" if usage["calls"] else "no call recorded")
        lines.append(f"| {i} | {measured} | {usage['calls_without_usage']} |")
    lines += ["", "## Deviations and privacy", "",
              "The initial 55 B prefill calls used the earlier identity sanitizer. Subsequent calls used "
              "a stronger sanitizer that also masks article-author front matter. The earlier calls cannot be "
              "retracted, and whether any article-author strings were included in them is unknown. No "
              "actual author strings are recorded here.",
              "The per-call usage remains in the token ledger. B/C comments and D evidence remain in ignored workbooks.", ""]
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.newly_included.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text("\n".join(lines), encoding="utf-8")
    args.newly_included.write_text("\n".join(new_included) + ("\n" if new_included else ""), encoding="utf-8")
    return {"screened": summary["n_complete_pairs"], "included": pf["included_reports"],
            "newly_included": len(new_included),
            "billable_tokens_measured_lower_bound": t.get("billable_tokens_measured_lower_bound", 0),
            "calls_without_usage": t.get("calls_without_usage", 0)}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="action", required=True)
    p = sub.add_parser("scope-ids")
    for name in ("ids-file", "scope-file", "fetch-manifest", "human-final", "text-dir", "out", "hold", "audit"):
        p.add_argument("--" + name, required=True, type=Path)
    p = sub.add_parser("prepare")
    for name in ("b", "c", "ids-file", "out"):
        p.add_argument("--" + name, required=True, type=Path)
    p.add_argument("--template", type=Path, help="original D workbook; defaults to ft_conflicts_for_D.xlsx beside --out")
    p = sub.add_parser("combine")
    for name in ("old", "batch3", "out"):
        p.add_argument("--" + name, required=True, type=Path)
    p = sub.add_parser("report")
    for name in ("b", "c", "ids-file", "source-ids-file", "held-ids-file", "merge-summary", "merged", "ledger", "out", "newly-included"):
        p.add_argument("--" + name, required=True, type=Path)
    p.add_argument("--prior-screened", type=int, default=190)
    args = ap.parse_args()
    try:
        result = {"scope-ids": scope_ids, "prepare": prepare, "combine": combine, "report": report}[args.action](args)
    except (ValueError, OSError) as exc:
        ap.error(str(exc))
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
