#!/usr/bin/env python3
"""Merge two independent reviewer screening exports and summarise agreement and PRISMA counts.

Subcommands
  merge    --stage TI|TA|FT --a A.(csv|xlsx) --b B.(csv|xlsx) [--reconciled R.csv] [--calibration-ids ids.txt]
           [--planned-size N] [--population merged_TI.csv (TA) | merged_TA.csv (FT)] --out-dir DIR
           -> DIR/merged_<stage>.csv, DIR/conflicts_<stage>.csv, DIR/agreement_<stage>.json
           TI (v3.1, PRE-005: title screening): the final disposition is computed by the advance-wins rule
           (build_workbooks.ti_consensus); conflicts are listed and counted but need no reconciliation.
  prisma   [--master records_master.(csv|xlsx)] [--ti merged_TI.csv] [--ta merged_TA.csv] [--ft merged_FT.csv]
           [--master-ta stage-2 records_master.(csv|xlsx)] --out prisma.csv
  selftest  synthetic example (temporary files only): checks Python kappa against a textbook value and
            against the screening workbook formulas recalculated by LibreOffice (skipped if soffice is absent).

Inputs: the reviewer sheets of templates_xlsx/screening_workbook.xlsx (screen_<stage>_reviewer_A/B), saved by
Excel/LibreOffice so formula results are cached, or CSV exports of those sheets. Column names come from
04_screening/screening_log_template.json (record_id, decision, primary_reason | primary_exclusion_code); the TI
closed-list codes come from 04_screening/fulltext_exclusion_codes.json (title_stage_closed_list).
Code lists come from the JSON templates via build_workbooks.screening_enums().

Rules: values outside the code lists stop the merge; the two exports must contain the same record_ids;
nothing is imputed. A count is written only when the stage it belongs to is complete; otherwise the
value is blank (never 0) and the status column says NOT_YET_PERFORMED or INCOMPLETE.
"""
from __future__ import annotations

import argparse
import contextlib
import csv
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_workbooks as bw  # noqa: E402

import openpyxl  # noqa: E402


class MergeError(Exception):
    pass


# ---------------------------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------------------------
def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = str(v).strip()
    return v or None


def read_table(path: Path, sheet: str | None = None) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(newline='', encoding='utf-8-sig') as fh:
            return [{k.strip(): _clean(v) for k, v in row.items() if k} for row in csv.DictReader(fh)]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [_clean(h) for h in next(it)]
    rows = []
    for vals in it:
        row = {h: _clean(v) for h, v in zip(head, vals) if h}
        if any(v is not None for v in row.values()):
            rows.append(row)
    wb.close()
    return rows


def default_sheet(path: Path, stage: str, reviewer: str) -> str | None:
    if Path(path).suffix.lower() == '.csv':
        return None
    names = openpyxl.load_workbook(path, read_only=True).sheetnames
    want = f'screen_{stage}_reviewer_{reviewer}'
    return want if want in names else None


def stage_cfg(E: dict, stage: str) -> dict:
    if stage == 'TI':
        if not E.get('ti'):
            raise MergeError('the screening templates define no TI stage (screening_disposition_stages)')
        return {'decisions': E['ti'], 'reason': 'primary_reason', 'reasons': E['ti_codes'],
                'group': E['ti_binary_map'], 'group_labels': E['ti_binary_labels'], 'sci_excl': []}
    if stage == 'TA':
        return {'decisions': E['ta'], 'reason': 'primary_reason', 'reasons': E['codes'],
                'group': E['binary_map'], 'group_labels': E['binary_labels'], 'sci_excl': []}
    gmap = {d: g for g, ds in E['ft_groups'].items() for d in ds}
    return {'decisions': E['ft'], 'reason': 'primary_exclusion_code', 'reasons': E['codes'],
            'group': gmap, 'group_labels': list(E['ft_groups']), 'sci_excl': E['ft_excl']}


# ---------------------------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------------------------
def cohen_kappa(pairs) -> dict:
    """Unweighted Cohen's kappa for a list of (a, b) category pairs (same algorithm as the workbook)."""
    n = len(pairs)
    if n == 0:
        return {'n': 0, 'agreements': 0, 'raw_agreement': None, 'p_e': None, 'kappa': None, 'note': 'no jointly coded records'}
    agree = sum(1 for a, b in pairs if a == b)
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    po = agree / n
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n ** 2
    if abs(pe - 1.0) < 1e-12:
        return {'n': n, 'agreements': agree, 'raw_agreement': po, 'p_e': pe, 'kappa': None,
                'note': 'kappa undefined (p_e = 1: no variation)'}
    return {'n': n, 'agreements': agree, 'raw_agreement': po, 'p_e': pe, 'kappa': (po - pe) / (1 - pe), 'note': ''}


def merge(rows_a, rows_b, stage: str, E: dict, reconciled=None, population=None):
    cfg = stage_cfg(E, stage)
    rf = cfg['reason']

    def index(rows, who):
        out = OrderedDict()
        for r in rows:
            rid = r.get('record_id')
            if not rid:
                continue
            if rid in out:
                raise MergeError(f'reviewer {who}: duplicate record_id {rid}')
            out[rid] = r
        return out

    A, B = index(rows_a, 'A'), index(rows_b, 'B')
    if set(A) != set(B):
        only_a, only_b = sorted(set(A) - set(B)), sorted(set(B) - set(A))
        raise MergeError(f'record_id sets differ: only in A {only_a[:20]} ({len(only_a)}); only in B {only_b[:20]} ({len(only_b)})')
    bad = []
    for who, X in (('A', A), ('B', B)):
        for rid, r in X.items():
            d, c = r.get('decision'), r.get(rf)
            if d is not None and d not in cfg['decisions']:
                bad.append(f'{who}:{rid}: decision {d!r} not in {stage} list')
            if c is not None and c not in cfg['reasons']:
                bad.append(f'{who}:{rid}: {rf} {c!r} not in code list')
            if stage == 'TI' and d in E['ti_excl'] and c is None:
                bad.append(f'{who}:{rid}: {d} without a closed-list reason')
            if stage == 'TI' and d is not None and d not in E['ti_excl'] and c is not None:
                bad.append(f'{who}:{rid}: {rf} {c!r} given with {d} (only {"/".join(E["ti_excl"])} takes a reason)')
    recon = {}
    for r in reconciled or []:
        rid = r.get('record_id')
        if not rid:
            continue
        if rid not in A:
            bad.append(f'reconciled: {rid} not in reviewer exports')
        fd, fc = r.get('final_disposition'), r.get('final_primary_exclusion_code')
        if fd is not None and fd not in cfg['decisions']:
            bad.append(f'reconciled:{rid}: final_disposition {fd!r} not in {stage} list')
        if fc is not None and fc not in cfg['reasons']:
            bad.append(f'reconciled:{rid}: final code {fc!r} not in code list')
        recon[rid] = (fd, fc)
    if bad:
        raise MergeError('invalid values:\n  ' + '\n  '.join(bad[:50]))

    ids = list(A)
    if population is not None and stage in ('TA', 'FT'):
        pop = set(population)
        missing = sorted(pop - set(ids))
        if missing:
            raise MergeError(f'population records missing from {stage} exports: {missing[:20]}')
        ids = [i for i in ids if i in pop]
    elif stage == 'FT':
        ids = [i for i in ids if A[i].get('decision') or B[i].get('decision')]

    merged, exact_pairs, group_pairs = [], [], []
    for rid in ids:
        a, b = A[rid], B[rid]
        da, db, ca, cb = a.get('decision'), b.get('decision'), a.get(rf), b.get(rf)
        both = da is not None and db is not None
        exact = (1 if da == db else 0) if both else None
        ga = cfg['group'].get(da) if da else None
        gb = cfg['group'].get(db) if db else None
        gagree = (1 if ga == gb else 0) if both else None
        code_agree = reason_agree = None
        if stage == 'FT' and da in cfg['sci_excl'] and db in cfg['sci_excl']:
            code_agree = 1 if ca == cb else 0
        if stage == 'TI' and da in E['ti_excl'] and db in E['ti_excl']:
            reason_agree = 1 if ca == cb else 0
        conflict = 1 if (exact == 0 or code_agree == 0 or reason_agree == 0) else None
        if both:
            exact_pairs.append((da, db))
            group_pairs.append((ga, gb))
        final, fcode, fsrc = None, None, None
        if rid in recon and recon[rid][0]:
            final, fcode, fsrc = recon[rid][0], recon[rid][1], 'reconciled'
        elif stage == 'TI' and both:
            final, fcode, fsrc = bw.ti_consensus(da, db, ca, cb, E)   # advance wins; never unresolved
        elif both and not conflict:
            final, fcode, fsrc = da, (ca if da in cfg['sci_excl'] else None), 'agreement'
        elif both:
            fsrc = 'unresolved_conflict'
        elif da or db:
            fsrc = 'awaiting_second_decision'
        else:
            fsrc = 'not_decided'
        row = OrderedDict([('record_id', rid), ('decision_A', da), ('decision_B', db), (f'{rf}_A', ca), (f'{rf}_B', cb),
                           ('both_decided', 1 if both else None), ('exact_agree', exact),
                           ('group_A', ga), ('group_B', gb), ('group_agree', gagree)])
        if stage == 'FT':
            row['code_agree'] = code_agree
        if stage == 'TI':
            row['reason_agree'] = reason_agree
        row.update([('conflict', conflict), ('final_disposition', final)])
        if stage == 'FT':
            row['final_primary_exclusion_code'] = fcode
        if stage == 'TI':
            row['final_primary_reason'] = fcode
        row['final_source'] = fsrc
        merged.append(row)

    stats = OrderedDict()
    stats['stage'] = stage
    stats['records_in_stage'] = len(merged)
    stats['decided_A'] = sum(1 for r in merged if r['decision_A'])
    stats['decided_B'] = sum(1 for r in merged if r['decision_B'])
    stats['jointly_decided'] = len(exact_pairs)
    stats['conflicts'] = sum(1 for r in merged if r['conflict'])
    stats['unresolved'] = sum(1 for r in merged if not r['final_disposition'])
    stats['counts_A'] = {d: sum(1 for r in merged if r['decision_A'] == d) for d in cfg['decisions']}
    stats['counts_B'] = {d: sum(1 for r in merged if r['decision_B'] == d) for d in cfg['decisions']}
    if stage == 'FT':
        for who in 'AB':
            stats[f'scientific_exclusions_by_code_{who}'] = {
                c: sum(1 for r in merged if r[f'decision_{who}'] in cfg['sci_excl'] and r[f'{rf}_{who}'] == c)
                for c in cfg['reasons']}
    if stage == 'TI':
        stats['consensus_rule'] = 'advance wins (screening_log_template field_schema.title_stage_rules.consensus_rule)'
        stats['consensus_counts'] = {d: sum(1 for r in merged if r['final_disposition'] == d) for d in cfg['decisions']}
        stats['exclusions_by_code_consensus'] = {
            c: sum(1 for r in merged if r['final_disposition'] in E['ti_excl'] and r['final_primary_reason'] == c)
            for c in cfg['reasons']}
        stats['closed_list_ft_codes'] = {c: E['ti_code_info'][c]['ft_code'] for c in cfg['reasons']}
    stats['exact_disposition'] = cohen_kappa(exact_pairs)
    stats['group_disposition'] = cohen_kappa(group_pairs)
    stats['group_labels'] = cfg['group_labels']
    return merged, stats


def calibration(merged, ids, E, stage: str = 'TA', planned: int | None = None) -> dict:
    by = {r['record_id']: r for r in merged}
    cal = E['ti_calibration'] if stage == 'TI' else E['calibration']
    thr = E['ti_threshold'] if stage == 'TI' else E['threshold']
    labels = E['ti_binary_labels'] if stage == 'TI' else E['binary_labels']
    planned = int(planned or cal['planned_sample_size'])
    found = [i for i in ids if i in by]
    pairs = [(by[i]['group_A'], by[i]['group_B']) for i in found if by[i]['group_A'] and by[i]['group_B']]
    k = cohen_kappa(pairs)
    out = {'planned_sample_size': planned, 'ids_given': len(ids), 'ids_not_in_pool': sorted(set(ids) - set(by)),
           'jointly_assessed': k['n'], 'agreements': k['agreements'], 'raw_agreement': k['raw_agreement'],
           'kappa_descriptive': k['kappa'], 'kappa_note': k['note'], 'threshold': thr,
           'binary_groups': labels, 'stage': stage}
    if k['n'] < planned:
        out['status'] = f'INCOMPLETE: {k["n"]}/{planned} jointly assessed'
    elif k['raw_agreement'] >= thr:
        out['status'] = ('RAW_AGREEMENT_CRITERION_MET (pass also requires every conceptual disagreement to be '
                         'resolved and documented by the reviewers)') if stage == 'TA' else 'PASS (raw agreement criterion)'

    else:
        out['status'] = 'NOT_PASSED: clarify manual and run a fresh round (new seed, unseen records)'
    return out


# ---------------------------------------------------------------------------------------------
# PRISMA summary
# ---------------------------------------------------------------------------------------------
def prisma(E, master_rows=None, ta_rows=None, ft_rows=None, ti_rows=None, ta_master_rows=None):
    """Return a list of (section, item, value, status, note); value is None unless status == COMPLETE.
    ti_rows: merged_TI rows (v3.1 title stage); ta_master_rows: stage-2 records_master (abstract_source)."""
    admin = set(E['admin'])
    ta_excl = E['ta_excl']
    ta_sought = [d for d in E['ta'] if d not in admin and d not in ta_excl]
    not_retrieved = [s for s in E['admin'] if 'NOT_RETRIEVED' in s]   # name must exist in fulltext_exclusion_codes.json
    if len(not_retrieved) != 1:
        raise MergeError('expected exactly one NOT_RETRIEVED administrative status in fulltext_exclusion_codes.json')
    out = []

    def add(section, item, value, status, note=''):
        out.append((section, item, value if status == 'COMPLETE' else None, status, note))

    # identification
    mrows = [r for r in (master_rows or []) if r.get('record_id')]
    st = 'COMPLETE' if mrows else 'NOT_YET_PERFORMED'
    add('identification', 'records_identified', len(mrows), st, 'rows in records_master (all imported records)')
    if mrows:
        by_src = Counter()
        for r in mrows:
            for s in (r.get('source_database') or 'source not recorded').split(';'):
                by_src[s.strip()] += 1
        for s, n in sorted(by_src.items()):
            add('identification', f'records_identified_by_source:{s}', n, st, 'a record listing several sources counts once per source')
    dups = [r for r in mrows if r.get('dedup_status') in E['duplicates']]
    add('identification', 'duplicates_removed', len(dups), st)
    for s in E['duplicates']:
        add('identification', f'duplicates_removed:{s}', sum(1 for r in dups if r['dedup_status'] == s), st)

    # title screening (stage 1; v3.1, PRE-005)
    ti_progress = []
    if E.get('ti'):
        ex = E['ti_excl'][0]
        ti_progress = [d for d in E['ti'] if d not in E['ti_excl']]
        if not ti_rows:
            ti_st, ti_note = 'NOT_YET_PERFORMED', ''
        else:
            unresolved = sum(1 for r in ti_rows if not r.get('final_disposition'))
            ti_st = 'COMPLETE' if unresolved == 0 else 'INCOMPLETE'
            ti_note = '' if unresolved == 0 else f'{unresolved} of {len(ti_rows)} records without a final TI disposition'
        ti_list = ti_rows or []
        fti = Counter(r.get('final_disposition') for r in ti_list)
        no_reason = sum(1 for r in ti_list if r.get('final_disposition') == ex and not r.get('final_primary_reason'))
        add('screening_title', 'records_screened_TI', len(ti_list), ti_st, ti_note or 'stage 1: title, source, year, document type only')
        add('screening_title', 'records_excluded_TI', fti[ex], ti_st, ti_note)
        for c in E['ti_codes']:
            add('screening_title', f'records_excluded_TI:{c}', sum(1 for r in ti_list if r.get('final_disposition') == ex
                                                                     and r.get('final_primary_reason') == c),
                ti_st if not no_reason else 'INCOMPLETE',
                ti_note or (f'{no_reason} title exclusions lack a code' if no_reason else f'maps to {E["ti_code_info"][c]["ft_code"]}'))
        add('screening_title', 'records_advanced_to_abstract_stage', sum(fti[d] for d in ti_progress), ti_st,
            ti_note or f'final TI disposition in {ti_progress}')
        add('screening_title', f'TI_{E["ti_fallback"]}', fti[E['ti_fallback']], ti_st,
            'proceeds to stage 2; not an exclusion')
        add('screening_title', 'conflicts_TI_logged', sum(1 for r in ti_list if r.get('conflict') in (1, '1')), ti_st,
            'resolved by the advance-wins rule; not adjudicated')

    # title/abstract
    if not ta_rows:
        ta_st, ta_note = 'NOT_YET_PERFORMED', ''
    else:
        unresolved = sum(1 for r in ta_rows if not r.get('final_disposition'))
        ta_st = 'COMPLETE' if unresolved == 0 else 'INCOMPLETE'
        ta_note = '' if unresolved == 0 else f'{unresolved} of {len(ta_rows)} records without a final TA disposition'
    ta_rows = ta_rows or []
    fin = Counter(r.get('final_disposition') for r in ta_rows)
    add('screening', 'records_screened_TA', len(ta_rows), ta_st, ta_note)
    add('screening', 'records_excluded_TA', sum(fin[d] for d in ta_excl), ta_st, ta_note)
    if E.get('ti'):
        if not ta_rows:
            add('screening', 'records_screened_TA_without_abstract', None, 'NOT_YET_PERFORMED')
        elif not ta_master_rows:
            add('screening', 'records_screened_TA_without_abstract', None, 'INCOMPLETE',
                'stage-2 records_master (--master-ta) not supplied')
        else:
            src = {r.get('record_id'): r.get('abstract_source') for r in ta_master_rows}
            add('screening', 'records_screened_TA_without_abstract',
                sum(1 for r in ta_rows if src.get(r.get('record_id')) == 'abstract_unavailable'), ta_st,
                ta_note or 'abstract_source = abstract_unavailable; screened on title and available text, not excluded')
    for d in E['ta']:
        if d not in ta_excl and d not in ta_sought:
            add('screening', f'TA_{d}', fin[d], ta_st, 'not a scientific exclusion; tracked separately')
    add('retrieval', 'reports_sought_for_retrieval', sum(fin[d] for d in ta_sought), ta_st,
        ta_note or f'final TA disposition in {ta_sought}')

    # full text
    if not ft_rows:
        ft_st, ft_note = 'NOT_YET_PERFORMED', ''
    else:
        unresolved = sum(1 for r in ft_rows if not r.get('final_disposition'))
        ft_st = 'COMPLETE' if unresolved == 0 else 'INCOMPLETE'
        ft_note = '' if unresolved == 0 else f'{unresolved} of {len(ft_rows)} reports without a final FT disposition'
    ft_rows = ft_rows or []
    ff = Counter(r.get('final_disposition') for r in ft_rows)
    nr = ff[not_retrieved[0]]
    add('retrieval', 'reports_not_retrieved', nr, ft_st, ft_note)
    add('eligibility', 'reports_assessed_FT', len(ft_rows) - nr, ft_st, ft_note or 'FT reports minus not retrieved')
    excl = [r for r in ft_rows if r.get('final_disposition') in E['ft_excl']]
    no_code = sum(1 for r in excl if not r.get('final_primary_exclusion_code'))
    add('eligibility', 'reports_excluded_FT_total', len(excl), ft_st if not no_code else 'INCOMPLETE',
        ft_note or (f'{no_code} exclusions lack a final FT code' if no_code else ''))
    for c in E['codes']:
        add('eligibility', f'reports_excluded_FT:{c}',
            sum(1 for r in excl if r.get('final_primary_exclusion_code') == c), ft_st if not no_code else 'INCOMPLETE', ft_note)
    for s in E['admin']:
        if s not in not_retrieved:
            add('eligibility', f'FT_{s}', ff[s], ft_st, 'administrative status, not a scientific exclusion')
    incl = E['ft_incl']
    add('included', 'reports_included_total', sum(ff[d] for d in incl), ft_st,
        ft_note or 'reports; study/cohort counts require report_cohort_links')
    for d in incl:
        add('included', f'reports_{d}', ff[d], ft_st, ft_note)
    if E.get('ti') and ti_rows and ti_st == 'COMPLETE' and ta_st == 'COMPLETE':
        adv = sum(fti[d] for d in ti_progress)
        if adv != len(ta_rows):
            add('checks', 'advanced_TI_vs_TA_rows', None, 'CHECK', f'advanced from TI={adv} but TA export has {len(ta_rows)} records')
    if ta_st == 'COMPLETE' and ft_st == 'COMPLETE':
        sought = sum(fin[d] for d in ta_sought)
        if sought != len(ft_rows):
            add('checks', 'sought_vs_FT_rows', None, 'CHECK', f'sought={sought} but FT export has {len(ft_rows)} reports')
    return out


# ---------------------------------------------------------------------------------------------
# Output helpers
# ---------------------------------------------------------------------------------------------
def write_csv(path: Path, rows, header=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        if rows and isinstance(rows[0], dict):
            header = header or list(rows[0].keys())
            w.writerow(header)
            for r in rows:
                w.writerow(['' if r.get(h) is None else r.get(h) for h in header])
        else:
            if header:
                w.writerow(header)
            for r in rows:
                w.writerow(['' if v is None else v for v in r])


def cmd_merge(args, E):
    sa = args.sheet_a or default_sheet(args.a, args.stage, 'A')
    sb = args.sheet_b or default_sheet(args.b, args.stage, 'B')
    ra, rb = read_table(args.a, sa), read_table(args.b, sb)
    for who, rows in (('A', ra), ('B', rb)):
        if rows and all(r.get('record_id') is None for r in rows):
            raise MergeError(f'reviewer {who}: record_id column is empty (formula values not cached?). '
                             'Open and save the workbook in Excel/LibreOffice, or export the sheet to CSV.')
    recon = read_table(args.reconciled) if args.reconciled else None
    pop = None
    if args.population:
        if args.stage == 'TI':
            raise MergeError('--population applies to TA (merged_TI.csv) and FT (merged_TA.csv) only')
        if args.stage == 'TA':   # stage-2 population = final TI disposition that progresses (v3.1)
            keep = [d for d in E['ti'] if d not in E['ti_excl']]
        else:
            keep = [d for d in E['ta'] if d not in E['admin'] and d not in E['ta_excl']]
        pop = [r['record_id'] for r in read_table(args.population) if r.get('final_disposition') in keep]
    merged, stats = merge(ra, rb, args.stage, E, recon, pop)
    if args.planned_size and not args.calibration_ids:
        raise MergeError('--planned-size requires --calibration-ids')
    if args.calibration_ids:
        p = Path(args.calibration_ids)
        ids = ([r['record_id'] for r in read_table(p) if r.get('record_id')] if p.suffix.lower() in ('.csv', '.xlsx')
               else [x.strip() for x in p.read_text(encoding='utf-8').splitlines() if x.strip()])
        if args.stage not in ('TA', 'TI'):
            raise MergeError('--calibration-ids applies to the TI and TA stages only')
        stats['calibration'] = calibration(merged, ids, E, args.stage, args.planned_size)
    out = Path(args.out_dir)
    write_csv(out / f'merged_{args.stage}.csv', merged)
    write_csv(out / f'conflicts_{args.stage}.csv', [r for r in merged if r['conflict']],
              header=list(merged[0].keys()) if merged else ['record_id'])
    (out / f'agreement_{args.stage}.json').write_text(json.dumps(stats, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    e, g = stats['exact_disposition'], stats['group_disposition']
    print(json.dumps({'stage': args.stage, 'records': stats['records_in_stage'], 'jointly_decided': stats['jointly_decided'],
                      'conflicts': stats['conflicts'], 'unresolved': stats['unresolved'],
                      'raw_agreement_exact': e['raw_agreement'], 'kappa_exact': e['kappa'],
                      'raw_agreement_group': g['raw_agreement'], 'kappa_group': g['kappa'],
                      'outputs': str(out)}, ensure_ascii=False))


def cmd_prisma(args, E):
    master = None
    if args.master:
        sheet = 'records_master' if Path(args.master).suffix.lower() == '.xlsx' else None
        master = read_table(args.master, sheet)
    ta_master = None
    if args.master_ta:
        ta_master = read_table(args.master_ta, 'records_master' if Path(args.master_ta).suffix.lower() == '.xlsx' else None)
    rows = prisma(E, master, read_table(args.ta) if args.ta else None, read_table(args.ft) if args.ft else None,
                  read_table(args.ti) if args.ti else None, ta_master)
    write_csv(Path(args.out), rows, header=['section', 'item', 'value', 'status', 'note'])
    for r in rows:
        print(f'{r[0]:15} {r[1]:45} {"" if r[2] is None else r[2]:>6}  {r[3]}  {r[4]}')


# ---------------------------------------------------------------------------------------------
# Self-test (synthetic data in a temporary directory only)
# ---------------------------------------------------------------------------------------------
def _recalc(path: Path, outdir: Path) -> Path | None:
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return None
    profile = Path(tempfile.mkdtemp(prefix='lo_profile_'))
    subprocess.run([soffice, f'-env:UserInstallation=file://{profile}', '--headless', '--calc', '--convert-to', 'xlsx',
                    '--outdir', str(outdir), str(path)], check=True, capture_output=True, timeout=180)
    shutil.rmtree(profile, ignore_errors=True)
    out = outdir / path.name
    return out if out.exists() else None


def _named(wb, name):
    dest = list(wb.defined_names[name].destinations)[0]
    return wb[dest[0]][dest[1].replace('$', '')].value


def _write_ids(path: Path, ids) -> Path:
    path.write_text('\n'.join(ids) + '\n', encoding='utf-8')
    return path


def _write_rows(path: Path, rows) -> Path:
    write_csv(path, rows)
    return path


def cmd_selftest(args, E):
    results = []

    def check(label, ok, detail=''):
        results.append((label, bool(ok), detail))

    # 1. textbook 2x2: [[20,5],[10,15]] -> p_o=.70, p_e=.50, kappa=.40
    pairs = [('y', 'y')] * 20 + [('y', 'n')] * 5 + [('n', 'y')] * 10 + [('n', 'n')] * 15
    k = cohen_kappa(pairs)
    check('python kappa textbook 2x2 = 0.40', abs(k['kappa'] - 0.4) < 1e-12 and abs(k['raw_agreement'] - 0.7) < 1e-12,
          f'kappa={k["kappa"]}')
    check('python kappa undefined when p_e = 1', cohen_kappa([('y', 'y')] * 5)['kappa'] is None)

    # 2. PRISMA blank rule
    rows = prisma(E)
    check('PRISMA: nothing performed -> all values blank', all(r[2] is None and r[3] == 'NOT_YET_PERFORMED' for r in rows))

    # 3. workbook formula cross-check on a synthetic example
    tmp = Path(tempfile.mkdtemp(prefix='screening_selftest_'))
    try:
        src = bw.load_sources()
        wbp = tmp / 'screening_workbook.xlsx'
        bw.build_screening(src, wbp, rows=30, conflict_rows=10)
        wb = openpyxl.load_workbook(wbp)
        ta, ft, codes, dup = E['ta'], E['ft'], E['codes'], E['duplicates'][0]
        n = 20
        ids = [f'SYNTH{i:03d}' for i in range(1, n + 1)]
        m = wb['records_master']
        for i, rid in enumerate(ids, start=2):
            m.cell(row=i, column=1, value=rid)
            m.cell(row=i, column=2, value='SYNTHETIC SELFTEST')
            m.cell(row=i, column=5, value='synthetic title')
        m.cell(row=n + 1, column=bw.MASTER_COLUMNS.index('dedup_status') + 1, value=dup)   # last record = duplicate
        # deterministic synthetic TA decisions with disagreements; record 19 undecided by B
        ta_a = [ta[(i * 7) % len(ta)] if i % 3 else ta[0] for i in range(n - 1)]
        ta_b = [ta_a[i] if i % 4 else ta[(i + 1) % len(ta)] for i in range(n - 1)]
        ta_b[-1] = None
        excl_ta = E['ta_excl'][0]

        def put(sheet, col_name, values, cols):
            ws = wb[sheet]
            c = cols.index(col_name) + 1
            for i, v in enumerate(values, start=2):
                ws.cell(row=i, column=c, value=v)

        SC = bw.screening_columns(E)
        ta_cols, ft_cols = SC['TA'], SC['FT']
        put('screen_TA_reviewer_A', 'decision', ta_a, ta_cols)
        put('screen_TA_reviewer_B', 'decision', ta_b, ta_cols)
        put('screen_TA_reviewer_A', 'primary_reason', [codes[i % len(codes)] if d == excl_ta else None for i, d in enumerate(ta_a)], ta_cols)
        put('screen_TA_reviewer_B', 'primary_reason', [codes[(i + 1) % len(codes)] if d == excl_ta else None for i, d in enumerate(ta_b)], ta_cols)
        # FT: first 14 records, include exclusions with codes and a code disagreement
        sci = E['ft_excl'][0]
        ft_a = [ft[(i * 3) % len(ft)] for i in range(14)] + [None] * 5
        ft_b = [ft_a[i] if i % 5 else ft[(i * 3 + 2) % len(ft)] for i in range(14)] + [None] * 5
        ft_a[2], ft_b[2] = sci, sci
        ft_a[3], ft_b[3] = sci, sci
        ft_a[6], ft_b[6] = sci, sci
        code_a = [codes[i % len(codes)] if d == sci else None for i, d in enumerate(ft_a)]
        code_b = [codes[i % len(codes)] if d == sci else None for i, d in enumerate(ft_b)]
        code_b[3] = codes[(3 + 2) % len(codes)]
        put('screen_FT_reviewer_A', 'decision', ft_a, ft_cols)
        put('screen_FT_reviewer_B', 'decision', ft_b, ft_cols)
        put('screen_FT_reviewer_A', 'primary_exclusion_code', code_a, ft_cols)
        put('screen_FT_reviewer_B', 'primary_exclusion_code', code_b, ft_cols)
        # TI (v3.1 title stage): advance-wins cases; record 19 (index 18) undecided by B
        has_ti = bool(E.get('ti'))
        if has_ti:
            adv, ex, fb, tc = E['ti_advance'], E['ti_excl'][0], E['ti_fallback'], E['ti_codes']
            ti_pairs = [(adv, None, adv, None), (ex, tc[0], ex, tc[0]), (ex, tc[2], ex, tc[1]), (adv, None, ex, tc[3]),
                        (ex, tc[4], fb, None), (fb, None, fb, None), (ex, tc[1], ex, tc[1]), (adv, None, adv, None),
                        (adv, None, fb, None), (ex, tc[3], ex, tc[3]), (adv, None, adv, None), (ex, tc[4], adv, None),
                        (adv, None, adv, None), (ex, tc[0], ex, tc[0]), (adv, None, adv, None), (fb, None, adv, None),
                        (adv, None, adv, None), (ex, tc[2], ex, tc[2]), (adv, None, None, None)]
            put('screen_TI_reviewer_A', 'decision', [x[0] for x in ti_pairs], SC['TI'])
            put('screen_TI_reviewer_A', 'primary_reason', [x[1] for x in ti_pairs], SC['TI'])
            put('screen_TI_reviewer_B', 'decision', [x[2] for x in ti_pairs], SC['TI'])
            put('screen_TI_reviewer_B', 'primary_reason', [x[3] for x in ti_pairs], SC['TI'])
        # calibration blocks: first 10 synthetic IDs (fewer than planned -> INCOMPLETE)
        for sheet in (['merge_TA', 'merge_TI'] if has_ti else ['merge_TA']):
            ws = wb[sheet]
            cal_cells = [c for c in ws.iter_rows() for c in c if c.protection.locked is False and c.number_format == '@']
            for cell, rid in zip(cal_cells, ids[:10]):
                cell.value = rid
        wb.save(wbp)
        rec = _recalc(wbp, tmp / 'recalc')
        if rec is None:
            check('workbook cross-check (LibreOffice not available)', True, 'SKIPPED')
        else:
            v = openpyxl.load_workbook(rec, data_only=True)
            # Python side reads the recalculated reviewer sheets (cached formula values) end-to-end
            for stage in ('TA', 'FT'):
                ra = read_table(rec, f'screen_{stage}_reviewer_A')
                rb = read_table(rec, f'screen_{stage}_reviewer_B')
                merged, st = merge(ra, rb, stage, E)
                grp = 'binary' if stage == 'TA' else 'group'
                for key, py in (('exact_kappa', st['exact_disposition']['kappa']),
                                ('exact_po', st['exact_disposition']['raw_agreement']),
                                (f'{grp}_kappa', st['group_disposition']['kappa']),
                                (f'{grp}_po', st['group_disposition']['raw_agreement']),
                                ('n_joint', st['jointly_decided']), ('n_conflict', st['conflicts'])):
                    xl = _named(v, f'{stage}_{key}')
                    ok = (py is None and xl in (None, '')) or (py is not None and isinstance(xl, (int, float)) and abs(xl - py) < 1e-9)
                    check(f'{stage} {key}: python vs workbook', ok, f'python={py} workbook={xl}')
                if stage == 'TA':
                    check('TA duplicate row hidden from reviewer sheets', st['records_in_stage'] == n - 1, f'{st["records_in_stage"]}')
                    cal = calibration(merged, ids[:10], E)
                    xl_k, xl_raw, xl_status = _named(v, 'TA_cal_kappa'), _named(v, 'TA_cal_raw'), _named(v, 'TA_cal_status')
                    pk = cal['kappa_descriptive']
                    okk = ((pk is None and (xl_k in (None, '') or str(xl_k).startswith('undefined'))) or
                           (pk is not None and isinstance(xl_k, (int, float)) and abs(xl_k - pk) < 1e-9))
                    check('calibration kappa: python vs workbook', okk, f'python={cal["kappa_descriptive"]} workbook={xl_k}')
                    check('calibration raw agreement: python vs workbook', abs((xl_raw or 0) - (cal['raw_agreement'] or 0)) < 1e-9,
                          f'python={cal["raw_agreement"]} workbook={xl_raw}')
                    check('calibration status INCOMPLETE below planned size', str(xl_status).startswith('INCOMPLETE')
                          and cal['status'].startswith('INCOMPLETE'), f'{xl_status} | {cal["status"]}')
                    ta_merged = merged
                else:
                    mism = [(who, c, st[f'scientific_exclusions_by_code_{who}'][c], _named(v, f'FT_excl_{who}_{c[:4]}'))
                            for c in codes for who in 'AB'
                            if st[f'scientific_exclusions_by_code_{who}'][c] != _named(v, f'FT_excl_{who}_{c[:4]}')]
                    total = sum(st['scientific_exclusions_by_code_A'].values()) + sum(st['scientific_exclusions_by_code_B'].values())
                    check('FT exclusions by code: python vs workbook', not mism and total > 0, f'mismatches={mism} total={total}')
                    ft_merged = merged
            # PRISMA with incomplete stages -> blanks
            master = read_table(rec, 'records_master')
            rows = prisma(E, master, ta_merged, ft_merged)
            d = {r[1]: r for r in rows}
            check('PRISMA: identified/duplicates complete from master',
                  d['records_identified'][2] == n and d['duplicates_removed'][2] == 1)
            check('PRISMA: incomplete TA -> blank, not 0', d['records_screened_TA'][2] is None
                  and d['records_screened_TA'][3] == 'INCOMPLETE')
            check('PRISMA: FT values blank while FT incomplete', all(r[2] is None for r in rows if r[0] in ('eligibility', 'included')))
            # resolve every open TA record (synthetic reconciliation) -> TA counts become available
            ra = read_table(rec, 'screen_TA_reviewer_A')
            recon = [{'record_id': r['record_id'], 'final_disposition': r['decision_A'] or r['decision_B']}
                     for r in ta_merged if not r['final_disposition']]
            ta_done, _ = merge(ra, read_table(rec, 'screen_TA_reviewer_B'), 'TA', E, recon)
            d = {r[1]: r for r in prisma(E, master, ta_done, None)}
            n_ex = sum(1 for r in ta_done if r['final_disposition'] in E['ta_excl'])
            check('PRISMA: reconciled TA -> counts filled', d['records_screened_TA'][2] == n - 1 and
                  d['records_excluded_TA'][2] == n_ex and d['records_screened_TA'][3] == 'COMPLETE',
                  f'screened={d["records_screened_TA"][2]} excluded={d["records_excluded_TA"][2]}')
            check('PRISMA: FT not supplied -> NOT_YET_PERFORMED blanks', d['reports_assessed_FT'][2] is None and
                  d['reports_assessed_FT'][3] == 'NOT_YET_PERFORMED')
            check('calibration --planned-size overrides the JSON sample size',
                  not calibration(ta_merged, ids[:10], E, 'TA', 10)['status'].startswith('INCOMPLETE'))
            if has_ti:
                ra_ti, rb_ti = read_table(rec, 'screen_TI_reviewer_A'), read_table(rec, 'screen_TI_reviewer_B')
                ti_merged, st = merge(ra_ti, rb_ti, 'TI', E)
                for key, py in (('exact_kappa', st['exact_disposition']['kappa']),
                                ('exact_po', st['exact_disposition']['raw_agreement']),
                                ('binary_kappa', st['group_disposition']['kappa']),
                                ('binary_po', st['group_disposition']['raw_agreement']),
                                ('n_joint', st['jointly_decided']), ('n_conflict', st['conflicts'])):
                    xl = _named(v, f'TI_{key}')
                    ok = (py is None and xl in (None, '')) or (py is not None and isinstance(xl, (int, float)) and abs(xl - py) < 1e-9)
                    check(f'TI {key}: python vs workbook', ok, f'python={py} workbook={xl}')
                mism = [(d_, st['consensus_counts'][d_], _named(v, f'TI_cons_{d_}')) for d_ in E['ti']
                        if st['consensus_counts'][d_] != _named(v, f'TI_cons_{d_}')]
                check('TI advance-wins consensus counts: python vs workbook', not mism, f'mismatches={mism} {st["consensus_counts"]}')
                mism = [(c, st['exclusions_by_code_consensus'][c], _named(v, f'TI_excl_{c[:4]}')) for c in tc
                        if st['exclusions_by_code_consensus'][c] != _named(v, f'TI_excl_{c[:4]}')]
                tot = sum(st['exclusions_by_code_consensus'].values())
                check('TI exclusions by closed-list code: python vs workbook', not mism and tot > 0, f'mismatches={mism} total={tot}')
                by = {r['record_id']: r for r in ti_merged}
                check('TI advance wins over an exclusion (conflict logged)',
                      by[ids[3]]['final_disposition'] == adv and by[ids[3]]['conflict'] == 1 and by[ids[11]]['final_disposition'] == adv)
                check('TI both excluded with different codes -> earlier closed-list code, conflict logged',
                      by[ids[2]]['final_disposition'] == ex and by[ids[2]]['final_primary_reason'] == tc[1] and by[ids[2]]['conflict'] == 1)
                check('TI exclusion vs awaiting -> awaiting classification (not excluded)',
                      by[ids[4]]['final_disposition'] == fb and by[ids[5]]['final_disposition'] == fb)
                check('TI conflicts never left unresolved', st['unresolved'] == 1 and
                      all(r['final_source'] != 'unresolved_conflict' for r in ti_merged), f'unresolved={st["unresolved"]}')
                cal_ti = calibration(ti_merged, ids[:10], E, 'TI')
                xl_raw, xl_k, xl_s = _named(v, 'TI_cal_raw'), _named(v, 'TI_cal_kappa'), _named(v, 'TI_cal_status')
                pk = cal_ti['kappa_descriptive']
                okk = ((pk is None and (xl_k in (None, '') or str(xl_k).startswith('undefined'))) or
                       (pk is not None and isinstance(xl_k, (int, float)) and abs(xl_k - pk) < 1e-9))
                check('TI calibration raw agreement and kappa: python vs workbook',
                      okk and abs((xl_raw or 0) - (cal_ti['raw_agreement'] or 0)) < 1e-9, f'python={cal_ti["raw_agreement"]}/{pk} workbook={xl_raw}/{xl_k}')
                check('TI calibration INCOMPLETE below the 100-record plan', str(xl_s).startswith('INCOMPLETE')
                      and cal_ti['status'].startswith('INCOMPLETE'), f'{xl_s} | {cal_ti["status"]}')
                d = {r[1]: r for r in prisma(E, master, None, None, ti_merged)}
                check('PRISMA: incomplete TI -> title box blank', d['records_screened_TI'][2] is None
                      and d['records_screened_TI'][3] == 'INCOMPLETE')
                rb_done = [dict(r) for r in rb_ti]
                for r in rb_done:
                    if r.get('record_id') == ids[18]:
                        r['decision'] = adv
                ti_done, _ = merge(ra_ti, rb_done, 'TI', E)
                d = {r[1]: r for r in prisma(E, master, None, None, ti_done)}
                n_ex = sum(1 for r in ti_done if r['final_disposition'] == ex)
                by_code = sum(d[f'records_excluded_TI:{c}'][2] for c in tc)
                check('PRISMA: complete TI -> title box filled by closed-list code',
                      d['records_screened_TI'][2] == n - 1 and d['records_excluded_TI'][2] == n_ex == by_code and
                      d['records_advanced_to_abstract_stage'][2] == n - 1 - n_ex and d['records_screened_TI'][3] == 'COMPLETE',
                      f'screened={d["records_screened_TI"][2]} excluded={n_ex} by_code={by_code}')
                ta_master = [{'record_id': ids[0], 'abstract_source': 'abstract_unavailable'},
                             {'record_id': ids[1], 'abstract_source': 'database_export'}]
                d = {r[1]: r for r in prisma(E, master, ta_done, None, ti_done, ta_master)}
                check('PRISMA: TA rows differing from advanced TI records -> CHECK row', 'advanced_TI_vs_TA_rows' in d)
                check('PRISMA: records screened at TA without abstract counted', d['records_screened_TA_without_abstract'][2] == 1)
            # CLI end-to-end on the recalculated workbook
            od = tmp / 'cli'
            rc = main(['merge', '--stage', 'TA', '--a', str(rec), '--b', str(rec), '--out-dir', str(od),
                       '--calibration-ids', str(_write_ids(tmp / 'ids.txt', ids[:10]))])
            agr = json.loads((od / 'agreement_TA.json').read_text())
            check('CLI merge writes merged/conflicts/agreement files', rc == 0 and (od / 'merged_TA.csv').exists()
                  and (od / 'conflicts_TA.csv').exists() and 'calibration' in agr)
            bad = _write_rows(tmp / 'bad.csv', [{'record_id': 'X1', 'decision': 'include'}])
            with contextlib.redirect_stderr(io.StringIO()):
                rc = main(['merge', '--stage', 'TA', '--a', str(bad), '--b', str(bad), '--out-dir', str(od)])
            check('CLI merge rejects values outside the JSON code list', rc == 2)
            if has_ti:
                oti = tmp / 'cli_ti'
                rc = main(['merge', '--stage', 'TI', '--a', str(rec), '--b', str(rec), '--out-dir', str(oti),
                           '--calibration-ids', str(_write_ids(tmp / 'ids_ti.txt', ids[:10]))])
                agr = json.loads((oti / 'agreement_TI.json').read_text())
                check('CLI merge --stage TI writes merged/conflicts/agreement files', rc == 0 and (oti / 'merged_TI.csv').exists()
                      and (oti / 'conflicts_TI.csv').exists() and agr['calibration']['stage'] == 'TI')
                prog = [d_ for d_ in E['ti'] if d_ not in E['ti_excl']]
                keep = [r['record_id'] for r in read_table(oti / 'merged_TI.csv') if r.get('final_disposition') in prog]
                ota = tmp / 'cli_ta_pop'
                rc = main(['merge', '--stage', 'TA', '--a', str(rec), '--b', str(rec), '--out-dir', str(ota),
                           '--population', str(oti / 'merged_TI.csv')])
                agr = json.loads((ota / 'agreement_TA.json').read_text())
                check('CLI merge --stage TA --population merged_TI.csv screens only TI-retained records',
                      rc == 0 and agr['records_in_stage'] == len(keep) < n - 1, f'{agr["records_in_stage"]} of {n - 1}')
                bad = _write_rows(tmp / 'bad_ti.csv', [{'record_id': 'X1', 'decision': E['ti_excl'][0], 'primary_reason': ''}])
                with contextlib.redirect_stderr(io.StringIO()):
                    rc = main(['merge', '--stage', 'TI', '--a', str(bad), '--b', str(bad), '--out-dir', str(oti)])
                check('CLI merge --stage TI rejects a title exclusion without a closed-list reason', rc == 2)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    width = max(len(r[0]) for r in results)
    for label, ok, detail in results:
        print(f'{"PASS" if ok else "FAIL"}  {label:{width}}  {detail}')
    failed = [r for r in results if not r[1]]
    print(f'selftest: {len(results) - len(failed)} passed, {len(failed)} failed')
    return 1 if failed else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    m = sub.add_parser('merge')
    m.add_argument('--stage', choices=['TI', 'TA', 'FT'], required=True)
    m.add_argument('--a', type=Path, required=True)
    m.add_argument('--b', type=Path, required=True)
    m.add_argument('--sheet-a')
    m.add_argument('--sheet-b')
    m.add_argument('--reconciled', type=Path, help='CSV: record_id, final_disposition[, final_primary_exclusion_code]')
    m.add_argument('--population', type=Path, help='TA: merged_TI.csv (population = final TI disposition that progresses); '
                                                    'FT: merged_TA.csv (population = final TA "sought" records)')
    m.add_argument('--calibration-ids', help='TI/TA: sampled record_ids (txt one per line, or CSV/XLSX with record_id)')
    m.add_argument('--planned-size', type=int, help='calibration round size if it differs from the JSON planned_sample_size '
                                                    '(e.g. 25 for the v3.1 TA re-calibration)')
    m.add_argument('--out-dir', type=Path, required=True)
    p = sub.add_parser('prisma')
    p.add_argument('--master', type=Path)
    p.add_argument('--ti', type=Path, help='merged_TI.csv (v3.1 title stage)')
    p.add_argument('--master-ta', type=Path, help='stage-2 records_master (abstract_source) for the without-abstract count')
    p.add_argument('--ta', type=Path)
    p.add_argument('--ft', type=Path)
    p.add_argument('--out', type=Path, required=True)
    sub.add_parser('selftest')
    args = ap.parse_args(argv)
    E = bw.screening_enums(bw.load_sources())
    try:
        if args.cmd == 'merge':
            cmd_merge(args, E)
        elif args.cmd == 'prisma':
            cmd_prisma(args, E)
        else:
            return cmd_selftest(args, E)
    except MergeError as e:
        print(f'ERROR: {e}', file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    sys.exit(main())
