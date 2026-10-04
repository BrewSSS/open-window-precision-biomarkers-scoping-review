#!/usr/bin/env python3
"""Build blank screening, extraction and critical-appraisal Excel workbooks from the
project's JSON templates (protocol v2 and later).

Rule: edit the JSON templates, then regenerate with this script. Never hand-edit the
structure (sheets, columns, dropdowns, protection) of the generated .xlsx files.

All enumerations are read from the JSON templates. Where the JSON is silent, a small number
of values are parsed from sentences in the Markdown manuals (TA dispositions, calibration
binary grouping and threshold); every such interpretation is listed in INTERPRETATIONS and
written into the README sheet of the workbook it affects. Only Excel-mechanics defaults
(TRUE/FALSE, the multi-value separator, group labels used inside formulas) are hard-coded.

No records, decisions or counts are written: every data row is empty.

Usage:
    python3 scripts/build_workbooks.py                       # screening + extraction + appraisal
    python3 scripts/build_workbooks.py --reviewer A --reviewer B   # also extraction_workbook_A/B
    python3 scripts/build_workbooks.py --only extraction --reviewer A
Options --rows (formula rows in the screening workbook, default 5000) and --out-dir.
After building, every workbook is reloaded and checked (sheet names, headers, data
validations, protection flags, empty data rows); use --no-verify to skip.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit('openpyxl is required: python3 -m pip install --user openpyxl')

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / 'templates_xlsx'
SCRIPT_VERSION = '1.0.0'
MAX_ROW = 1048576
MULTI_SEP = '; '                      # Excel-mechanics: separator for multi-valued cells
BOOLEAN_VALUES = ['TRUE', 'FALSE']    # Excel-mechanics default for JSON booleans

SOURCES = {
    'screening_log': '04_screening/screening_log_template.json',
    'ft_codes': '04_screening/fulltext_exclusion_codes.json',
    'screening_manual': '04_screening/screening_manual.md',
    'calibration_plan': '04_screening/calibration_plan.md',
    'search_log': '03_search/search_log_template.json',
    'dictionary': '05_extraction/data_dictionary.json',
    'extraction_template': '05_extraction/extraction_template.json',
    'extraction_manual': '05_extraction/extraction_manual.md',
    'appraisal_template': '05_extraction/critical_appraisal_template.json',
    'appraisal_manual': '05_extraction/critical_appraisal_manual.md',
}
SCREENING_SOURCES = ['screening_log', 'ft_codes', 'screening_manual', 'calibration_plan', 'search_log', 'dictionary']
EXTRACTION_SOURCES = ['dictionary', 'extraction_template', 'extraction_manual', 'screening_log', 'ft_codes']
APPRAISAL_SOURCES = ['appraisal_template', 'appraisal_manual']

# ---------------------------------------------------------------------------------------------
# Documented interpretations of JSON fields (written into README sheets).
# ---------------------------------------------------------------------------------------------
INTERPRETATIONS = {
    'screening': [
        'TA decision list = screening_log_template.field_schema.screening_dispositions filtered to the dispositions '
        'printed in bold in screening_manual.md section "B. Title/abstract" (the JSON does not tag dispositions by stage). '
        'The task wording "include / exclude / unclear" is implemented with these JSON values: ADVANCE ~ include/progress, '
        'EXCLUDE_TA ~ exclude, AWAITING_CLASSIFICATION ~ unclear, FRONTIER_PREPRINT = preprint routing.',
        'FT decision list = all screening_dispositions except the TA-only ones (TA dispositions that are not '
        'administrative statuses in fulltext_exclusion_codes.json).',
        'TA primary_reason dropdown re-uses the FT01-FT08 eligibility hierarchy (the manual requires "one primary reason" '
        'at TA but defines no separate TA reason codes).',
        'FT primary_exclusion_code dropdown contains only FT01-FT08. The administrative statuses (DUPLICATE_*, NOT_RETRIEVED, '
        'AWAITING_*, FRONTIER_PREPRINT) are offered in the FT decision dropdown instead, because screening_log_template '
        'record_validation_rules state that administrative statuses require a null primary exclusion code.',
        'Calibration binary grouping (progress / do not exclude vs EXCLUDE_TA) and the 80% raw-agreement threshold are parsed '
        'from calibration_plan.md; planned sample size, seed and sampled IDs come from screening_log_template.calibration.',
        'FT agreement is also summarised on a 3-group collapse: INCLUDE (non-administrative, non-exclusion FT dispositions), '
        'SCIENTIFIC_EXCLUSION (non-administrative dispositions containing "EXCLUDE"), ADMINISTRATIVE_STATUS '
        '(statuses listed in fulltext_exclusion_codes.json). The group labels are formula mechanics, not protocol codes.',
        'dedup_status uses the administrative statuses whose name starts with DUPLICATE; blank means retained/unique '
        '(no JSON label exists for a retained record). Rows flagged as duplicates are hidden from the reviewer sheets.',
        'records_master columns: record_id/source/route conventions follow screening_log_template and search_log_template; '
        'the bibliographic columns (title, authors, year, journal, doi, pmid, abstract) and dedup_group_id are not defined '
        'in any JSON template and follow the task specification.',
        'A_eligible / B_eligible use TRUE/FALSE (JSON booleans) plus NR/NA/UNCLEAR from data_dictionary.json.',
        'FT "page/location" is the JSON field evidence_locations; "supporting passage" and "note" are added columns.',
    ],
    'extraction': [
        'Nested JSON objects (analyte_id, extractor_a_b, ai_assistance) are flattened to "parent.child" columns.',
        'JSON arrays (source_locator, immune_link, scope_streams, sample_set_ids, relationship_tags, registry_ids, ...) are '
        'stored in one cell with values separated by "; "; their dropdowns use a warning (not stop) so several values can be entered.',
        'Vocabulary-to-column mapping is by identical name, plus three aliases: report_cohort_links.scope_streams <- scope_stream, '
        'measurements.record_type <- measurement_record_type, measurements.result_status <- candidate_result_status.',
        'reports.primary_fulltext_exclusion_reason <- fulltext_exclusion_codes.json hierarchy and '
        'reports.consensus_screening_decision <- screening_dispositions (cross-file; warning style).',
        'NR/NA/UNCLEAR are appended to a dropdown when the field is a legacy field whose missing_value_rule allows them, or '
        'the vocabulary already contains them; they are not appended to structural lists (precision_domain, record_type, '
        'scope_streams). Primary/foreign-key columns reject NR/NA/UNCLEAR (missing_value_semantics.identifiers).',
        'Fields with legacy data_type enum (or new fields with a vocabulary) get a strict dropdown; enum_or_text / '
        'string_or_object fields with a vocabulary get a warning dropdown that still allows free text.',
        'Count columns (legacy integer_or_missing_code / separate_counts, plus new fields named n_*, *_count, *_n) accept a '
        'non-negative integer or NR/NA/UNCLEAR. Every other column is Text format (@) to stop Excel auto-conversion of IDs, '
        'DOIs, PMIDs, gene names (SEPT7, MARCH1), ratios (1-2) and leading zeros.',
        'Header comments: definition from legacy_61_fields (exact table.column, or a same-named legacy field in another table, '
        'marked as such), new_operational_fields_not_in_legacy_61 (slash shorthand expanded), primary/foreign-key specs, or the '
        'controlled-vocabulary name; columns with none of these are marked "no definition in data_dictionary.json".',
    ],
    'appraisal': [
        'Item wording is NOT reproduced: critical_appraisal_manual.md requires the archived official form as the source of item '
        'wording, so each tool sheet is a long item table (item_id_as_printed is typed by the appraiser).',
        'Tool sheets are named by tool_id; "_CURRENT" is dropped only when the id exceeds Excel\'s 31-character limit.',
        'Record-level scalar fields -> appraisal_records; list fields (items, official_tool_outputs, initial_disagreements, '
        'consensus_items) -> one sheet each (items split per tool); precision_profile -> one wide row per appraisal record.',
        'Dropdown mapping: identical names to status_codes keys; response_normalized <- tool_response_normalized; '
        'analysis_phase_if_prediction_model <- prediction_model_phase; validation_type_if_prediction_model <- '
        'prediction_validation_type; precision_profile *.status <- precision_profile_status; *.linkage_type <- '
        'functional_linkage_type; decision_use_evidence.use_tested <- decision_use_status; initial_disagreements '
        'appraiser_1/2_response <- tool_response_normalized (warning: verbatim answers allowed); tool_id <- tool_registry.',
        'precision_profile_note (a fixed explanatory string in the template) is shown in README, not as a column.',
    ],
}

# Vocabulary aliases for extraction columns whose name differs from the vocabulary key.
VOCAB_ALIASES = {
    ('report_cohort_links', 'scope_streams'): 'scope_stream',
    ('measurements', 'record_type'): 'measurement_record_type',
    ('measurements', 'result_status'): 'candidate_result_status',
}
# Appraisal field -> status_codes key, where the names are not identical.
APPRAISAL_ALIASES = {
    'response_normalized': 'tool_response_normalized',
    'analysis_phase_if_prediction_model': 'prediction_model_phase',
    'validation_type_if_prediction_model': 'prediction_validation_type',
    'status': 'precision_profile_status',
    'linkage_type': 'functional_linkage_type',
    'use_tested': 'decision_use_status',
    'appraiser_1_response': 'tool_response_normalized',
    'appraiser_2_response': 'tool_response_normalized',
}
APPRAISAL_WARNING_FIELDS = {'appraiser_1_response', 'appraiser_2_response'}

# ---------------------------------------------------------------------------------------------
# Styling / Excel helpers
# ---------------------------------------------------------------------------------------------
HEADER_FONT = Font(bold=True, color='FFFFFF')
FILL = {
    'header': PatternFill('solid', fgColor='1F4E78'),
    'formula': PatternFill('solid', fgColor='7F7F7F'),
    'input': PatternFill('solid', fgColor='2E75B6'),
    'section': PatternFill('solid', fgColor='DDEBF7'),
    'result': PatternFill('solid', fgColor='FFF2CC'),
    'inputcell': PatternFill('solid', fgColor='E2EFDA'),
}
LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)
WRAP = Alignment(wrap_text=True, vertical='top')


def L(col: int) -> str:
    return get_column_letter(col)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_sources() -> dict:
    src = {}
    for key, rel in SOURCES.items():
        p = ROOT / rel
        raw = p.read_bytes()
        src[key] = {'path': rel, 'sha256': hashlib.sha256(raw).hexdigest(),
                    'data': json.loads(raw) if rel.endswith('.json') else raw.decode('utf-8')}
    return src


def version_of(key: str, data) -> str:
    """Short human-readable version string for the README."""
    if not isinstance(data, dict):
        m = re.search(r'\*\*Protocol:\*\*\s*([^\n·]+)', data) or re.search(r'v(\d+\.\d+)', data)
        return m.group(1).strip() if m else 'markdown'
    picks = []
    for path in ['template_version', 'protocol_version', 'protocol_date', 'strategy_version', 'draft_updated',
                 'metadata.version', 'metadata.schema_version', 'metadata.date', 'metadata.created_date',
                 'metadata.schema_verified_date', 'accessed_on', 'template_status', 'taxonomy_status']:
        cur = data
        for part in path.split('.'):
            cur = cur.get(part) if isinstance(cur, dict) else None
        if cur is not None:
            picks.append(f'{path}={cur}')
    return '; '.join(picks)


def truncate(text: str, n: int) -> str:
    text = text or ''
    return text if len(text) <= n else text[: n - 1] + '…'


def protect(ws, allow_row_edit: bool = False):
    """Protect a sheet without password; allow filtering, sorting and column formatting."""
    p = ws.protection
    p.sheet = True
    p.autoFilter = False
    p.sort = False
    p.formatColumns = False
    p.formatRows = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    if allow_row_edit:
        p.insertRows = False
        p.deleteRows = False


def write_header(ws, headers, comments=None, kinds=None, width=18, freeze='A2'):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        kind = (kinds or {}).get(h, 'header')
        c.font = HEADER_FONT
        c.fill = FILL.get(kind, FILL['header'])
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.protection = LOCKED
        if comments and comments.get(h):
            cm = Comment(comments[h], 'build_workbooks.py')
            cm.width, cm.height = 420, 260
            c.comment = cm
        ws.column_dimensions[L(i)].width = width
    ws.row_dimensions[1].height = 32
    if freeze:
        ws.freeze_panes = freeze


def input_column(ws, col: int, text: bool = True):
    """Unlock a whole column (column default style) and set Text format to block auto-conversion."""
    cd = ws.column_dimensions[L(col)]
    cd.protection = UNLOCKED
    if text:
        cd.number_format = '@'


def add_dv(ws, col: int, formula1: str, kind: str = 'list', style: str = 'stop', prompt: str = '',
           title: str = '', first_row: int = 2, last_row: int = MAX_ROW, registry=None, label=''):
    dv = DataValidation(type=kind, formula1=formula1, allow_blank=True, showErrorMessage=True,
                        errorStyle=style, showInputMessage=bool(prompt))
    if kind == 'list':
        dv.showDropDown = False   # openpyxl: False means the in-cell dropdown arrow IS shown
    dv.promptTitle = truncate(title, 32)
    dv.prompt = truncate(prompt, 250)
    dv.errorTitle = truncate('Value not in code list' if kind == 'list' else 'Invalid value', 32)
    if kind == 'list' and style == 'stop':
        dv.error = 'Choose a value from the dropdown (codes sheet). Codes come from the JSON templates.'
    elif kind == 'list':
        dv.error = truncate('Not a listed code. For multi-valued fields separate codes with "; ". Press Yes to keep.', 220)
    else:
        dv.error = truncate(prompt or 'Invalid value', 220)
    rng = f'{L(col)}{first_row}:{L(col)}{last_row}'
    dv.add(rng)
    ws.add_data_validation(dv)
    if registry is not None:
        registry.setdefault(ws.title, []).append({'col': L(col), 'formula1': formula1, 'label': label,
                                                  'style': style, 'type': kind})
    return dv


class Codes:
    """A 'codes' sheet where every enumeration is a column backed by a workbook-level defined name."""

    def __init__(self, wb, ws):
        self.wb, self.ws, self.col, self.names = wb, ws, 1, {}

    def add(self, name: str, values, header: str | None = None, extra: list | None = None, note: str = ''):
        values = [str(v) for v in values]
        if not values:
            raise SystemExit(f'Empty enumeration for {name}: check the JSON templates.')
        c = self.col
        cells = [(header or name, values)] + list(extra or [])
        for j, (h, vals) in enumerate(cells):
            hc = self.ws.cell(row=1, column=c + j, value=h)
            hc.font, hc.fill = HEADER_FONT, FILL['header']
            hc.alignment = Alignment(wrap_text=True, vertical='center')
            for i, v in enumerate(vals, start=2):
                cell = self.ws.cell(row=i, column=c + j, value=v)
                cell.alignment = WRAP
                cell.number_format = '@'
            self.ws.column_dimensions[L(c + j)].width = 28 if j == 0 else 48
        if note:
            hc = self.ws.cell(row=1, column=c)
            cm = Comment(note, 'build_workbooks.py')
            cm.width, cm.height = 380, 160
            hc.comment = cm
        last = len(values) + 1
        ref = f"'{self.ws.title}'!${L(c)}$2:${L(c)}${last}"
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        if extra:
            self.wb.defined_names[name + '_table'] = DefinedName(
                name + '_table', attr_text=f"'{self.ws.title}'!${L(c)}$2:${L(c + len(extra))}${last}")
        self.names[name] = values
        self.col = c + len(cells) + 1
        return '=' + name


def define_cell(wb, name: str, ws, ref: str):
    col, row = re.match(r'([A-Z]+)(\d+)', ref).groups()
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!${col}${row}")


def write_readme(ws, title: str, sections):
    """sections: list of (heading, rows); rows are strings or (key, value) tuples."""
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 120
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    r = 3
    for heading, rows in sections:
        hc = ws.cell(row=r, column=1, value=heading)
        hc.font = Font(bold=True)
        hc.fill = FILL['section']
        ws.cell(row=r, column=2).fill = FILL['section']
        r += 1
        for row in rows:
            if isinstance(row, tuple):
                ws.cell(row=r, column=1, value=str(row[0])).alignment = WRAP
                ws.cell(row=r, column=2, value=str(row[1])).alignment = WRAP
            else:
                ws.cell(row=r, column=2, value=str(row)).alignment = WRAP
            r += 1
        r += 1
    protect(ws)


def source_rows(src, keys):
    return [(src[k]['path'], f"{version_of(k, src[k]['data'])}; sha256={src[k]['sha256'][:16]}…") for k in keys]


# ---------------------------------------------------------------------------------------------
# Screening enumerations (all from JSON / manuals)
# ---------------------------------------------------------------------------------------------
def screening_enums(src) -> dict:
    log = src['screening_log']['data']
    ftj = src['ft_codes']['data']
    manual = src['screening_manual']['data']
    plan = src['calibration_plan']['data']
    dic = src['dictionary']['data']
    search = src['search_log']['data']

    disp = list(log['field_schema']['screening_dispositions'])
    sec = re.search(r'^### B\. Title/abstract.*?(?=^### )', manual, re.S | re.M)
    if not sec:
        raise SystemExit('screening_manual.md: section "### B. Title/abstract" not found')
    bold = set(re.findall(r'\*\*([A-Z][A-Z_]+)\*\*', sec.group(0)))
    ta = [d for d in disp if d in bold]
    admin = [s['status'] for s in ftj['not_scientific_exclusion_statuses']]
    ta_only = [d for d in ta if d not in admin]
    ft = [d for d in disp if d not in ta_only]
    hierarchy = list(ftj['primary_reason_rule']['hierarchy'])
    code_info = {c['code']: c for c in ftj['codes']}
    ft_excl = [d for d in ft if d not in admin and 'EXCLUDE' in d]
    ft_incl = [d for d in ft if d not in admin and d not in ft_excl]
    ta_excl = [d for d in ta_only if 'EXCLUDE' in d]

    m = re.search(r'\*\*([^*]+)\*\* \(([^)]*)\) versus \*\*([A-Z_]+)\*\*', plan)
    if not m:
        raise SystemExit('calibration_plan.md: binary disposition sentence not found')
    prog_label, prog_txt, excl_label = m.group(1).strip(), m.group(2), m.group(3)
    progress = [d for d in re.findall(r'[A-Z][A-Z_]+', prog_txt) if d in disp]
    thr = re.search(r'raw agreement is \*\*at least (\d+(?:\.\d+)?)%\*\*', plan)
    if not thr:
        raise SystemExit('calibration_plan.md: raw agreement threshold not found')
    problems = []
    if set(progress) | {excl_label} != set(ta):
        problems.append(f'calibration binary groups {progress}+[{excl_label}] differ from TA dispositions {ta}')
    if [c['code'] for c in ftj['codes']] != hierarchy:
        problems.append('fulltext_exclusion_codes.json: codes[] order differs from primary_reason_rule.hierarchy')
    for s in admin:
        if s not in disp:
            problems.append(f'administrative status {s} is not in screening_dispositions')
    missing_codes = list(dic['missing_value_semantics']['codes'])
    dups = [s for s in admin if s.startswith('DUPLICATE')]
    cal = log['calibration']
    rule = cal.get('repeat_seed_rule', '')
    seed_offset_ok = bool(re.search(r'\+\s*round_number\s*-\s*1', rule))
    return {
        'dispositions': disp, 'ta': ta, 'ft': ft, 'admin': admin, 'admin_info': ftj['not_scientific_exclusion_statuses'],
        'codes': hierarchy, 'code_info': code_info, 'ft_excl': ft_excl, 'ft_incl': ft_incl, 'ta_excl': ta_excl,
        'binary_labels': [prog_label, excl_label], 'binary_map': {d: (excl_label if d == excl_label else prog_label) for d in ta},
        'threshold': float(thr.group(1)) / 100.0, 'missing_codes': missing_codes,
        'booleans': BOOLEAN_VALUES + missing_codes, 'duplicates': dups,
        'not_exclusion_reasons': ftj['explicitly_not_exclusion_reasons'],
        'searches': search['searches'], 'calibration': cal, 'seed_offset_ok': seed_offset_ok,
        'ta_fields': list(log['field_schema']['title_abstract']['reviewer_1'].keys()),
        'ft_fields': list(log['field_schema']['full_text']['reviewer_1'].keys()),
        'field_schema': log['field_schema'], 'problems': problems,
        'ft_groups': {'INCLUDE': ft_incl, 'SCIENTIFIC_EXCLUSION': ft_excl, 'ADMINISTRATIVE_STATUS': [d for d in ft if d in admin]},
    }


# ---------------------------------------------------------------------------------------------
# Screening workbook
# ---------------------------------------------------------------------------------------------
MASTER_COLUMNS = ['record_id', 'source_database', 'search_id', 'route', 'title', 'authors', 'year', 'journal',
                  'doi', 'pmid', 'abstract', 'dedup_group_id', 'dedup_status', 'retained_record_id']


def agreement_block(wb, ws, top: int, left: int, cats, rng_a: str, rng_b: str, prefix: str, title: str):
    """k x k agreement table with Cohen's kappa computed from worksheet formulas.
    rng_a / rng_b: absolute ranges holding category labels for reviewer A / B (blank = not decided)."""
    k = len(cats)
    ws.cell(row=top, column=left, value=title).font = Font(bold=True)
    ws.cell(row=top, column=left).fill = FILL['section']
    hr = top + 1
    ws.cell(row=hr, column=left, value='A (rows) \\ B (cols)').font = Font(bold=True)
    for j, c in enumerate(cats):
        ws.cell(row=hr, column=left + 1 + j, value=c).font = Font(bold=True)
        ws.cell(row=hr, column=left + 1 + j).alignment = Alignment(wrap_text=True)
    ws.cell(row=hr, column=left + 1 + k, value='row total (A)').font = Font(bold=True)
    for i, c in enumerate(cats):
        r = hr + 1 + i
        ws.cell(row=r, column=left, value=c).font = Font(bold=True)
        for j in range(k):
            ws.cell(row=r, column=left + 1 + j,
                    value=f'=COUNTIFS({rng_a},${L(left)}{r},{rng_b},{L(left + 1 + j)}${hr})')
        ws.cell(row=r, column=left + 1 + k, value=f'=SUM({L(left + 1)}{r}:{L(left + k)}{r})')
    tr = hr + 1 + k
    ws.cell(row=tr, column=left, value='column total (B)').font = Font(bold=True)
    for j in range(k + 1):
        cl = L(left + 1 + j)
        ws.cell(row=tr, column=left + 1 + j, value=f'=SUM({cl}{hr + 1}:{cl}{hr + k})')
    diag = ','.join(f'{L(left + 1 + i)}{hr + 1 + i}' for i in range(k))
    prods = '+'.join(f'{L(left + 1 + k)}{hr + 1 + i}*{L(left + 1 + i)}{tr}' for i in range(k))
    n_ref = f'{L(left + 1 + k)}{tr}'
    sr = tr + 1
    lab, val = left, left + 1
    rows = [
        ('n (jointly coded with listed codes)', f'={n_ref}', 'n'),
        ('observed agreements (diagonal)', f'=SUM({diag})', 'agree'),
        ('raw agreement p_o', f'=IF({L(val)}{sr}=0,"",{L(val)}{sr + 1}/{L(val)}{sr})', 'po'),
        ('chance agreement p_e', f'=IF({L(val)}{sr}=0,"",({prods})/{L(val)}{sr}^2)', 'pe'),
        ("Cohen's kappa (descriptive)",
         f'=IF({L(val)}{sr}=0,"",IF({L(val)}{sr + 3}=1,"undefined (p_e=1)",({L(val)}{sr + 2}-{L(val)}{sr + 3})/(1-{L(val)}{sr + 3})))',
         'kappa'),
    ]
    out = {}
    for i, (label, formula, key) in enumerate(rows):
        ws.cell(row=sr + i, column=lab, value=label)
        c = ws.cell(row=sr + i, column=val, value=formula)
        c.fill = FILL['result']
        if key in ('po', 'pe'):
            c.number_format = '0.0%'
        elif key == 'kappa':
            c.number_format = '0.000'
        ref = f'{L(val)}{sr + i}'
        define_cell(wb, f'{prefix}_{key}', ws, ref)
        out[key] = ref
    return sr + len(rows) + 1, out


def build_screening(src, out_path: Path, rows: int = 5000, conflict_rows: int = 1000, registry=None):
    E = screening_enums(src)
    registry = {} if registry is None else registry
    wb = openpyxl.Workbook()
    names = ['README', 'records_master', 'screen_TA_reviewer_A', 'screen_TA_reviewer_B',
             'screen_FT_reviewer_A', 'screen_FT_reviewer_B', 'merge_TA', 'merge_FT', 'codes', 'log']
    ws_readme = wb.active
    ws_readme.title = names[0]
    S = {'README': ws_readme}
    for n in names[1:]:
        S[n] = wb.create_sheet(n)
    last = rows + 1

    # ---- codes -------------------------------------------------------------------------------
    codes = Codes(wb, S['codes'])
    ci = E['code_info']
    codes.add('lst_TA_decision', E['ta'], 'TA decision (screening_dispositions, TA stage)')
    codes.add('lst_TA_reason', E['codes'], 'TA primary_reason (FT01–FT08 hierarchy)',
              extra=[('label', [ci[c]['label'] for c in E['codes']])])
    codes.add('lst_FT_decision', E['ft'], 'FT decision (screening_dispositions, FT stage)')
    codes.add('lst_FT_code', E['codes'], 'FT primary_exclusion_code (apply in hierarchy order)',
              extra=[('label', [ci[c]['label'] for c in E['codes']]), ('definition', [ci[c]['definition'] for c in E['codes']])])
    codes.add('lst_FT_sci_excl', E['ft_excl'], 'FT dispositions that require an FT code')
    codes.add('lst_boolean', E['booleans'], 'boolean + missing codes')
    codes.add('lst_TA_binary_map', E['ta'], 'TA disposition → calibration binary group',
              extra=[('binary group', [E['binary_map'][d] for d in E['ta']])])
    codes.add('lst_TA_binary', E['binary_labels'], 'calibration binary groups (calibration_plan.md)')
    grp_rows = [(d, g) for g, ds in E['ft_groups'].items() for d in ds]
    ft_order = {d: i for i, d in enumerate(E['ft'])}
    grp_rows.sort(key=lambda x: ft_order[x[0]])
    codes.add('lst_FT_group_map', [d for d, _ in grp_rows], 'FT disposition → 3-group collapse',
              extra=[('group', [g for _, g in grp_rows])])
    codes.add('lst_FT_group', list(E['ft_groups']), 'FT 3-group labels (formula mechanics)')
    codes.add('lst_admin_status', E['admin'], 'administrative statuses (not scientific exclusions)',
              extra=[('meaning', [s['meaning'] for s in E['admin_info']]), ('handling', [s['handling'] for s in E['admin_info']])])
    codes.add('lst_dedup_duplicate', E['duplicates'], 'dedup_status values (blank = retained)')
    codes.add('lst_search_id', [s['id'] for s in E['searches']], 'search_id (search_log_template)',
              extra=[('database', [s['database'] for s in E['searches']]), ('platform', [s['platform'] for s in E['searches']]),
                     ('route', [s['route'] for s in E['searches']])])
    codes.add('lst_database', sorted({s['database'] for s in E['searches']}), 'source_database values')
    codes.add('lst_route', sorted({s['route'] for s in E['searches']}), 'route values')
    codes.add('lst_not_exclusion', E['not_exclusion_reasons'], 'explicitly NOT exclusion reasons (reference)')
    protect(S['codes'])
    S['codes'].sheet_properties.tabColor = '808080'

    # ---- records_master (read-only) ----------------------------------------------------------
    fs = E['field_schema']
    dbs = sorted({s['database'] for s in E['searches']})
    m_comments = {
        'record_id': fs['record_id'] + ' Text format; never reuse an ID.',
        'source_database': fs['source_database_or_route'] + f' Values (search_log_template): {"; ".join(dbs)}. '
                           f'Several sources for one record: separate with "{MULTI_SEP}".',
        'search_id': 'search_log_template.searches[].id of the export(s) containing this record; several: "; ".',
        'route': 'search_log_template route (E AND I / E AND O / supplemental). Several: "; ".',
        'doi': 'Text format: keep exactly as exported.', 'pmid': 'Text format (leading zeros preserved).',
        'dedup_group_id': 'Same value for all rows that are copies/versions of one report.',
        'dedup_status': f'Blank = retained/unique record. {" / ".join(E["duplicates"])} (fulltext_exclusion_codes.json) '
                        'for removed duplicates; such rows are hidden from the reviewer sheets.',
        'retained_record_id': 'screening_log_template administrative.retained_record_id: record_id kept for a duplicate.',
    }
    ws = S['records_master']
    write_header(ws, MASTER_COLUMNS, m_comments)
    for i in range(1, len(MASTER_COLUMNS) + 1):
        ws.column_dimensions[L(i)].number_format = '@'
    ws.column_dimensions[L(MASTER_COLUMNS.index('abstract') + 1)].width = 60
    ws.column_dimensions[L(MASTER_COLUMNS.index('title') + 1)].width = 50
    col = {h: i + 1 for i, h in enumerate(MASTER_COLUMNS)}
    add_dv(ws, col['dedup_status'], '=lst_dedup_duplicate', registry=registry, label='dedup_status',
           prompt='Blank = retained. Duplicate statuses from fulltext_exclusion_codes.json.', title='dedup_status')
    add_dv(ws, col['search_id'], '=lst_search_id', style='warning', registry=registry, label='search_id',
           prompt='search_log_template id; several separated by "; "', title='search_id')
    add_dv(ws, col['source_database'], '=lst_database', style='warning', registry=registry, label='source_database',
           prompt='Database (search_log_template); several separated by "; "', title='source_database')
    add_dv(ws, col['route'], '=lst_route', style='warning', registry=registry, label='route',
           prompt='Route (search_log_template)', title='route')
    ws.auto_filter.ref = f'A1:{L(len(MASTER_COLUMNS))}1'
    protect(ws)
    ws.sheet_properties.tabColor = '1F4E78'
    mc = {h: L(i) for h, i in col.items()}

    # ---- reviewer sheets -----------------------------------------------------------------------
    def rid_formula(r):
        return (f'=IF(OR(records_master!${mc["record_id"]}{r}="",ISNUMBER(MATCH(records_master!${mc["dedup_status"]}{r},'
                f'lst_dedup_duplicate,0))),"",records_master!${mc["record_id"]}{r})')

    ta_cols = ['record_id', 'title'] + E['ta_fields'] + ['note']
    ft_cols = ['record_id', 'title'] + E['ft_fields'] + ['supporting_passage', 'note', 'check_code_rule']
    ta_comments = {
        'record_id': 'Formula from records_master (locked). Do not sort or insert rows: row n = records_master row n.',
        'title': 'Formula from records_master (locked); read the abstract in records_master.',
        'decision': 'Independent TA disposition (screening_manual.md §3B). When uncertain: ADVANCE or AWAITING_CLASSIFICATION; '
                    'never exclude by assumption.',
        'primary_reason': 'Required only for EXCLUDE_TA: the first failed eligibility dimension (FT hierarchy). '
                          'Full-text-only criteria are not TA exclusion grounds.',
        'rationale_or_quote': 'Concise abstract/title quote or basis.',
        'decided_at': 'YYYY-MM-DD (text).', 'note': 'Free text.',
    }
    ft_comments = {
        'record_id': ta_comments['record_id'], 'title': ta_comments['title'],
        'decision': 'Independent FT disposition. Administrative statuses (duplicate, not retrieved, awaiting translation/'
                    'classification, frontier preprint) are NOT scientific exclusions and take no FT code.',
        'primary_exclusion_code': 'Only with ' + '/'.join(E['ft_excl']) + ': exactly one code, the FIRST established reason '
                                  'in hierarchy order (fulltext_exclusion_codes.json). Leave blank otherwise.',
        'secondary_failed_dimensions': 'Other failed dimensions (FT codes), separated by "; ".',
        'A_eligible': 'Core A criteria met (TRUE/FALSE) or NR/NA/UNCLEAR.', 'B_eligible': 'Support B criteria met.',
        'evidence_locations': 'Page / section / table / figure / supplement location of the supporting evidence.',
        'rationale': 'Rule applied and reasoning.', 'decided_at': 'YYYY-MM-DD (text).',
        'supporting_passage': 'Verbatim supporting passage from the full text.', 'note': 'Free text.',
        'check_code_rule': 'Formula: flags EXCLUDE without code, or a code without EXCLUDE (screening_log record_validation_rules).',
    }
    for stage, cols, cmts in (('TA', ta_cols, ta_comments), ('FT', ft_cols, ft_comments)):
        for rv in ('A', 'B'):
            ws = S[f'screen_{stage}_reviewer_{rv}']
            kinds = {h: ('formula' if h in ('record_id', 'title', 'check_code_rule') else 'input') for h in cols}
            write_header(ws, cols, cmts, kinds, freeze='C2')
            ws.column_dimensions['B'].width = 50
            c = {h: i + 1 for i, h in enumerate(cols)}
            for h in cols:
                if kinds[h] == 'input':
                    input_column(ws, c[h])
            for r in range(2, last + 1):
                ws.cell(row=r, column=1, value=rid_formula(r))
                ws.cell(row=r, column=2, value=f'=IF($A{r}="","",records_master!${mc["title"]}{r})')
            if stage == 'TA':
                add_dv(ws, c['decision'], '=lst_TA_decision', registry=registry, label='TA decision',
                       prompt='TA disposition (screening_dispositions). Uncertain -> ADVANCE / AWAITING_CLASSIFICATION.',
                       title='TA decision')
                add_dv(ws, c['primary_reason'], '=lst_TA_reason', registry=registry, label='TA primary_reason',
                       prompt='Only for EXCLUDE_TA: first failed dimension (FT01-FT08 hierarchy).', title='primary reason')
            else:
                add_dv(ws, c['decision'], '=lst_FT_decision', registry=registry, label='FT decision',
                       prompt='FT disposition. Admin statuses are not scientific exclusions.', title='FT decision')
                add_dv(ws, c['primary_exclusion_code'], '=lst_FT_code', registry=registry, label='FT code',
                       prompt='Only with EXCLUDE_FT: first established reason in FT01->FT08 order.', title='FT code')
                add_dv(ws, c['secondary_failed_dimensions'], '=lst_FT_code', style='warning', registry=registry,
                       label='FT secondary', prompt='FT codes separated by "; "', title='secondary dimensions')
                for h in ('A_eligible', 'B_eligible'):
                    add_dv(ws, c[h], '=lst_boolean', registry=registry, label=h, prompt='TRUE / FALSE / NR / NA / UNCLEAR',
                           title=h)
                d, k = L(c['decision']), L(c['primary_exclusion_code'])
                for r in range(2, last + 1):
                    ws.cell(row=r, column=c['check_code_rule'], value=(
                        f'=IF($A{r}="","",IF(AND(ISNUMBER(MATCH({d}{r},lst_FT_sci_excl,0)),{k}{r}=""),"MISSING FT CODE",'
                        f'IF(AND(NOT(ISNUMBER(MATCH({d}{r},lst_FT_sci_excl,0))),{k}{r}<>""),"CODE NOT ALLOWED FOR THIS DECISION","")))'))
            ws.auto_filter.ref = f'A1:{L(len(cols))}1'
            protect(ws)
            ws.sheet_properties.tabColor = '2E75B6' if rv == 'A' else 'C55A11'

    # ---- merge sheets --------------------------------------------------------------------------
    def merge_sheet(stage: str):
        ws = S[f'merge_{stage}']
        ra, rb = f'screen_{stage}_reviewer_A', f'screen_{stage}_reviewer_B'
        rcols = ta_cols if stage == 'TA' else ft_cols
        rc = {h: L(i + 1) for i, h in enumerate(rcols)}
        reason = 'primary_reason' if stage == 'TA' else 'primary_exclusion_code'
        grp = 'binary' if stage == 'TA' else 'group'
        gmap = 'lst_TA_binary_map_table' if stage == 'TA' else 'lst_FT_group_map_table'
        heads = ['record_id', 'decision_A', 'decision_B', f'{reason}_A', f'{reason}_B', 'both_decided', 'exact_agree',
                 f'{grp}_A', f'{grp}_B', f'{grp}_agree']
        if stage == 'FT':
            heads.append('code_agree')
        heads += ['conflict', 'conflict_seq']
        write_header(ws, heads, {
            'record_id': 'Formula. Rows align with records_master; duplicates are blank.',
            'exact_agree': '1 = same disposition, 0 = different, blank = not jointly decided.',
            f'{grp}_A': 'Collapsed group used for the ' + ('calibration (binary)' if stage == 'TA' else '3-group') + ' table.',
            'conflict': '1 when the dispositions differ' + (' or both chose an exclusion with different FT codes.' if stage == 'FT' else '.'),
            'code_agree': 'Both scientific exclusion: 1 same FT code, 0 different.',
        }, {h: 'formula' for h in heads}, width=14, freeze='B2')
        h = {x: L(i + 1) for i, x in enumerate(heads)}
        for r in range(2, last + 1):
            f = {
                'record_id': f'=IF({ra}!$A{r}="","",{ra}!$A{r})',
                'decision_A': f'=IF(OR($A{r}="",{ra}!${rc["decision"]}{r}=""),"",{ra}!${rc["decision"]}{r})',
                'decision_B': f'=IF(OR($A{r}="",{rb}!${rc["decision"]}{r}=""),"",{rb}!${rc["decision"]}{r})',
                f'{reason}_A': f'=IF(OR($A{r}="",{ra}!${rc[reason]}{r}=""),"",{ra}!${rc[reason]}{r})',
                f'{reason}_B': f'=IF(OR($A{r}="",{rb}!${rc[reason]}{r}=""),"",{rb}!${rc[reason]}{r})',
                'both_decided': f'=IF(AND({h["decision_A"]}{r}<>"",{h["decision_B"]}{r}<>""),1,"")',
                'exact_agree': f'=IF({h["both_decided"]}{r}=1,IF({h["decision_A"]}{r}={h["decision_B"]}{r},1,0),"")',
                f'{grp}_A': f'=IF({h["decision_A"]}{r}="","",IFERROR(VLOOKUP({h["decision_A"]}{r},{gmap},2,FALSE),"INVALID"))',
                f'{grp}_B': f'=IF({h["decision_B"]}{r}="","",IFERROR(VLOOKUP({h["decision_B"]}{r},{gmap},2,FALSE),"INVALID"))',
                f'{grp}_agree': f'=IF({h["both_decided"]}{r}=1,IF({h[grp + "_A"]}{r}={h[grp + "_B"]}{r},1,0),"")',
            }
            if stage == 'FT':
                f['code_agree'] = (f'=IF(AND(ISNUMBER(MATCH({h["decision_A"]}{r},lst_FT_sci_excl,0)),'
                                   f'ISNUMBER(MATCH({h["decision_B"]}{r},lst_FT_sci_excl,0))),'
                                   f'IF({h[reason + "_A"]}{r}={h[reason + "_B"]}{r},1,0),"")')
                f['conflict'] = f'=IF(OR({h["exact_agree"]}{r}=0,{h["code_agree"]}{r}=0),1,"")'
            else:
                f['conflict'] = f'=IF({h["exact_agree"]}{r}=0,1,"")'
            f['conflict_seq'] = f'=IF({h["conflict"]}{r}=1,COUNTIF(${h["conflict"]}$2:{h["conflict"]}{r},1),"")'
            for x in heads:
                ws.cell(row=r, column=heads.index(x) + 1, value=f[x])

        def rng(x):
            return f'${h[x]}$2:${h[x]}${last}'

        # conflict list
        cl0 = len(heads) + 2
        clh = ['conflict_no', 'row_in_table', 'record_id', 'decision_A', 'decision_B', f'{reason}_A', f'{reason}_B']
        for j, x in enumerate(clh):
            c = ws.cell(row=1, column=cl0 + j, value=x)
            c.font, c.fill = HEADER_FONT, FILL['formula']
            ws.column_dimensions[L(cl0 + j)].width = 14
        for i in range(conflict_rows):
            r = i + 2
            ws.cell(row=r, column=cl0, value=i + 1)
            ws.cell(row=r, column=cl0 + 1, value=f'=IFERROR(MATCH({L(cl0)}{r},{rng("conflict_seq")},0),"")')
            for j, x in enumerate(clh[2:]):
                ws.cell(row=r, column=cl0 + 2 + j,
                        value=f'=IF({L(cl0 + 1)}{r}="","",INDEX({rng(x)},{L(cl0 + 1)}{r}))')

        # summary
        s0 = cl0 + len(clh) + 1
        lab, val = s0, s0 + 1
        ws.column_dimensions[L(lab)].width = 46
        r = 1
        ws.cell(row=r, column=lab, value=f'SUMMARY merge_{stage} (worksheet formulas; do not edit)').font = Font(bold=True, size=12)
        r += 1
        if stage == 'TA':
            ws.cell(row=r, column=lab, value=(
                f'Pass rule (calibration_plan.md): raw agreement on binary {E["binary_labels"][0]} vs {E["binary_labels"][1]} '
                f'≥ {E["threshold"]:.0%} AND every conceptual disagreement resolved. Kappa is descriptive only.'))
        else:
            ws.cell(row=r, column=lab, value=('Full-text agreement is descriptive; the TA calibration threshold is not '
                                              'reinterpreted as a full-text threshold (calibration_plan.md, Full-text stage).'))
        ws.cell(row=r, column=lab).alignment = WRAP
        r += 2
        stats = [
            ('records in records_master (non-blank record_id)', f'=SUMPRODUCT(--(records_master!$A$2:$A${last}<>""))', 'master_n'),
            ('records available to screen (duplicates hidden)', f'=SUMPRODUCT(--({rng("record_id")}<>""))', 'pool_n'),
            ('decided by reviewer A', f'=SUMPRODUCT(--({rng("decision_A")}<>""))', 'n_A'),
            ('decided by reviewer B', f'=SUMPRODUCT(--({rng("decision_B")}<>""))', 'n_B'),
            ('jointly decided', f'=SUM({rng("both_decided")})', 'n_joint'),
            ('conflicts', f'=SUM({rng("conflict")})', 'n_conflict'),
            ('values outside the code list (INVALID)', f'=COUNTIF({rng(grp + "_A")},"INVALID")+COUNTIF({rng(grp + "_B")},"INVALID")', 'n_invalid'),
            ('capacity warning', f'=IF(COUNTA(records_master!$A:$A)-1>{rows},"MORE RECORDS THAN FORMULA ROWS ({rows}): regenerate with --rows",'
                                 f'IF({L(val)}{r + 5}>{conflict_rows},"CONFLICTS EXCEED LIST ({conflict_rows}): filter column conflict",""))', 'warning'),
        ]
        for label, formula, key in stats:
            ws.cell(row=r, column=lab, value=label)
            c = ws.cell(row=r, column=val, value=formula)
            c.fill = FILL['result']
            define_cell(wb, f'{stage}_{key}', ws, f'{L(val)}{r}')
            r += 1
        r += 1
        # per-reviewer counts
        cats = E['ta'] if stage == 'TA' else E['ft']
        ws.cell(row=r, column=lab, value='Counts per reviewer (all decided records)').font = Font(bold=True)
        ws.cell(row=r, column=lab).fill = FILL['section']
        ws.cell(row=r, column=val, value='reviewer A').font = Font(bold=True)
        ws.cell(row=r, column=val + 1, value='reviewer B').font = Font(bold=True)
        r += 1
        for cat in cats:
            ws.cell(row=r, column=lab, value=cat)
            ws.cell(row=r, column=val, value=f'=COUNTIF({rng("decision_A")},{L(lab)}{r})')
            ws.cell(row=r, column=val + 1, value=f'=COUNTIF({rng("decision_B")},{L(lab)}{r})')
            r += 1
        r += 1
        if stage == 'FT':
            ws.cell(row=r, column=lab, value='Scientific exclusions by primary FT code').font = Font(bold=True)
            ws.cell(row=r, column=lab).fill = FILL['section']
            ws.cell(row=r, column=val, value='reviewer A').font = Font(bold=True)
            ws.cell(row=r, column=val + 1, value='reviewer B').font = Font(bold=True)
            r += 1
            for code in E['codes']:
                ws.cell(row=r, column=lab, value=code)
                ws.cell(row=r, column=val, value=f'=COUNTIFS({rng("group_A")},"SCIENTIFIC_EXCLUSION",{rng(reason + "_A")},{L(lab)}{r})')
                ws.cell(row=r, column=val + 1, value=f'=COUNTIFS({rng("group_B")},"SCIENTIFIC_EXCLUSION",{rng(reason + "_B")},{L(lab)}{r})')
                define_cell(wb, f'FT_excl_A_{code[:4]}', ws, f'{L(val)}{r}')
                define_cell(wb, f'FT_excl_B_{code[:4]}', ws, f'{L(val + 1)}{r}')
                r += 1
            r += 1
        r, _ = agreement_block(wb, ws, r, lab, cats, rng('decision_A'), rng('decision_B'), f'{stage}_exact',
                               f'Agreement on exact {stage} disposition ({len(cats)}×{len(cats)})')
        gcats = E['binary_labels'] if stage == 'TA' else list(E['ft_groups'])
        r, _ = agreement_block(wb, ws, r, lab, gcats, rng(grp + '_A'), rng(grp + '_B'), f'{stage}_{grp}',
                               f'Agreement on {"calibration binary" if stage == "TA" else "3-group"} disposition '
                               f'({len(gcats)}×{len(gcats)})')
        unlocked = []
        if stage == 'TA':
            cal = E['calibration']
            n_cal = int(cal['planned_sample_size'])
            r += 1
            ws.cell(row=r, column=lab, value=(f'calibration_{n_cal} — initial pilot (calibration_plan.md): {n_cal} records '
                                              f'drawn with random.Random(seed).sample(sorted(pool), {n_cal}); paste sampled '
                                              f'record_ids below (green cells).')).font = Font(bold=True)
            ws.cell(row=r, column=lab).fill = FILL['section']
            r += 1
            params = [
                ('planned_sample_size (screening_log_template)', n_cal, 'cal_planned', False),
                ('base random_seed (screening_log_template)', cal['random_seed'], 'cal_seed', False),
                ('round_number (1 = initial; input)', None, 'cal_round', True),
                ('seed for this round (= seed + round_number - 1)',
                 f'=IF({L(val)}{r + 2}="","",{L(val)}{r + 1}+{L(val)}{r + 2}-1)' if E['seed_offset_ok'] else 'see repeat_seed_rule', 'cal_round_seed', False),
                ('raw-agreement threshold (calibration_plan.md)', E['threshold'], 'cal_threshold', False),
                ('all conceptual disagreements resolved? (TRUE/FALSE; input)', None, 'cal_resolved', True),
            ]
            pr = {}
            for label, value, key, is_input in params:
                ws.cell(row=r, column=lab, value=label)
                c = ws.cell(row=r, column=val, value=value)
                if key == 'cal_threshold':
                    c.number_format = '0%'
                if is_input:
                    c.protection = UNLOCKED
                    c.fill = FILL['inputcell']
                    unlocked.append(f'{L(val)}{r}')
                define_cell(wb, f'TA_{key}', ws, f'{L(val)}{r}')
                pr[key] = f'{L(val)}{r}'
                r += 1
            add_dv(ws, val, '=lst_boolean', first_row=int(pr['cal_resolved'][len(L(val)):]),
                   last_row=int(pr['cal_resolved'][len(L(val)):]), registry=registry, label='cal_resolved',
                   prompt='TRUE only when every conceptual disagreement is resolved and documented.', title='resolved?')
            # results placeholder rows, then the per-ID table
            res_top = r
            r += 7
            th = ['sample_no', 'record_id (input)', 'in_pool', 'binary_A', 'binary_B', 'jointly', 'agree']
            for j, x in enumerate(th):
                c = ws.cell(row=r, column=lab + j, value=x)
                c.font, c.fill = HEADER_FONT, FILL['header']
            t0 = r + 1
            sampled = list(cal.get('sampled_record_ids') or [])
            for i in range(n_cal):
                rr = t0 + i
                idc = f'{L(lab + 1)}{rr}'
                ws.cell(row=rr, column=lab, value=i + 1)
                c = ws.cell(row=rr, column=lab + 1, value=sampled[i] if i < len(sampled) else None)
                c.protection, c.fill, c.number_format = UNLOCKED, FILL['inputcell'], '@'
                unlocked.append(idc)
                ws.cell(row=rr, column=lab + 2, value=f'=IF({idc}="","",IF(ISNUMBER(MATCH({idc},{rng("record_id")},0)),1,"NOT IN POOL"))')
                ws.cell(row=rr, column=lab + 3, value=f'=IF({L(lab + 2)}{rr}<>1,"",INDEX({rng("binary_A")},MATCH({idc},{rng("record_id")},0)))')
                ws.cell(row=rr, column=lab + 4, value=f'=IF({L(lab + 2)}{rr}<>1,"",INDEX({rng("binary_B")},MATCH({idc},{rng("record_id")},0)))')
                ws.cell(row=rr, column=lab + 5, value=f'=IF(AND({L(lab + 3)}{rr}<>"",{L(lab + 4)}{rr}<>""),1,"")')
                ws.cell(row=rr, column=lab + 6, value=f'=IF({L(lab + 5)}{rr}=1,IF({L(lab + 3)}{rr}={L(lab + 4)}{rr},1,0),"")')
            t1 = t0 + n_cal - 1
            ra_ = f'${L(lab + 3)}${t0}:${L(lab + 3)}${t1}'
            rb_ = f'${L(lab + 4)}${t0}:${L(lab + 4)}${t1}'
            _, kap = agreement_block(wb, ws, t1 + 2, lab, E['binary_labels'], ra_, rb_, 'TA_cal',
                                     f'calibration_{n_cal}: binary agreement table on the sampled records')
            rows_res = [
                ('sample IDs entered', f'=SUMPRODUCT(--({L(lab + 1)}{t0}:{L(lab + 1)}{t1}<>""))', 'cal_ids'),
                ('IDs not found in pool', f'=COUNTIF({L(lab + 2)}{t0}:{L(lab + 2)}{t1},"NOT IN POOL")', 'cal_notfound'),
                ('jointly assessed (denominator)', f'=SUM({L(lab + 5)}{t0}:{L(lab + 5)}{t1})', 'cal_den'),
                ('agreements (numerator)', f'=SUM({L(lab + 6)}{t0}:{L(lab + 6)}{t1})', 'cal_num'),
                ('raw agreement', f'=IF({L(val)}{res_top + 2}=0,"",{L(val)}{res_top + 3}/{L(val)}{res_top + 2})', 'cal_raw'),
                ("Cohen's kappa (descriptive; see table below)", f'={kap["kappa"]}', 'cal_kappa'),
                ('calibration status',
                 f'=IF({L(val)}{res_top + 2}<{pr["cal_planned"]},"INCOMPLETE: "&{L(val)}{res_top + 2}&"/"&{pr["cal_planned"]}&" jointly assessed",'
                 f'IF(AND({L(val)}{res_top + 4}>={pr["cal_threshold"]},{pr["cal_resolved"]}="TRUE"),"PASS",'
                 f'"NOT PASSED: clarify manual and run a fresh round (new seed, unseen records)"))', 'cal_status'),
            ]
            for i, (label, formula, key) in enumerate(rows_res):
                ws.cell(row=res_top + i, column=lab, value=label)
                c = ws.cell(row=res_top + i, column=val, value=formula)
                c.fill = FILL['result']
                if key == 'cal_raw':
                    c.number_format = '0.0%'
                define_cell(wb, f'TA_{key}', ws, f'{L(val)}{res_top + i}')
        protect(ws)
        ws.sheet_properties.tabColor = '548235'
        registry.setdefault('_unlocked_cells', {})[ws.title] = unlocked

    merge_sheet('TA')
    merge_sheet('FT')

    # ---- log -----------------------------------------------------------------------------------
    ws = S['log']
    write_header(ws, ['date', 'who', 'action', 'free_text'], {
        'date': 'YYYY-MM-DD (text).', 'who': 'Role ID (reviewer_1 / reviewer_2 / adjudicator / data manager); no invented names.',
        'action': 'e.g. import, dedup, lock TA, merge, rule clarification, regenerate workbook.', 'free_text': 'Details.'})
    ws.column_dimensions['D'].width = 90
    for i in range(1, 5):
        input_column(ws, i)

    # ---- README --------------------------------------------------------------------------------
    gen = _dt.date.today().isoformat()
    cal = E['calibration']
    write_readme(S['README'], '筛选工作簿 Screening workbook (blank template — no records, no decisions)', [
        ('生成信息 Generation', [('generated_on', gen), ('generator', f'scripts/build_workbooks.py v{SCRIPT_VERSION}'),
                                ('formula rows (--rows)', rows), ('conflict list rows', conflict_rows),
                                ('regenerate', 'python3 scripts/build_workbooks.py --reviewer A --reviewer B'),
                                ('status', f'formal_screening_status={src["screening_log"]["data"].get("formal_screening_status")}; '
                                           f'screening_counts_status={src["screening_log"]["data"].get("screening_counts_status")}')]),
        ('源模板版本 Source templates', source_rows(src, SCREENING_SOURCES)),
        ('使用步骤 How to use', [
            '1. 数据管理员：撤销 records_master 工作表保护（无密码：审阅→撤销工作表保护），以“仅粘贴数值”导入去重后的正式检索记录（含重复记录行并标注 dedup_status），然后重新保护。不得加入摸底/校准示例记录。',
            '2. 每位审阅者使用自己的副本，只填写本人 sheet（screen_TA_reviewer_A 或 _B）；锁定本人决定前不得查看对方 sheet。决定列为下拉选择，结构列已锁定。不要排序、插入或删除行：第 n 行与 records_master 第 n 行对应。',
            '3. 双方锁定后，数据管理员把两份决定合入同一工作簿（复制 decision 等输入列的数值），merge_TA 自动给出一致标记、冲突清单、每人计数、原始一致率与 Cohen kappa；或用 scripts/merge_screening.py 合并两份导出文件。',
            f'4. 校准：merge_TA 中 calibration_{cal["planned_sample_size"]} 区块——粘贴按 seed={cal["random_seed"]} 抽取的记录号，填写是否全部概念性分歧已解决；通过条件为原始一致率 ≥{E["threshold"]:.0%} 且分歧全部解决，kappa 仅描述。',
            '5. 冲突讨论后仍未解决者交预先指定的裁决者（姓名待定，不得虚构）。讨论记录写入 log。',
            '6. 全文阶段同理：screen_FT_reviewer_A/B → merge_FT。管理状态（重复、未获取、待翻译、待分类、预印本前沿登记）不是科学排除，不填 FT 代码；check_code_rule 列会提示违规。',
            '7. 任何未执行阶段的计数一律留空或记为 NOT YET AVAILABLE，绝不写 0。',
        ]),
        ('维护规则 Maintenance rule', [
            '先修改 JSON 模板，再运行脚本重新生成；不得手工修改生成的 xlsx 结构（sheet、列、下拉、保护）。',
            '工作表保护不设密码，仅用于防止误改结构。', '下拉列表全部来自 codes 工作表，codes 由 JSON 生成。']),
        ('模板解释 Interpretations of JSON fields', INTERPRETATIONS['screening']),
        ('模板不一致提示 Template issues detected', E['problems'] or ['none detected by the generator']),
    ])
    S['README'].sheet_properties.tabColor = '000000'

    wb.calculation.fullCalcOnLoad = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    spec = {
        'sheets': names,
        'headers': {'records_master': MASTER_COLUMNS, 'screen_TA_reviewer_A': ta_cols, 'screen_TA_reviewer_B': ta_cols,
                    'screen_FT_reviewer_A': ft_cols, 'screen_FT_reviewer_B': ft_cols, 'log': ['date', 'who', 'action', 'free_text']},
        'protected': ['README', 'records_master', 'screen_TA_reviewer_A', 'screen_TA_reviewer_B', 'screen_FT_reviewer_A',
                      'screen_FT_reviewer_B', 'merge_TA', 'merge_FT', 'codes'],
        'unprotected': ['log'],
        'unlocked_cols': {s: [L(i + 1) for i, hh in enumerate(cols) if hh not in ('record_id', 'title', 'check_code_rule')]
                          for s, cols in (('screen_TA_reviewer_A', ta_cols), ('screen_TA_reviewer_B', ta_cols),
                                          ('screen_FT_reviewer_A', ft_cols), ('screen_FT_reviewer_B', ft_cols))},
        'locked_cols': {s: ['A', 'B'] for s in ('screen_TA_reviewer_A', 'screen_TA_reviewer_B', 'screen_FT_reviewer_A', 'screen_FT_reviewer_B')},
        'empty_input_sheets': ['records_master', 'log'],
        'formula_input_sheets': ['screen_TA_reviewer_A', 'screen_TA_reviewer_B', 'screen_FT_reviewer_A', 'screen_FT_reviewer_B'],
        'dv': registry, 'names': list(codes.names), 'unlocked_cells': registry.get('_unlocked_cells', {}),
    }
    return spec, E


# ---------------------------------------------------------------------------------------------
# Extraction workbook
# ---------------------------------------------------------------------------------------------
def flatten(blank: dict, prefix: str = ''):
    for k, v in blank.items():
        if isinstance(v, dict):
            yield from flatten(v, prefix + k + '.')
        else:
            yield prefix + k, v


def split_targets(f):
    tabs = [x.strip() for x in f['destination_table'].split(';')]
    cols = [x.strip() for x in f['destination_column'].split(';')]
    if len(tabs) == 1:
        return [(tabs[0], c) for c in cols]
    if len(cols) == 1:
        return [(t, cols[0]) for t in tabs]
    return list(zip(tabs, cols))


def parse_new_field_defs(dic, specs):
    """Expand keys such as 'measurements.candidate_selection_rule / basis / prespecification_status'."""
    out = {}
    for key, text in dic.get('new_operational_fields_not_in_legacy_61', {}).items():
        table, rest = key.split('.', 1)
        cols = specs[table]['blank_record'].keys()
        m = re.match(r'(\w+) and (\w+) fields$', rest)
        if m:   # "effect_value_source and digitization fields"
            targets = [m.group(1)] + [c for c in cols if c.startswith(m.group(2)[:7])]
        else:
            parts = [p.strip() for p in rest.split('/')]
            targets = [parts[0]]
            stem = parts[0].split('_')
            for p in parts[1:]:
                # prefer the stem-prefixed expansion (".../ source_locator" after candidate_* means
                # candidate_source_locator), longest stem first; fall back to the literal column name
                for n in range(len(stem), 0, -1):
                    cand = '_'.join(stem[:n]) + '_' + p
                    if cand in cols:
                        targets.append(cand)
                        break
                else:
                    if p in cols:
                        targets.append(p)
        for t in targets:
            if t in cols:
                out[(table, t)] = (text, f'new_operational_fields_not_in_legacy_61["{key}"]')
    return out


def extraction_fields(src):
    dic = src['dictionary']['data']
    tmpl = src['extraction_template']['data']
    specs = dic['relational_model']['table_specs']
    vocab = dic['controlled_vocabulary']
    codes = list(dic['missing_value_semantics']['codes'])
    E = screening_enums(src)
    problems = []
    if list(specs) != list(tmpl['tables']):
        problems.append('Table order/set differs between data_dictionary.json and extraction_template.json')
    for t in specs:
        if t in tmpl['tables'] and specs[t]['blank_record'] != tmpl['tables'][t]['blank_record']:
            problems.append(f'blank_record of {t} differs between dictionary and template (dictionary order used)')
    legacy = {}
    for f in dic['legacy_61_fields']:
        for pair in split_targets(f):
            legacy[pair] = f
    legacy_by_col = {}
    for (t, c), f in legacy.items():
        legacy_by_col.setdefault(c, []).append((t, f))
    newdefs = parse_new_field_defs(dic, specs)
    cross = {('reports', 'primary_fulltext_exclusion_reason'): ('fulltext_exclusion_codes.json hierarchy', E['codes']),
             ('reports', 'consensus_screening_decision'): ('screening_log_template screening_dispositions', E['dispositions'])}
    used_vocab = set()
    tables = {}
    for t, spec in specs.items():
        pk = spec['primary_key']
        fks = {}
        for fk in spec.get('foreign_keys', []):
            m = re.match(r'(\w+)(\[\])?\s*->\s*([\w.]+)(.*)', fk)
            if m:
                fks[m.group(1)] = (m.group(3), 'optional' in m.group(4))
        fields = []
        for path, value in flatten(spec['blank_record']):
            col = path.split('.')[-1]
            parent = path.split('.')[0]
            is_child = '.' in path
            meta = {'table': t, 'field': path, 'is_list': isinstance(value, list), 'is_pk': path == pk,
                    'fk': fks.get(path), 'legacy': None}
            # ---- definition
            defn, dsrc = None, None
            lf = legacy.get((t, parent))
            if lf:
                meta['legacy'] = lf
                defn = lf['description'] + (f' — component `{col}`' if is_child else '')
                dsrc = f'legacy_61_fields["{lf["field"]}"] ({lf["module"]})'
            elif (t, path) in newdefs:
                defn, dsrc = newdefs[(t, path)]
            elif meta['is_pk']:
                defn, dsrc = f'Primary key of {t} ({spec["row_unit"]}).', 'table_specs.primary_key'
            elif meta['fk']:
                defn = f'Foreign key → {meta["fk"][0]}' + (' (optional)' if meta['fk'][1] else '') + '.'
                dsrc = 'table_specs.foreign_keys'
            elif path in legacy_by_col:
                ot, of = legacy_by_col[path][0]
                defn = of['description'] + f' (same-name legacy field, mapped to {ot})'
                dsrc = f'legacy_61_fields["{of["field"]}"] (same name, other table)'
            vkey = VOCAB_ALIASES.get((t, path), path if path in vocab else None)
            if not defn and vkey:
                defn, dsrc = f'Coded with controlled_vocabulary.{vkey}.', f'controlled_vocabulary.{vkey}'
            if not defn:
                defn, dsrc = 'No definition in data_dictionary.json (see extraction_manual.md); field name is self-describing.', 'NONE'
            if meta['is_pk'] and dsrc != 'table_specs.primary_key':
                defn += ' Primary key.'
            if meta['fk'] and dsrc != 'table_specs.foreign_keys':
                defn += f' Foreign key → {meta["fk"][0]}.'
            meta['definition'], meta['definition_source'] = defn, dsrc
            # ---- type
            ltype = lf['data_type'] if lf else None
            if ltype and is_child:
                ltype = f'string (component of {lf["data_type"]})'
            is_id = meta['is_pk'] or bool(meta['fk'])
            is_count = (ltype in ('integer_or_missing_code', 'separate_counts') or
                        (not lf and (path.startswith('n_') or path.endswith('_count') or path.endswith('_n'))))
            if ltype:
                typ = ltype
            elif is_id:
                typ = 'identifier' + (' array' if meta['is_list'] else '')
            elif meta['is_list']:
                typ = 'multi_enum (array)' if vkey else 'array'
            elif vkey:
                typ = 'enum'
            elif is_count:
                typ = 'integer_or_missing_code'
            else:
                typ = 'string'
            meta['type'], meta['is_count'], meta['is_id'] = typ, is_count, is_id
            # ---- allowed values / dropdown
            allowed, mode, list_name = None, 'none', None
            missing_ok = not is_id
            if lf and 'forbidden' in lf['missing_value_rule']:
                missing_ok = False
            if vkey:
                used_vocab.add(vkey)
                vals = list(vocab[vkey])
                if missing_ok and (lf or any(c in vals for c in codes)):
                    vals += [c for c in codes if c not in vals]
                allowed = vals
                if meta['is_list'] or (ltype == 'multi_enum'):
                    mode = 'warning (multi-valued, "; ")'
                elif ltype and ltype != 'enum':
                    mode = 'warning (free text allowed)'
                else:
                    mode = 'strict'
                list_name = f'voc_{t[:4]}_{re.sub(r"[^A-Za-z0-9]", "_", path)}'
                meta['vocab_key'] = vkey
            elif (t, path) in cross:
                label, vals = cross[(t, path)]
                allowed = list(vals) + [c for c in codes if c not in vals]
                mode = 'warning (cross-file list)'
                list_name = f'voc_{t[:4]}_{path}'
                meta['vocab_key'] = label
            else:
                meta['vocab_key'] = ''
            meta['allowed'], meta['mode'], meta['list_name'] = allowed, mode, list_name
            meta['missing_policy'] = ('forbidden (identifier)' if not missing_ok else
                                      f'{"/".join(codes)} after source review; blank = not yet extracted')
            meta['excel_format'] = 'General + integer/missing-code validation' if is_count else 'Text (@)'
            fields.append(meta)
        tables[t] = {'spec': spec, 'fields': fields}
    unmapped = [k for k in vocab if k not in used_vocab and k != 'missing_codes']
    for k in unmapped:
        problems.append(f'controlled_vocabulary.{k} has no matching column in any table (listed in codes sheet only)')
    undefined = [f'{t}.{m["field"]}' for t, v in tables.items() for m in v['fields'] if m['definition_source'] == 'NONE']
    return tables, codes, vocab, unmapped, undefined, problems


def build_extraction(src, out_path: Path, reviewer: str | None = None, registry=None):
    registry = {} if registry is None else registry
    tables, mcodes, vocab, unmapped, undefined, problems = extraction_fields(src)
    dic = src['dictionary']['data']
    tmpl = src['extraction_template']['data']
    n_tables = len(tables)
    wb = openpyxl.Workbook()
    ws_readme = wb.active
    ws_readme.title = 'README'
    sheets = {t: wb.create_sheet(t) for t in tables}
    ws_dd = wb.create_sheet('data_dictionary')
    ws_codes = wb.create_sheet('codes')
    codes = Codes(wb, ws_codes)
    ids_neq = ','.join(f'"{c}"' for c in mcodes)
    headers = {}
    for t, info in tables.items():
        ws = sheets[t]
        cols = [m['field'] for m in info['fields']]
        headers[t] = cols
        comments = {}
        for m in info['fields']:
            txt = [m['definition'], '', f'Type: {m["type"]}', f'Missing values: {m["missing_policy"]}',
                   f'Excel format: {m["excel_format"]}']
            if m['allowed']:
                txt.append(f'Allowed ({m["mode"]}): ' + ' | '.join(m['allowed']))
            if m['is_list']:
                txt.append(f'Multiple values: separate with "{MULTI_SEP}"')
            txt.append(f'Source: {m["definition_source"]}')
            comments[m['field']] = '\n'.join(txt)
        write_header(ws, cols, comments, {m['field']: ('formula' if m['is_id'] else 'header') for m in info['fields']})
        for i, m in enumerate(info['fields'], start=1):
            input_column(ws, i, text=not m['is_count'])
            if m['list_name']:
                if m['list_name'] not in codes.names:
                    codes.add(m['list_name'], m['allowed'], f'{t}.{m["field"]}',
                              note=f'Source: {m["vocab_key"]}; mode: {m["mode"]}')
                add_dv(ws, i, '=' + m['list_name'], style='stop' if m['mode'] == 'strict' else 'warning',
                       prompt=('; '.join(m['allowed']) if len('; '.join(m['allowed'])) < 240 else
                               f'Choose from codes sheet ({m["vocab_key"]}).') + ('  Multi: "; "' if m['is_list'] else ''),
                       title=m['field'], registry=registry, label=f'{t}.{m["field"]}')
            elif m['is_count']:
                c = f'{L(i)}2'
                add_dv(ws, i, f'IF(ISNUMBER({c}),AND({c}>=0,{c}=INT({c})),OR(' + ','.join(f'{c}="{x}"' for x in mcodes) + '))',
                       kind='custom', prompt=f'Integer ≥ 0, or {"/".join(mcodes)} after source review.', title=m['field'],
                       registry=registry, label=f'{t}.{m["field"]}')
            elif m['is_id']:
                c = f'{L(i)}2'
                add_dv(ws, i, f'AND(' + ','.join(f'{c}<>"{x}"' for x in mcodes) + ')', kind='custom',
                       prompt=f'Stable non-empty ID; {"/".join(mcodes)} are forbidden for identifiers.', title=m['field'],
                       registry=registry, label=f'{t}.{m["field"]}')
            width = 30 if any(k in m['field'] for k in ('notes', 'title', 'locator', 'definition', 'rule')) else 18
            ws.column_dimensions[L(i)].width = width
        ws.auto_filter.ref = f'A1:{L(len(cols))}1'
        protect(ws, allow_row_edit=True)
        ws.sheet_properties.tabColor = '2E75B6'
    # vocabularies not mapped to any column: still listed for reference
    for k in unmapped:
        codes.add(f'voc_unmapped_{k}', vocab[k], f'{k} (UNMAPPED: no column)',
                  note='controlled_vocabulary entry with no matching extraction column; reported as a template issue.')
    codes.add('voc_missing_codes', [f'{k}' for k in mcodes], 'missing codes',
              extra=[('meaning', [dic['missing_value_semantics']['codes'][k] for k in mcodes])])
    protect(ws_codes)
    ws_codes.sheet_properties.tabColor = '808080'

    dd_head = ['table', 'order', 'field', 'type', 'allowed_values', 'vocabulary_source', 'dropdown_mode',
               'missing_value_policy', 'excel_format', 'identifier', 'definition', 'definition_source',
               'legacy_field', 'legacy_module', 'legacy_missing_value_rule']
    write_header(ws_dd, dd_head)
    r = 2
    for t, info in tables.items():
        for i, m in enumerate(info['fields'], start=1):
            lf = m['legacy'] or {}
            row = [t, i, m['field'], m['type'], ' | '.join(m['allowed']) if m['allowed'] else '', m['vocab_key'],
                   m['mode'], m['missing_policy'], m['excel_format'],
                   'PK' if m['is_pk'] else ('FK → ' + m['fk'][0] if m['fk'] else ''), m['definition'],
                   m['definition_source'], lf.get('field', ''), lf.get('module', ''), lf.get('missing_value_rule', '')]
            for j, v in enumerate(row, start=1):
                c = ws_dd.cell(row=r, column=j, value=v)
                c.alignment = WRAP
            r += 1
    for j, w in enumerate([20, 6, 30, 18, 50, 24, 22, 30, 22, 26, 60, 34, 18, 12, 40], start=1):
        ws_dd.column_dimensions[L(j)].width = w
    ws_dd.auto_filter.ref = f'A1:{L(len(dd_head))}{r - 1}'
    protect(ws_dd)

    gen = _dt.date.today().isoformat()
    label = reviewer or 'UNASSIGNED (master copy: generate per-extractor files with --reviewer A / --reviewer B)'
    tbl_rows = [(t, f'row unit: {i["spec"]["row_unit"]} | PK: {i["spec"]["primary_key"]} | FK: '
                    f'{"; ".join(i["spec"].get("foreign_keys", [])) or "—"} | {len(i["fields"])} columns | {i["spec"].get("description", "")}')
                for t, i in tables.items()]
    write_readme(ws_readme, f'提取工作簿 Extraction workbook — reviewer {reviewer or "(unassigned)"}', [
        ('生成信息 Generation', [('extractor / reviewer', label), ('generated_on', gen),
                                ('generator', f'scripts/build_workbooks.py v{SCRIPT_VERSION}'),
                                ('tables', f'{n_tables} tables (confirmed from data_dictionary.json relational_model.table_specs)'),
                                ('status', tmpl['metadata'].get('status', '')),
                                ('regenerate', 'python3 scripts/build_workbooks.py --reviewer A --reviewer B')]),
        ('源模板版本 Source templates', source_rows(src, EXTRACTION_SOURCES)),
        ('使用说明 How to use', [
            '每位提取者只使用本人文件（extraction_workbook_A.xlsx / _B.xlsx），独立完成后再由数据管理员比对、讨论、形成共识；分歧交预先指定的裁决者（姓名待定）。',
            '录入顺序：study_families → reports → cohorts → report_cohort_links（填 scope_streams）→ sample_sets → measurements（sample_set_ids）→ precision_validation（七个父域每域一行，不计总分）→ extraction_provenance。',
            '空白 = 尚未提取；核查原文后才填 NR（未报告）/ NA（不适用）/ UNCLEAR（全文仍不清楚）。主键与必需关联键必须是稳定非空 ID，禁止使用 NR/NA/UNCLEAR。',
            f'数组字段（JSON 列表）在一个单元格内用 "{MULTI_SEP}" 分隔；嵌套对象拆成 "父.子" 列（如 analyte_id.original_label）。',
            '除计数列外，所有列均为文本格式（@），防止 Excel 把 SEPT7/MARCH1 等基因名、1-2 等比值、DOI/PMID 及前导零自动转换。粘贴时请用“仅粘贴数值”；不要直接双击打开 CSV，应通过“数据→自文本/CSV”并将列设为文本。',
            '计数列只接受非负整数或 NR/NA/UNCLEAR；参与者、供者、生物样本、细胞、技术重复、分装数分别记录，不能互相替代。',
            '表头批注给出字段定义、类型、允许值与来源；data_dictionary 工作表列出全部字段；codes 工作表列出全部下拉值。',
            '外键一致性（report-study-cohort 三元组、sample_set_ids 等）需在合并后用脚本检查；工作簿内不做跨表校验。',
        ]),
        ('表 Tables', tbl_rows),
        ('关系规则 Relational rules (data_dictionary.json)', dic['relational_model'].get('rules', [])),
        ('完整性检查 Referential integrity checks (extraction_template.json)', tmpl.get('referential_integrity_checks', [])),
        ('护栏 Guardrails', tmpl.get('guardrails', [])),
        ('维护规则 Maintenance rule', ['先修改 JSON 模板，再运行脚本重新生成；不得手工修改生成的 xlsx 结构。工作表保护无密码，仅防误改结构；行可插入/删除，可筛选排序。']),
        ('模板解释 Interpretations of JSON fields', INTERPRETATIONS['extraction']),
        ('无定义字段 Columns without a dictionary definition', [', '.join(undefined)] if undefined else ['none']),
        ('模板不一致提示 Template issues detected', problems or ['none detected by the generator']),
    ])
    ws_readme.sheet_properties.tabColor = '000000'
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    spec = {'sheets': ['README'] + list(tables) + ['data_dictionary', 'codes'], 'headers': headers,
            'protected': ['README'] + list(tables) + ['data_dictionary', 'codes'], 'unprotected': [],
            'unlocked_cols': {t: [L(i + 1) for i in range(len(c))] for t, c in headers.items()},
            'locked_cols': {}, 'empty_input_sheets': list(tables), 'formula_input_sheets': [],
            'dv': registry, 'names': list(codes.names), 'unlocked_cells': {},
            'text_cols': {t: [L(i + 1) for i, m in enumerate(tables[t]['fields']) if not m['is_count']] for t in tables},
            'reviewer': reviewer}
    return spec, {'n_tables': n_tables, 'unmapped': unmapped, 'undefined': undefined, 'problems': problems,
                  'n_fields': sum(len(v['fields']) for v in tables.values())}


# ---------------------------------------------------------------------------------------------
# Appraisal workbook
# ---------------------------------------------------------------------------------------------
def tool_sheet_name(tool_id: str, used: set) -> str:
    name = tool_id
    if len(name) > 31:
        name = re.sub(r'_CURRENT$', '', name)
    name = name[:31]
    base, i = name, 2
    while name in used:
        name = f'{base[:28]}_{i}'
        i += 1
    used.add(name)
    return name


def flatten_profile(profile: dict, prefix=''):
    for k, v in profile.items():
        if isinstance(v, dict):
            yield from flatten_profile(v, prefix + k + '.')
        else:
            yield prefix + k


def build_appraisal(src, out_path: Path, registry=None):
    registry = {} if registry is None else registry
    a = src['appraisal_template']['data']
    sc = a['status_codes']
    rec = a['appraisal_record_template']
    tools = a['tool_selection']['tool_registry']
    wb = openpyxl.Workbook()
    ws_readme = wb.active
    ws_readme.title = 'README'
    keys = []
    for k in rec:
        if k.endswith('_id'):
            keys.append(k)
        else:
            break
    scalar = [k for k, v in rec.items() if not isinstance(v, (list, dict)) and not (isinstance(v, str) and v)]
    notes = {k: v for k, v in rec.items() if isinstance(v, str) and v}
    lists = {k: v for k, v in rec.items() if isinstance(v, list)}
    profile_cols = list(flatten_profile(rec['precision_profile']))
    used = {'README', 'appraisal_records', 'official_tool_outputs', 'initial_disagreements', 'consensus_items',
            'precision_profile', 'tool_registry', 'codes'}
    tool_sheets = {t['tool_id']: tool_sheet_name(t['tool_id'], used) for t in tools}
    ws_rec = wb.create_sheet('appraisal_records')
    ws_tools = {tid: wb.create_sheet(name) for tid, name in tool_sheets.items()}
    other = {k: wb.create_sheet(k) for k in lists if k != 'items'}
    ws_prof = wb.create_sheet('precision_profile')
    ws_reg = wb.create_sheet('tool_registry')
    ws_codes = wb.create_sheet('codes')
    codes = Codes(wb, ws_codes)
    for k, vals in sc.items():
        if isinstance(vals, list):
            codes.add(f'ap_{k}', vals, k, note=sc.get('coding_note', ''))
    codes.add('ap_tool_id', [t['tool_id'] for t in tools], 'tool_id (tool_registry)',
              extra=[('tool_name', [t['tool_name'] for t in tools])])
    for t in tools:
        codes.add(f'ap_only_{re.sub(r"[^A-Za-z0-9]", "_", t["tool_id"])}', [t['tool_id']], f'tool_id for sheet {tool_sheets[t["tool_id"]]}')
    protect(ws_codes)

    def code_for(field: str):
        leaf = field.split('.')[-1]
        if leaf in sc and isinstance(sc[leaf], list):
            return leaf
        alias = APPRAISAL_ALIASES.get(leaf)
        return alias if alias in sc else None

    tool_ids = {t['tool_id'] for t in tools}
    hints = {}
    for lk, lv in lists.items():
        for f, v in (lv[0] if lv else {}).items():
            if isinstance(v, str) and v:
                hints[f] = v
    headers = {}

    def data_sheet(ws, cols, comments=None, tool_id=None):
        headers[ws.title] = cols
        cm = {c: (comments or {}).get(c) or hints.get(c) for c in cols}
        for c in cols:
            ck = code_for(c)
            if ck:
                cm[c] = ((cm[c] + '\n') if cm[c] else '') + f'Values: {" | ".join(sc[ck])} (status_codes.{ck}). ' + sc.get('coding_note', '')
        write_header(ws, cols, cm)
        for i, c in enumerate(cols, start=1):
            input_column(ws, i)
            ck = code_for(c)
            if c == 'tool_id':
                nm = f'ap_only_{re.sub(r"[^A-Za-z0-9]", "_", tool_id)}' if tool_id else 'ap_tool_id'
                add_dv(ws, i, '=' + nm, registry=registry, label=f'{ws.title}.{c}', title='tool_id',
                       prompt='Tool from critical_appraisal_template tool_registry.')
            elif ck:
                warn = c.split('.')[-1] in APPRAISAL_WARNING_FIELDS
                add_dv(ws, i, f'=ap_{ck}', style='warning' if warn else 'stop', registry=registry, label=f'{ws.title}.{c}',
                       title=c.split('.')[-1], prompt=' | '.join(sc[ck]))
            elif c.endswith('_id') and c != 'tool_id':
                ws.column_dimensions[L(i)].width = 16
        ws.auto_filter.ref = f'A1:{L(len(cols))}1'
        protect(ws, allow_row_edit=True)

    data_sheet(ws_rec, scalar, {'tool_id': 'Choose the tool from the design (tool_selection).'})
    item_fields = list(lists['items'][0].keys()) if lists.get('items') else []
    for tid, ws in ws_tools.items():
        data_sheet(ws, keys + ['tool_id'] + item_fields, tool_id=tid)
        ws.sheet_properties.tabColor = '2E75B6'
    for k, ws in other.items():
        data_sheet(ws, keys + ['tool_id'] + list(lists[k][0].keys()))
    data_sheet(ws_prof, keys + profile_cols)

    reg_cols = []
    for t in tools:
        for k in t:
            if k not in reg_cols:
                reg_cols.append(k)
    reg_cols.append('sheet')
    write_header(ws_reg, reg_cols)
    for i, t in enumerate(tools, start=2):
        for j, k in enumerate(reg_cols, start=1):
            v = tool_sheets[t['tool_id']] if k == 'sheet' else t.get(k, '')
            ws_reg.cell(row=i, column=j, value=v).alignment = WRAP
    r = len(tools) + 3
    ws_reg.cell(row=r, column=1, value='tool_selection (design → tool)').font = Font(bold=True)
    for k, v in a['tool_selection'].items():
        if isinstance(v, str) and k not in ('source_directory', 'source_accessed_on'):
            r += 1
            ws_reg.cell(row=r, column=1, value=k).alignment = WRAP
            ws_reg.cell(row=r, column=2, value=v).alignment = WRAP
    for j in range(1, len(reg_cols) + 1):
        ws_reg.column_dimensions[L(j)].width = 40
    protect(ws_reg)

    problems = [f'tool_selection maps to unknown tool_id {v}' for k, v in a['tool_selection'].items()
                if isinstance(v, str) and k.isidentifier() and v.isupper() and v not in tool_ids and '_' in v]
    no_code = [f'{s}.{c}' for s, cols in headers.items() for c in cols if c.split('.')[-1] in
               ('status', 'use_tested') and not code_for(c)]
    problems += [f'no status code list for {x}' for x in no_code]
    gen = _dt.date.today().isoformat()
    write_readme(ws_readme, '严格评价工作簿 Critical appraisal workbook (blank)', [
        ('生成信息 Generation', [('generated_on', gen), ('generator', f'scripts/build_workbooks.py v{SCRIPT_VERSION}'),
                                ('template_version', a.get('template_version')), ('protocol_version', a.get('protocol_version')),
                                ('regenerate', 'python3 scripts/build_workbooks.py --reviewer A --reviewer B')]),
        ('源模板版本 Source templates', source_rows(src, APPRAISAL_SOURCES)),
        ('目的 Purpose', [a.get('purpose', ''), a.get('record_unit', '')]),
        ('使用说明 How to use', [
            '先在 appraisal_records 建一行（report_id/study_id/cohort_id/analysis_id + 设计与工具 + 两位评价者）。工具按设计选择（tool_registry 工作表）；不要对所有研究套用 RCT 工具。',
            '在对应工具工作表（每个工具一张）逐条记录条目：item_id_as_printed 按存档的官方表格原样填写（本工作簿不复制条目原文），tool_response_verbatim 保留原表答案，response_normalized 用 YES/NO/UNCLEAR/NA，reporting_status 用 REPORTED/NR/CONFLICTING/NOT_APPLICABLE。未报告 = UNCLEAR + NR，不是 NO；NA 需写理由。',
            '两位评价者各自独立填写（appraiser_id 区分），然后在 initial_disagreements 记录分歧，consensus_items 记录共识；无法解决者交裁决者（姓名待定）。',
            'official_tool_outputs 记录工具自身定义的领域/总体判断（如 PROBAST+AI 领域）；precision_profile 每条评价记录一行，七个维度分别描述，不求和、不排名、不评分。',
            '不计总分、不排序、不据评价结果排除研究、不做 GRADE。',
        ]),
        ('固定说明 Template notes', [f'{k}: {v}' for k, v in notes.items()] + [f'process.{k}: {v}' for k, v in a.get('process', {}).items()]),
        ('维护规则 Maintenance rule', ['先修改 JSON 模板，再运行脚本重新生成；不得手工修改生成的 xlsx 结构。工作表保护无密码。']),
        ('模板解释 Interpretations of JSON fields', INTERPRETATIONS['appraisal']),
        ('模板不一致提示 Template issues detected', problems or ['none detected by the generator']),
    ])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    data_sheets = ['appraisal_records'] + list(tool_sheets.values()) + list(other) + ['precision_profile']
    spec = {'sheets': [ws.title for ws in wb.worksheets], 'headers': headers,
            'protected': [ws.title for ws in wb.worksheets], 'unprotected': [],
            'unlocked_cols': {s: [L(i + 1) for i in range(len(c))] for s, c in headers.items()}, 'locked_cols': {},
            'empty_input_sheets': data_sheets, 'formula_input_sheets': [], 'dv': registry,
            'names': list(codes.names), 'unlocked_cells': {}}
    return spec, {'tools': tool_sheets, 'problems': problems}


# ---------------------------------------------------------------------------------------------
# Verification (reload with openpyxl)
# ---------------------------------------------------------------------------------------------
def verify(path: Path, spec: dict) -> dict:
    wb = openpyxl.load_workbook(path)
    errs = []
    if wb.sheetnames != spec['sheets']:
        errs.append(f'sheet names {wb.sheetnames} != {spec["sheets"]}')
    for s, heads in spec['headers'].items():
        ws = wb[s]
        got = [ws.cell(row=1, column=i + 1).value for i in range(len(heads))]
        if got != heads:
            errs.append(f'{s}: header mismatch')
        if ws.cell(row=1, column=len(heads) + 1).value not in (None, '') and s not in ('merge_TA', 'merge_FT'):
            errs.append(f'{s}: extra header cells')
        for i in range(len(heads)):
            if ws.cell(row=1, column=i + 1).protection.locked is not True:
                errs.append(f'{s}: header {heads[i]} not locked')
    n_dv = 0
    for s, items in spec['dv'].items():
        if s.startswith('_'):
            continue
        ws = wb[s]
        have = [(str(dv.sqref), dv.formula1, dv.errorStyle or 'stop') for dv in ws.data_validations.dataValidation]
        for it in items:
            ok = any(f1 == it['formula1'] and any(part.startswith(it['col']) and re.match(rf'^{it["col"]}\d', part)
                                                  for part in sq.split()) for sq, f1, _ in have)
            if not ok:
                errs.append(f'{s}: missing data validation on column {it["col"]} ({it["label"]})')
            n_dv += 1
    for s in spec['protected']:
        if not wb[s].protection.sheet:
            errs.append(f'{s}: not protected')
        if wb[s].protection.password:
            errs.append(f'{s}: has a password')
    for s in spec['unprotected']:
        if wb[s].protection.sheet:
            errs.append(f'{s}: unexpectedly protected')
    for s, cols in spec['unlocked_cols'].items():
        for c in cols:
            if wb[s].column_dimensions[c].protection.locked is not False:
                errs.append(f'{s}: input column {c} not unlocked')
    for s, cols in spec.get('text_cols', {}).items():
        for c in cols:
            if wb[s].column_dimensions[c].number_format != '@':
                errs.append(f'{s}: column {c} not Text format')
    for s, cols in spec['locked_cols'].items():
        for c in cols:
            if wb[s][f'{c}2'].protection.locked is not True:
                errs.append(f'{s}: formula column {c} not locked')
    for s, cells in spec['unlocked_cells'].items():
        for ref in cells:
            if wb[s][ref].protection.locked is not False:
                errs.append(f'{s}: input cell {ref} not unlocked')
    for s in spec['empty_input_sheets']:
        if wb[s].max_row > 1:
            errs.append(f'{s}: has data rows ({wb[s].max_row - 1})')
    for s in spec['formula_input_sheets']:
        ws = wb[s]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value is not None and not (isinstance(c.value, str) and c.value.startswith('=')):
                    errs.append(f'{s}: non-formula value at {c.coordinate}')
                    break
    for n in spec['names']:
        if n not in wb.defined_names:
            errs.append(f'defined name {n} missing')
    return {'file': path.name, 'sheets': len(wb.sheetnames), 'data_validations_checked': n_dv,
            'protected_sheets': sum(1 for s in wb.sheetnames if wb[s].protection.sheet), 'errors': errs}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--out-dir', type=Path, default=DEFAULT_OUT)
    ap.add_argument('--rows', type=int, default=5000, help='formula rows in the screening workbook (>= records in pool)')
    ap.add_argument('--conflict-rows', type=int, default=1000)
    ap.add_argument('--reviewer', action='append', choices=['A', 'B'], default=[],
                    help='also build extraction_workbook_<R>.xlsx stamped with this extractor label (repeatable)')
    ap.add_argument('--only', choices=['screening', 'extraction', 'appraisal'], action='append', default=[])
    ap.add_argument('--no-verify', action='store_true')
    args = ap.parse_args(argv)
    src = load_sources()
    only = set(args.only) or {'screening', 'extraction', 'appraisal'}
    built, failed = [], False
    if 'screening' in only:
        p = args.out_dir / 'screening_workbook.xlsx'
        spec, E = build_screening(src, p, rows=args.rows, conflict_rows=args.conflict_rows)
        built.append((p, spec, {'TA decisions': E['ta'], 'FT decisions': E['ft'], 'issues': E['problems']}))
    if 'extraction' in only:
        targets = [None] if not (args.reviewer and args.only) else []
        targets += sorted(set(args.reviewer))
        for rv in targets:
            p = args.out_dir / (f'extraction_workbook_{rv}.xlsx' if rv else 'extraction_workbook.xlsx')
            spec, info = build_extraction(src, p, reviewer=rv)
            built.append((p, spec, {'tables': info['n_tables'], 'fields': info['n_fields'],
                                    'undefined_columns': len(info['undefined']), 'issues': info['problems']}))
    if 'appraisal' in only:
        p = args.out_dir / 'appraisal_workbook.xlsx'
        spec, info = build_appraisal(src, p)
        built.append((p, spec, {'tool_sheets': list(info['tools'].values()), 'issues': info['problems']}))
    for p, spec, info in built:
        line = {'written': str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p)}
        if not args.no_verify:
            v = verify(p, spec)
            line.update({k: v[k] for k in ('sheets', 'data_validations_checked', 'protected_sheets')})
            line['verify'] = 'OK' if not v['errors'] else v['errors']
            failed |= bool(v['errors'])
        line.update({k: v for k, v in info.items() if k != 'issues'})
        line['template_issues'] = len(info.get('issues', []))
        print(json.dumps(line, ensure_ascii=False))
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
