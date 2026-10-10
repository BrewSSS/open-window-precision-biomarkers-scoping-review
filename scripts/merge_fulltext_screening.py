#!/usr/bin/env python3
"""Merge the two locked full-text screening workbooks (ft_screen_B.xlsx, ft_screen_C.xlsx).

Reads the 'screen' sheet of both workbooks (built by scripts/build_fulltext_workbooks.py), checks completeness and
coding rules (a primary code only with EXCLUDE, and EXCLUDE always with a code), and reports:
  * raw agreement and Cohen's kappa on the disposition, on the primary FT code (both EXCLUDE), on the age-rule
    check and on validation-element confirmation; kappa is descriptive;
  * a titles-only conflict list for D (ft_conflicts_for_D.csv); reviewer comments stay in ignored workbooks;
  * PRISMA-ready counts over screened reports, with unscreened records reported separately.
Outputs (in --out-dir, default the workbooks' folder): ft_merged.csv, ft_conflicts_for_D.csv, ft_merge_summary.json,
ft_merge_summary.md.

Usage:
  python3 scripts/merge_fulltext_screening.py --b <ft_screen_B.xlsx> --c <ft_screen_C.xlsx> [--resolutions <csv>]
  python3 scripts/merge_fulltext_screening.py --self-test     # synthetic fills under /tmp; writes nothing in the repo
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_fulltext_workbooks as bfw  # noqa: E402

INCLUDE = {"INCLUDE_A", "INCLUDE_B", "INCLUDE_A_AND_B"}
ADMIN = {"RETAIN_BACKGROUND", "AWAITING_CLASSIFICATION", "NOT_RETRIEVED", "DUPLICATE_REPORT"}


def read_wb(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["screen"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    if head != bfw.COLUMNS:
        sys.exit(f"{path}: unexpected header")
    out = {}
    for row in it:
        if row[0] is None:
            continue
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, row)}
        out[r["record_id"]] = r
    return out


def kappa(pairs: list[tuple[str, str]]) -> float | None:
    n = len(pairs)
    if not n:
        return None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else round((po - pe) / (1 - pe), 4)


def agree(pairs):
    n = len(pairs)
    k = sum(a == b for a, b in pairs)
    return {"n": n, "agree": k, "raw_agreement": round(k / n, 4) if n else None, "kappa": kappa(pairs)}


def rule_problems(r: dict) -> list[str]:
    p = []
    d, c = r["your_disposition"], r["your_primary_code"]
    if not d:
        p.append("disposition missing")
    elif d not in bfw.DISPOSITIONS:
        p.append(f"unknown disposition {d}")
    if c and d != "EXCLUDE":
        p.append("primary code without EXCLUDE")
    if d == "EXCLUDE" and not c:
        p.append("EXCLUDE without primary code")
    if c and c not in bfw.ft_codes():
        p.append(f"unknown code {c}")
    if d == "DUPLICATE_REPORT" and not r["your_comment"]:
        p.append("DUPLICATE_REPORT without retained record in comment")
    return p


def decision(r: dict) -> str:
    return r["your_disposition"] + ("|" + r["your_primary_code"] if r["your_disposition"] == "EXCLUDE" else "")


def prisma(dec: dict[str, tuple[str, str]]) -> dict:
    disp = Counter(d for d, _ in dec.values())
    codes = Counter(c for d, c in dec.values() if d == "EXCLUDE")
    order = bfw.ft_codes()
    return {
        "n": len(dec),
        "by_disposition": {k: disp.get(k, 0) for k in bfw.DISPOSITIONS},
        "excluded_by_ft_code": {k: codes.get(k, 0) for k in order},
        "included_reports": sum(disp[k] for k in INCLUDE),
        "stream_A_reports": disp["INCLUDE_A"] + disp["INCLUDE_A_AND_B"],
        "stream_B_reports": disp["INCLUDE_B"] + disp["INCLUDE_A_AND_B"],
        "sought_not_retrieved": disp["NOT_RETRIEVED"],
        "assessed_full_text": len(dec) - disp["NOT_RETRIEVED"] - disp["DUPLICATE_REPORT"],
        "note": "A/B stream totals overlap through INCLUDE_A_AND_B; cohort-level counts come from extraction linkage.",
    }


def merge(b_path: Path, c_path: Path, out_dir: Path, resolutions: Path | None = None) -> dict:
    B, C = read_wb(b_path), read_wb(c_path)
    if set(B) != set(C):
        sys.exit(f"record sets differ: only B {sorted(set(B) - set(C))[:5]}, only C {sorted(set(C) - set(B))[:5]}")
    ids = sorted(B)
    complete = [i for i in ids if B[i]["your_disposition"] and C[i]["your_disposition"]]
    complete_set = set(complete)
    incomplete = [i for i in ids if i not in complete_set]
    problems = {rv: {i: rule_problems(W[i]) for i in complete if rule_problems(W[i])} for rv, W in (("B", B), ("C", C))}
    pairs_disp = [(B[i]["your_disposition"], C[i]["your_disposition"]) for i in complete]
    pairs_dec = [(decision(B[i]), decision(C[i])) for i in complete]
    bin_ = lambda d: "include" if d in INCLUDE else ("exclude" if d == "EXCLUDE" else "other")  # noqa: E731
    both_ex = [i for i in complete if B[i]["your_disposition"] == C[i]["your_disposition"] == "EXCLUDE"]
    stats = {
        "disposition": agree(pairs_disp),
        "disposition_with_ft_code": agree(pairs_dec),
        "include_exclude_other": agree([(bin_(a), bin_(b)) for a, b in pairs_disp]),
        "ft_code_when_both_exclude": agree([(B[i]["your_primary_code"], C[i]["your_primary_code"]) for i in both_ex]),
        "age_rule_check": agree([(B[i]["your_age_rule_check"], C[i]["your_age_rule_check"]) for i in complete
                                 if B[i]["your_age_rule_check"] and C[i]["your_age_rule_check"]]),
        "validation_element_confirmed": agree([(B[i]["your_validation_element_confirmed"], C[i]["your_validation_element_confirmed"])
                                               for i in complete if B[i]["your_validation_element_confirmed"] and C[i]["your_validation_element_confirmed"]]),
    }
    res = {}
    if resolutions and not resolutions.exists():
        sys.exit(f"resolution file missing: {resolutions}")
    if resolutions:
        for r in csv.DictReader(resolutions.open(newline="", encoding="utf-8")):
            if r.get("final_disposition"):
                rid, disp = r["record_id"], r["final_disposition"]
                code = r.get("final_primary_code", "") or ""
                if rid not in B or disp not in bfw.DISPOSITIONS or (disp == "EXCLUDE") != bool(code) or (code and code not in bfw.ft_codes()):
                    sys.exit(f"invalid D resolution for {rid}: {disp} {code}")
                if rid in res and res[rid] != (disp, code):
                    sys.exit(f"conflicting D resolutions for {rid}")
                res[rid] = (disp, code)
    agreed, final, conflicts, merged = {}, {}, [], []
    for i in ids:
        b, c = B[i], C[i]
        same = decision(b) == decision(c) and b["your_disposition"] != ""
        if same and not (problems["B"].get(i) or problems["C"].get(i)):
            agreed[i] = (b["your_disposition"], b["your_primary_code"] if b["your_disposition"] == "EXCLUDE" else "")
        flags = [x for x in ("age_rule_check", "validation_element_confirmed") if b["your_" + x] != c["your_" + x]]
        if i in complete and (not same or problems["B"].get(i) or problems["C"].get(i)):
            conflicts.append({"record_id": i, "title": b["title"], "doi": b["doi"],
                              "B_disposition": b["your_disposition"], "B_code": b["your_primary_code"],
                              "C_disposition": c["your_disposition"], "C_code": c["your_primary_code"],
                              "rule_problems": "; ".join([f"B: {x}" for x in problems["B"].get(i, [])] + [f"C: {x}" for x in problems["C"].get(i, [])]),
                              "final_disposition": res.get(i, ("", ""))[0], "final_primary_code": res.get(i, ("", ""))[1]})
        f = res.get(i) or agreed.get(i)
        if (f and not (problems["B"].get(i) or problems["C"].get(i))) or i in res:
            final[i] = f
        merged.append({"record_id": i, "title": b["title"],
                       **{f"{rv}_{k}": W[i]["your_" + k] for rv, W in (("B", B), ("C", C))
                          for k in ("disposition", "primary_code", "age_rule_check", "validation_element_confirmed")},
                       "agreed": "yes" if same else "no", "secondary_disagreement": ";".join(flags),
                       "final_disposition": final.get(i, ("", ""))[0], "final_primary_code": final.get(i, ("", ""))[1],
                       "final_source": "D" if i in res else ("agreed" if i in final else "")})
    out_dir.mkdir(parents=True, exist_ok=True)
    unresolved_ids = {r["record_id"] for r in conflicts if not r["final_disposition"]}
    if set(res) - {r["record_id"] for r in conflicts}:
        sys.exit(f"D resolutions are not screened conflicts: {sorted(set(res) - {r['record_id'] for r in conflicts})[:5]}")
    conflict_path = out_dir / "ft_conflicts_for_D.csv"
    if conflict_path.exists():
        with conflict_path.open(newline="", encoding="utf-8") as fh:
            old_rows = list(csv.DictReader(fh))
        conflict_ids = {r["record_id"] for r in conflicts}
        external_complete_resolutions = (resolutions is not None and resolutions.resolve() != conflict_path.resolve()
                                         and not unresolved_ids and set(res) == conflict_ids)
        if any(r.get("final_disposition") for r in old_rows) and not external_complete_resolutions:
            conflict_path = out_dir / "ft_conflicts_for_D_pending.csv"
        else:
            backup_dir = out_dir / "_backup_ft_merge_canonical"
            backup_dir.mkdir(parents=True, exist_ok=True)
            digest = hashlib.sha256(conflict_path.read_bytes()).hexdigest()[:16]
            backup_path = backup_dir / f"ft_conflicts_for_D_{digest}.csv"
            if not backup_path.exists():
                shutil.copy2(conflict_path, backup_path)
    for path, rows in ((out_dir / "ft_merged.csv", merged), (conflict_path, conflicts)):
        with path.open("w", newline="", encoding="utf-8") as fh:
            cols = list(rows[0]) if rows else ["record_id"]
            w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)
    n_partial = sum(bool(B[i]["your_disposition"]) != bool(C[i]["your_disposition"]) for i in incomplete)
    screened_final_complete = len(final) == len(complete) and not unresolved_ids and not n_partial
    summary = {
        "inputs": {rv: {"path": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for rv, p in (("B", b_path), ("C", c_path))},
        "n_records": len(ids),
        "n_screened_pairs": len(complete),
        "n_without_paired_decision": len(incomplete),
        "n_with_one_reviewer_decision": n_partial,
        "n_complete_pairs": len(complete),
        "rule_problems": {rv: len(v) for rv, v in problems.items()},
        "agreement": stats,
        "n_conflicts_for_D": len(conflicts),
        "n_resolved_by_D": len(res),
        "n_unresolved_screened_conflicts": len(unresolved_ids),
        "conflicts_csv": conflict_path.name,
        "prisma_agreed_only": prisma(agreed),
        "prisma_decided_to_date": prisma(final),
        "prisma_final": prisma(final) if screened_final_complete else None,
        "final_complete": screened_final_complete,
        "final_scope": "paired screened reports only; records without paired decisions are outside these disposition counts",
        "kappa_note": "Cohen's kappa is descriptive (prevalence-sensitive); full-text agreement has no pass threshold (calibration_plan.md).",
    }
    (out_dir / "ft_merge_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n")
    s = stats
    md = [f"# Full-text screening merge", "",
          f"Scope {len(ids)} records; paired screened reports {len(complete)}; without paired decisions {len(incomplete)} "
          f"(one-sided {n_partial}). The latter are not counted as exclusions or adjudication conflicts. "
          f"Rule problems among screened reports: B {summary['rule_problems']['B']}, C {summary['rule_problems']['C']}.",
          f"Disposition agreement {s['disposition']['agree']}/{s['disposition']['n']} (kappa {s['disposition']['kappa']}); "
          f"with FT code {s['disposition_with_ft_code']['agree']}/{s['disposition_with_ft_code']['n']}; "
          f"FT code when both exclude {s['ft_code_when_both_exclude']['agree']}/{s['ft_code_when_both_exclude']['n']}.",
          f"Age-rule check {s['age_rule_check']['agree']}/{s['age_rule_check']['n']}; validation element {s['validation_element_confirmed']['agree']}/{s['validation_element_confirmed']['n']}.",
          f"Conflicts for D: {len(conflicts)} (resolved {len(res)}, pending {len(unresolved_ids)}). "
          f"Final counts cover paired screened reports only.", "",
          "| Disposition | agreed | final |", "|---|---|---|"]
    pf = summary["prisma_final"] or {"by_disposition": {}, "excluded_by_ft_code": {}}
    for k in bfw.DISPOSITIONS:
        md.append(f"| {k} | {summary['prisma_agreed_only']['by_disposition'][k]} | {pf['by_disposition'].get(k, 'pending')} |")
    md += ["", "| FT code (EXCLUDE) | agreed | final |", "|---|---|---|"]
    for k in bfw.ft_codes():
        md.append(f"| {k} | {summary['prisma_agreed_only']['excluded_by_ft_code'][k]} | {pf['excluded_by_ft_code'].get(k, 'pending')} |")
    (out_dir / "ft_merge_summary.md").write_text("\n".join(md) + "\n")
    return summary


def self_test() -> None:
    rnd = random.Random(20261008)
    codes = bfw.ft_codes()
    with tempfile.TemporaryDirectory(prefix="ft_selftest_", dir="/tmp") as td:
        td = Path(td)
        src = td / "v.csv"
        with src.open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["record_id", "title", "journal", "year", "doi", "pmid", "final_layer", "final_E5_subtypes"])
            for k in range(60):
                w.writerow([f"FS-{k:06d}", f"Synthetic title {k}", "J", 2020, f"10.0/x{k}", str(k), "V", "omics_discovery"])
        bfw.main(["--input", str(src), "--filter", "final_layer=V", "--out-dir", str(td), "--meta", str(td / "none.csv")])
        truth = {}
        for k in range(50):
            d = rnd.choice(bfw.DISPOSITIONS)
            truth[f"FS-{k:06d}"] = (d, rnd.choice(codes) if d == "EXCLUDE" else "")
        planted_conflicts = set()
        for rv in ("B", "C"):
            wb = load_workbook(td / f"ft_screen_{rv}.xlsx")
            ws = wb["screen"]
            col = {c.value: c.column for c in ws[1]}
            for row in range(2, ws.max_row + 1):
                rid = ws.cell(row, 1).value
                if int(rid[3:]) >= 50:  # planted unscreened records: never D conflicts or exclusions
                    continue
                d, c = truth[rid]
                k = int(rid[3:])
                if rv == "C" and k % 7 == 0:  # planted disagreement
                    d, c = ("EXCLUDE", codes[3]) if d != "EXCLUDE" else ("INCLUDE_A", "")
                    planted_conflicts.add(rid)
                if rv == "B" and k == 3:  # planted rule problem
                    d, c = "INCLUDE_A", codes[0]
                    planted_conflicts.add(rid)
                ws.cell(row, col["your_disposition"], d)
                ws.cell(row, col["your_primary_code"], c or None)
                ws.cell(row, col["your_age_rule_check"], "adults_confirmed")
                ws.cell(row, col["your_validation_element_confirmed"], "yes" if (k + (rv == "C")) % 5 else "unclear")
                ws.cell(row, col["your_comment"], "p.3" if d in ("DUPLICATE_REPORT", "EXCLUDE") else None)
            wb.save(td / f"ft_screen_{rv}.xlsx")
        s = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td)
        conf = {r["record_id"] for r in csv.DictReader((td / "ft_conflicts_for_D.csv").open())}
        assert s["n_records"] == 60 and s["n_complete_pairs"] == 50, s
        assert s["n_without_paired_decision"] == 10 and s["rule_problems"] == {"B": 1, "C": 0}, s
        assert conf == planted_conflicts, (sorted(conf ^ planted_conflicts))
        assert s["rule_problems"]["B"] == 1 and s["rule_problems"]["C"] == 0
        assert s["prisma_final"] is None and not s["final_complete"]
        agreed_n = 50 - len(planted_conflicts)
        assert s["prisma_agreed_only"]["n"] == agreed_n
        # D resolves every conflict -> final counts complete and consistent
        with (td / "res.csv").open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["record_id", "final_disposition", "final_primary_code"])
            for rid in sorted(conf):
                w.writerow([rid, *truth[rid]])
        s2 = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td, td / "res.csv")
        pf = s2["prisma_final"]
        exp = Counter(d for d, _ in truth.values())
        assert s2["final_complete"] and pf["by_disposition"] == {k: exp.get(k, 0) for k in bfw.DISPOSITIONS}, pf
        assert sum(pf["excluded_by_ft_code"].values()) == exp["EXCLUDE"]
        assert pf["included_reports"] == sum(exp[k] for k in INCLUDE)
        canonical = td / "ft_conflicts_for_D.csv"
        assert all(r["final_disposition"] for r in csv.DictReader(canonical.open()))
        canonical_before = canonical.read_bytes()
        # A complete, separate source can update canonical decisions, with the old file backed up.
        first = sorted(conf)[0]
        with (td / "res_alt.csv").open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["record_id", "final_disposition", "final_primary_code"])
            for rid in sorted(conf):
                altered = (("INCLUDE_A", "") if truth[rid][0] == "EXCLUDE" else ("EXCLUDE", codes[0])) if rid == first else truth[rid]
                w.writerow([rid, *altered])
        s3 = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td, td / "res_alt.csv")
        assert s3["conflicts_csv"] == canonical.name and canonical.read_bytes() != canonical_before
        assert any(p.read_bytes() == canonical_before for p in (td / "_backup_ft_merge_canonical").glob("ft_conflicts_for_D_*.csv"))
        canonical_before = canonical.read_bytes()
        # Incomplete and same-path sources never replace a resolved canonical file.
        with (td / "partial.csv").open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["record_id", "final_disposition", "final_primary_code"])
            w.writerow([sorted(conf)[0], *truth[sorted(conf)[0]]])
        s4 = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td, td / "partial.csv")
        assert s4["conflicts_csv"] == "ft_conflicts_for_D_pending.csv" and canonical.read_bytes() == canonical_before
        s5 = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td, canonical)
        assert s5["conflicts_csv"] == "ft_conflicts_for_D_pending.csv" and canonical.read_bytes() == canonical_before
        s6 = merge(td / "ft_screen_B.xlsx", td / "ft_screen_C.xlsx", td)
        assert s6["conflicts_csv"] == "ft_conflicts_for_D_pending.csv" and canonical.read_bytes() == canonical_before
        kp = kappa([("a", "a"), ("b", "b"), ("a", "b"), ("b", "b")])
        assert kp == 0.5, kp
        print(json.dumps({"self_test": "passed", "records": 60, "screened": 50,
                          "planted_conflicts": len(planted_conflicts),
                          "disposition_agreement": s["agreement"]["disposition"], "tmp": str(td)}))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--b", type=Path)
    ap.add_argument("--c", type=Path)
    ap.add_argument("--resolutions", type=Path, help="D's filled ft_conflicts_for_D.csv (final_disposition, final_primary_code)")
    ap.add_argument("--out-dir", type=Path)
    ap.add_argument("--self-test", action="store_true")
    a = ap.parse_args()
    if a.self_test:
        return self_test()
    if not (a.b and a.c):
        ap.error("--b and --c are required (or --self-test)")
    s = merge(a.b, a.c, a.out_dir or a.b.parent, a.resolutions)
    print(json.dumps({k: s[k] for k in ("n_records", "n_complete_pairs", "rule_problems", "n_conflicts_for_D", "final_complete")}))


if __name__ == "__main__":
    main()
