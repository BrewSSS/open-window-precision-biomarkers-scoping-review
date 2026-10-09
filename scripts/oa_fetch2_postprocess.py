#!/usr/bin/env python3
"""D role — post-processing for the second OA fetch pass, 2026-10-09 (scripts/oa_fetch2_run.py).

Run after scripts/oa_fetch2_run.py has finished writing fulltexts/oa_fetch2_ledger.jsonl and
04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_attempts_2026-10-09.csv (and after the
16 new PDFs have had scripts/extract_pdf_text.py run on them into fulltext/V_text/).

What this does:
  1. Updates 04_screening/formal_2026-10-05_v0.9/fulltext/fulltext_fetch_manifest.csv rows for
     the 76 scripted targets (status/source/url/sha256/bytes/verified_text/note), preserving the
     existing scope/final_layer columns. The 6 already-registered manual-browser rows (A's prior
     browser downloads of 6 of the 7 pre-blocked-publisher records) and the still-outstanding
     7th (FS-002317, karger.com) are left untouched/reported as-is.
  2. Writes 04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_2026-10-09.{csv,md}: one row
     per record among all 83 verdict==A_click_download records (route(s) tried, last http
     status, outcome, final verdict).
  3. Updates fulltext_requests/fulltext_requests.csv status for all 83 reference_ids (16 new
     scripted retrievals, 6 pre-existing manual retrievals not yet reflected there, 61 still
     requested).
  4. Rebuilds fulltext_requests/fulltext_acquisition_2026-10-09.xlsx: sheet 1 is cut down to the
     61 records A still has to click (1 pre-blocked-publisher + 60 not retrieved this round);
     sheets 2-4 (maybe_free/unknown_check/library_request) are regenerated from the current
     oa_sweep categories (unchanged membership this round, refreshed formatting/status).
  5. Appends a dated section to fulltext_requests/README.md with the new totals.
"""
from __future__ import annotations

import csv
import json
import urllib.parse
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SWEEP = ROOT / "04_screening/formal_2026-10-05_v0.9/fulltext/oa_sweep_2026-10-09.csv"
MANIFEST_CSV = ROOT / "04_screening/formal_2026-10-05_v0.9/fulltext/fulltext_fetch_manifest.csv"
ATTEMPTS_CSV = ROOT / "04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_attempts_2026-10-09.csv"
LEDGER = ROOT / "fulltexts/oa_fetch2_ledger.jsonl"
REPORT_CSV = ROOT / "04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_2026-10-09.csv"
REPORT_MD = ROOT / "04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_2026-10-09.md"
REQ_CSV = ROOT / "fulltext_requests/fulltext_requests.csv"
README = ROOT / "fulltext_requests/README.md"
XLSX = ROOT / "fulltext_requests/fulltext_acquisition_2026-10-09.xlsx"

PRE_BLOCKED_HOSTS = {
    "bjsm.bmj.com", "journals.humankinetics.com", "journals.physiology.org", "www.physiology.org",
    "www.sciencedirect.com", "onlinelibrary.wiley.com", "karger.com",
    "bmjopensem.bmj.com", "downloads.hindawi.com", "europepmc.org",
    "scholarworks.montana.edu", "www.mdpi.com",
}
MANUAL_DONE = {"FS-013865", "FS-003581", "FS-004116", "FS-006184", "FS-004083", "FS-004580"}
KARGER_OUTSTANDING = "FS-002317"

ACQ_COLUMNS = ["priority", "record_id", "title", "journal", "year", "doi", "free_pdf_link",
               "landing_link", "doi_link", "save_as", "status"]


def host_of(u):
    return urllib.parse.urlparse(u).netloc if u else ""


def load_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_sweep_83():
    return [r for r in load_csv(SWEEP) if r["verdict"] == "A_click_download"]


def load_ledger():
    out = {}
    if LEDGER.exists():
        for line in LEDGER.open(encoding="utf-8"):
            line = line.strip()
            if not line:
                continue
            d = json.loads(line)
            out[d["record_id"]] = d
    return out


def load_attempts():
    by_rec = {}
    if ATTEMPTS_CSV.exists():
        for r in load_csv(ATTEMPTS_CSV):
            by_rec.setdefault(r["record_id"], []).append(r)
    return by_rec


# ---------------------------------------------------------------------------------------------
# Step 1: manifest update
# ---------------------------------------------------------------------------------------------
def update_manifest(ledger: dict):
    rows = load_csv(MANIFEST_CSV)
    fieldnames = list(rows[0].keys())
    updated = 0
    for r in rows:
        rid = r["record_id"]
        d = ledger.get(rid)
        if not d:
            continue
        if d["status"] == "retrieved_oa":
            r["status"] = "retrieved_oa"
            r["source"] = "oa_fetch2_2026-10-09"
            r["url"] = d.get("url") or ""
            r["sha256"] = d.get("sha256") or ""
            r["bytes"] = str(d.get("bytes") or "")
            r["verified_text"] = "True"
            r["note"] = (d.get("note") or "")[:200]
            updated += 1
        else:
            r["status"] = "not_retrieved_oa"
            r["note"] = f"oa_fetch2_2026-10-09: {(d.get('note') or 'no route found')[:170]}"
    with MANIFEST_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    return updated


# ---------------------------------------------------------------------------------------------
# Step 2: per-record report
# ---------------------------------------------------------------------------------------------
def build_report(sweep83, ledger, attempts):
    rows = []
    for r in sweep83:
        rid = r["record_id"]
        h = host_of(r["best_pdf_url"] or r["best_landing_url"])
        if rid in MANUAL_DONE:
            rows.append({"record_id": rid, "priority": r["priority"], "title": r["title"],
                         "route_tried": "manual_browser (done by A before this round)",
                         "http_status": "n/a", "outcome": "already_retrieved",
                         "final_verdict": "retrieved_manual"})
            continue
        if rid == KARGER_OUTSTANDING or h in PRE_BLOCKED_HOSTS:
            rows.append({"record_id": rid, "priority": r["priority"], "title": r["title"],
                         "route_tried": f"none (pre-known blocked publisher host={h})",
                         "http_status": "n/a", "outcome": "blocked_publisher_preknown",
                         "final_verdict": "left_for_A_click"})
            continue
        atts = attempts.get(rid, [])
        d = ledger.get(rid, {})
        route_tried = "; ".join(f"{a['route']}->{a['outcome']}" for a in atts) or "no_attempt_logged"
        last_status = atts[-1]["http_status"] if atts else ""
        verdict = "retrieved_oa" if d.get("status") == "retrieved_oa" else "left_for_A_click"
        outcome = "retrieved" if d.get("status") == "retrieved_oa" else "not_retrieved"
        rows.append({"record_id": rid, "priority": r["priority"], "title": r["title"],
                     "route_tried": route_tried, "http_status": last_status,
                     "outcome": outcome, "final_verdict": verdict})
    rows.sort(key=lambda x: int(x["priority"]) if x["priority"] else 999)

    with REPORT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["record_id", "priority", "title", "route_tried",
                                          "http_status", "outcome", "final_verdict"])
        w.writeheader()
        w.writerows(rows)

    n_retrieved_script = sum(1 for r in rows if r["final_verdict"] == "retrieved_oa")
    n_manual = sum(1 for r in rows if r["final_verdict"] == "retrieved_manual")
    n_left = sum(1 for r in rows if r["final_verdict"] == "left_for_A_click")
    # The attempts CSV's 'host' column is the *requested* URL's host, not necessarily the host
    # that actually served the block (e.g. a doi.org redirect landing on the blocked host) --
    # the accurate list is read from the run's own stderr BLOCK log lines instead.
    blocked_hosts_seen = sorted({
        "dspace.lboro.ac.uk", "figshare.com", "escholarship.org", "ncbi.nlm.nih.gov",
        "cris.maastrichtuniversity.nl", "air.unimi.it", "journals.physiology.org",
        "iris.unipa.it", "www.mdpi.com", "doaj.org", "www.degruyterbrill.com",
    })
    md = [
        f"# OA fetch round 2 (D role, 2026-10-09) — scripts/oa_fetch2_run.py",
        "",
        f"- Scope: 83 records with verdict `A_click_download` in `oa_sweep_2026-10-09.csv`.",
        f"- 7 on pre-known blocked-publisher hosts skipped upfront (6 already retrieved by A in"
        f" a browser before this round; 1, FS-002317/karger.com, still outstanding).",
        f"- Remaining 76 worked in priority order via scripted resolution (DOI-redirect +"
        f" `citation_pdf_url`/obvious-PDF-link HTML parsing; PMC id-converter for the one"
        f" pubmed.ncbi.nlm.nih.gov-only record): **{n_retrieved_script}/76 retrieved**.",
        f"- Already done by A (manual browser, pre-dates this round): {n_manual}.",
        f"- Left for A to click: {n_left} (60 from the 76 + FS-002317).",
        "",
        "Hosts that showed a 403/429/captcha/WAF-challenge block during this run (single-strike,"
        " stopped immediately, never retried, no header/fingerprint spoofing attempted): "
        + (", ".join(blocked_hosts_seen) or "none"),
        "",
        "Deviation note: the first captcha hit was on a `*.ncbi.nlm.nih.gov` subdomain; the "
        "single-strike block was applied to the whole `ncbi.nlm.nih.gov` family (not just the "
        "exact subdomain) since all PMC/PubMed hosts share the same bot-detection system — "
        "retrying a sibling subdomain would be functionally the same as retrying the blocked "
        "host. This cost the 18 www.ncbi.nlm.nih.gov-landing + 2 pmc.ncbi.nlm.nih.gov-direct + 1 "
        "pubmed.ncbi.nlm.nih.gov record in this batch; they are on A's click-list.",
        "",
        "Full per-record log: `oa_fetch2_2026-10-09.csv`. Raw attempt-by-attempt log (route/host/"
        "http_status/outcome, no full-text content): `oa_fetch2_attempts_2026-10-09.csv`.",
    ]
    REPORT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")
    return n_retrieved_script, n_manual, n_left


# ---------------------------------------------------------------------------------------------
# Step 3: fulltext_requests.csv status update
# ---------------------------------------------------------------------------------------------
def update_requests_csv(sweep83, ledger):
    rows = load_csv(REQ_CSV)
    fieldnames = list(rows[0].keys())
    n_updated = 0
    sweep_ids = {r["record_id"] for r in sweep83}
    for r in rows:
        rid = r["reference_id"]
        if rid not in sweep_ids:
            continue
        if rid in MANUAL_DONE:
            if r["status"] != "retrieved_2026-10-09_manual_browser":
                r["status"] = "retrieved_2026-10-09_manual_browser"
                n_updated += 1
            continue
        d = ledger.get(rid)
        if d and d.get("status") == "retrieved_oa":
            r["status"] = "retrieved_2026-10-09_oa_fetch2"
            n_updated += 1
        # else: leave as 'requested' (still on A's click-list)
    with REQ_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    return n_updated


# ---------------------------------------------------------------------------------------------
# Step 4: rebuild the acquisition workbook
# ---------------------------------------------------------------------------------------------
def acq_row(r, status=""):
    doi = (r.get("doi") or "").strip()
    return {
        "priority": r.get("priority", ""), "record_id": r["record_id"], "title": r.get("title", ""),
        "journal": r.get("journal", ""), "year": r.get("year", ""), "doi": doi,
        "free_pdf_link": r.get("best_pdf_url", ""), "landing_link": r.get("best_landing_url", ""),
        "doi_link": f"https://doi.org/{doi}" if doi else "",
        "save_as": f"{r['record_id']}.pdf", "status": status,
    }


def build_workbook(sweep83_leftover, sweep_rest):
    sheets = {
        "1_click_to_download": [acq_row(r) for r in sweep83_leftover],
        "2_maybe_free": [acq_row(r) for r in sweep_rest["likely_free_check_manually"]],
        "3_unknown_check": [acq_row(r) for r in sweep_rest["openalex_unknown"]],
        "4_library_request": [acq_row(r) for r in sweep_rest["closed_library_request"]],
    }
    for rows in sheets.values():
        rows.sort(key=lambda r: int(r["priority"]) if r["priority"] else 999)

    n1, n2, n3, n4 = (len(sheets["1_click_to_download"]), len(sheets["2_maybe_free"]),
                      len(sheets["3_unknown_check"]), len(sheets["4_library_request"]))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)

    readme_ws = wb.create_sheet("README")
    readme_ws.append([
        "全文获取清单（2026-10-09 第二轮开放获取抓取后重建）。见本文件夹 README.md 的最新小节获取"
        "总计数；本工作簿 sheet1 为 A 仍需点击下载的记录。"
    ])
    for row in [
        [],
        ["sheet", "count", "说明"],
        ["1_click_to_download", n1, "出版社/仓储挡脚本或本轮未解析出直链 PDF，需人工点击下载"],
        ["2_maybe_free", n2, "按期刊规则推测可能免费，需人工确认（未经 API 确认）"],
        ["3_unknown_check", n3, "OpenAlex 未能确认开放获取状态，需人工核实"],
        ["4_library_request", n4, "已确认非开放获取，需机构图书馆/馆际互借"],
    ]:
        readme_ws.append(row)
    for c in range(1, 4):
        readme_ws.cell(row=2, column=c).font = bold

    sheet_names = {"1_click_to_download": f"1_click_to_download_{n1}",
                   "2_maybe_free": f"2_maybe_free_{n2}",
                   "3_unknown_check": f"3_unknown_check_{n3}",
                   "4_library_request": f"4_library_request_{n4}"}
    for key, sheet_name in sheet_names.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(ACQ_COLUMNS)
        for c in range(1, len(ACQ_COLUMNS) + 1):
            ws.cell(row=1, column=c).font = bold
        for row in sheets[key]:
            ws.append([row[c] for c in ACQ_COLUMNS])
            r_idx = ws.max_row
            for col_name in ("free_pdf_link", "landing_link", "doi_link"):
                col_idx = ACQ_COLUMNS.index(col_name) + 1
                val = row[col_name]
                if val:
                    cell = ws.cell(row=r_idx, column=col_idx)
                    cell.hyperlink = val
                    cell.style = "Hyperlink"
        for i, col_name in enumerate(ACQ_COLUMNS, 1):
            width = max(12, min(60, max((len(str(r[col_name])) for r in sheets[key]), default=10) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    wb.save(XLSX)
    return n1, n2, n3, n4


# ---------------------------------------------------------------------------------------------
# Step 5: README update
# ---------------------------------------------------------------------------------------------
def update_readme(n_retrieved_script, n_manual_already, n_left, n1, n2, n3, n4, total_retrieved_scope):
    text = README.read_text(encoding="utf-8")
    addition = (
        f"\n## 第二轮开放获取抓取（D 角色，2026-10-09）\n\n"
        f"对 `oa_sweep_2026-10-09.csv` 中 83 条 `A_click_download` 记录做了第二轮抓取："
        f"7 条出版社已知挡脚本（6 条此前 A 已用浏览器下载并入库，1 条 FS-002317/karger.com 仍待点击）"
        f"跳过，其余 76 条按 `priority` 顺序解析（DOI 跳转 + 落地页 `citation_pdf_url`/明显 PDF 链接 "
        f"解析；PMC id-converter 处理唯一的 pubmed 专用链接），**{n_retrieved_script}/76 条取到全文并通过"
        f"校验**（PDF 魔数 + 正文 >2000 字 + DOI 或 ≥3 个标题特征词命中）。\n\n"
        f"本轮遇到的脚本封锁（首次 403/429/验证码/WAF 挑战即停该主机，不重试、不伪造浏览器特征）："
        f"figshare.com、dspace.lboro.ac.uk、escholarship.org、ncbi.nlm.nih.gov（整个子域族，详见 "
        f"`oa_fetch2_2026-10-09.md` 的偏差说明）、cris.maastrichtuniversity.nl、air.unimi.it、"
        f"iris.unipa.it、doaj.org、www.degruyterbrill.com、journals.physiology.org、www.mdpi.com。\n\n"
        f"范围 477 篇全文中，累计已取得 **{total_retrieved_scope}** 篇（含本轮新增 {n_retrieved_script} "
        f"篇脚本抓取 + 此前 {n_manual_already} 篇 A 浏览器下载）。A 仍需点击 **{n_left}** 篇"
        f"（`fulltext_acquisition_2026-10-09.xlsx` sheet1 `1_click_to_download_{n1}`）；另有 "
        f"sheet2 `2_maybe_free_{n2}`、sheet3 `3_unknown_check_{n3}`（均需人工核实，非本轮改动）、"
        f"sheet4 `4_library_request_{n4}`（机构图书馆/馆际互借，非本轮改动）。\n\n"
        f"逐条记录：`04_screening/formal_2026-10-05_v0.9/fulltext/oa_fetch2_2026-10-09.csv`"
        f"（含 route/http_status/outcome/最终判定）与同目录 `oa_fetch2_2026-10-09.md`。本轮新取到"
        f"的 16 篇全文已完成 sha256/字节数登记（`fulltext_fetch_manifest.csv`）与正文抽取"
        f"（`fulltext/V_text/`），尚未提交 AI 预筛（由另一次任务分派）。\n"
    )
    README.write_text(text + addition, encoding="utf-8")


def main():
    sweep83 = load_sweep_83()
    ledger = load_ledger()
    attempts = load_attempts()

    n_manifest_updated = update_manifest(ledger)
    n_retrieved_script, n_manual, n_left = build_report(sweep83, ledger, attempts)
    n_req_updated = update_requests_csv(sweep83, ledger)

    sweep_all = load_csv(SWEEP)
    sweep_rest = {
        "likely_free_check_manually": [r for r in sweep_all if r["verdict"] == "likely_free_check_manually"],
        "openalex_unknown": [r for r in sweep_all if r["verdict"] == "openalex_unknown"],
        "closed_library_request": [r for r in sweep_all if r["verdict"] == "closed_library_request"],
    }
    leftover = [r for r in sweep83 if r["record_id"] not in MANUAL_DONE
                and not (ledger.get(r["record_id"], {}).get("status") == "retrieved_oa")]
    n1, n2, n3, n4 = build_workbook(leftover, sweep_rest)

    # total retrieved within the confirmed 477 scope, post-update
    manifest_rows = load_csv(MANIFEST_CSV)
    total_retrieved_scope = sum(1 for r in manifest_rows if r["scope"] == "confirmed"
                                 and r["status"] in ("retrieved_oa", "retrieved_manual"))

    update_readme(n_retrieved_script, n_manual, n_left, n1, n2, n3, n4, total_retrieved_scope)

    print(f"manifest_rows_updated={n_manifest_updated}")
    print(f"n_retrieved_script={n_retrieved_script} n_manual_already={n_manual} n_left_for_A={n_left}")
    print(f"requests_csv_rows_updated={n_req_updated}")
    print(f"workbook sheets: click={n1} maybe_free={n2} unknown={n3} library={n4}")
    print(f"total_retrieved_within_scope_477={total_retrieved_scope}")
    print(f"left_for_A_to_click={n1} library_request={n4}")


if __name__ == "__main__":
    main()
