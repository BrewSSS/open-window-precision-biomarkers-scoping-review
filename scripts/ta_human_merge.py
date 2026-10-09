#!/usr/bin/env python3
"""Merge reviewer B and C's filled stage-2 exclusion-sample verification workbooks
(ta_review_B.xlsx / ta_review_C.xlsx, built by scripts/ta_ai_merge.py; pre-filled by an AI
stand-in per ai_stage2/README.md "Exclusion-sample verification workbooks pre-filled
(2026-10-09)"). Models scripts/triage_human_merge.py for the earlier ai_triage workbooks, but
reads the sample_verify (519 rows, blind 10% sample of AI-agreed exclusions) and awaiting
(2 rows) sheets, and their your_disposition / your_primary_code / your_comment columns, instead
of the V/U/MX/ADJ _E5/_E4 layout.

For each record (sample_verify + awaiting, both workbooks):
  - B/C values come straight from your_disposition / your_primary_code / your_comment.
  - consensus is the shared value when B and C agree, else "DISAGREE"; if only one reviewer
    has answered, that reviewer's value is the consensus (single_rater).
  - ai_final_code is joined from ai_stage2/ta_ai_merged.csv (the AI exclusion code the sample
    was drawn to verify); only present for sample_verify rows (awaiting is a different
    population -- abstract missing/uninformative, not an agreed-exclusion sample).

Outputs (no abstracts): ta_human_verification_merged.csv, ta_human_verification_summary.json,
ta_human_verification_summary.md:
  - B vs C raw agreement + Cohen's kappa on your_disposition (sample_verify + awaiting
    combined, and sample_verify alone).
  - each reviewer's override rate against their own hidden _prefill log (fraction of records
    where the current your_disposition or your_primary_code differs from what the AI stand-in
    originally wrote).
  - confirmed-ADVANCE rate per AI exclusion code (ai_final_code, sample_verify only) with
    Wilson 95% CIs -- the PRE-007 rule-4 trigger table: any ADVANCE found inside the sample for
    a given code is evidence that code's full ai_excluded_unverified population should be
    re-screened.
  - a disagreement list for D (record_id, sheet, both reviewers' disposition/code/comment).

Self-test (--selftest): copies the real (currently AI-pre-filled, not yet human-edited)
ta_review_B.xlsx / ta_review_C.xlsx under /tmp/triage/selftest_ta_human/, applies synthetic
random overrides to a subset of rows (fixed seed) so the merge has real agreement/override
signal to check, and runs the full merge on the synthetic copies. Nothing under /tmp is
committed.
"""
import argparse
import csv
import json
import math
import os
import random
import shutil
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

CONTENT_SHEETS = ["sample_verify", "awaiting"]
PREFILL_COLUMNS = ("your_disposition", "your_primary_code", "your_comment")


def cohen_kappa(pairs):
    """pairs: list of (rater_a_label, rater_b_label). Returns (raw_agreement, kappa, n).
    Identical algorithm to scripts/triage_compare.py::cohen_kappa / triage_human_merge.py."""
    n = len(pairs)
    if n == 0:
        return (None, None, 0)
    agree = sum(1 for a, b in pairs if a == b)
    po = agree / n
    labels = sorted(set(a for a, _ in pairs) | set(b for _, b in pairs))
    a_counts = Counter(a for a, _ in pairs)
    b_counts = Counter(b for _, b in pairs)
    pe = sum((a_counts[l] / n) * (b_counts[l] / n) for l in labels)
    if pe == 1.0:
        kappa = 1.0 if po == 1.0 else 0.0
    else:
        kappa = (po - pe) / (1 - pe)
    return (po, kappa, n)


def wilson_ci(x, n, z=1.96):
    """Wilson score interval for a binomial proportion x/n. Returns (lo, hi) or (None, None)."""
    if n == 0:
        return (None, None)
    phat = x / n
    denom = 1 + z * z / n
    center = (phat + z * z / (2 * n)) / denom
    half = (z * math.sqrt(phat * (1 - phat) / n + z * z / (4 * n * n))) / denom
    return (max(0.0, center - half), min(1.0, center + half))


def read_workbook(path):
    """Returns (answers, prefill). answers: record_id -> dict(sheet, disposition,
    primary_code, comment). prefill: record_id -> dict(your_disposition=..., your_primary_code=...,
    your_comment=...) taken from the hidden _prefill log (last entry wins per column)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    answers = {}
    for sheet_name in CONTENT_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        rid_col = header.index("record_id")
        disp_col = header.index("your_disposition")
        code_col = header.index("your_primary_code")
        cm_col = header.index("your_comment")
        for row in rows:
            row = tuple(row) + (None,) * (len(header) - len(row))
            rid = row[rid_col]
            if rid is None:
                continue
            answers[rid] = {
                "sheet": sheet_name,
                "disposition": str(row[disp_col]).strip() if row[disp_col] else "",
                "primary_code": str(row[code_col]).strip() if row[code_col] else "",
                "comment": str(row[cm_col]).strip() if row[cm_col] else "",
            }
    prefill = defaultdict(dict)
    if "_prefill" in wb.sheetnames:
        ws = wb["_prefill"]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        idx = {h: i for i, h in enumerate(header)}
        for row in rows:
            if row[idx["record_id"]] is None:
                continue
            rid = row[idx["record_id"]]
            col = row[idx["column"]]
            val = row[idx["value"]]
            prefill[rid][col] = "" if val is None else str(val)
    wb.close()
    return answers, dict(prefill)


def load_ai_final_codes(ai_merged_csv):
    """record_id -> ai_final_code, from ai_stage2/ta_ai_merged.csv. Returns {} if the file is
    not given/found (ai-code trigger table is then skipped, not an error)."""
    out = {}
    if not ai_merged_csv or not os.path.isfile(ai_merged_csv):
        return out
    with open(ai_merged_csv, encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            rid = row.get("record_id")
            if rid:
                out[rid] = (row.get("ai_final_code") or "").strip()
    return out


def override_rate(answers, prefill):
    """Fraction of records whose current your_disposition or your_primary_code differs from
    the logged AI pre-fill value for that column. Returns (n_overridden, n_with_prefill)."""
    n_overridden = 0
    n_with_prefill = 0
    for rid, a in answers.items():
        pf = prefill.get(rid)
        if not pf:
            continue
        n_with_prefill += 1
        pf_disp = pf.get("your_disposition", "")
        pf_code = pf.get("your_primary_code", "")
        if a["disposition"] != pf_disp or a["primary_code"] != pf_code:
            n_overridden += 1
    return n_overridden, n_with_prefill


def merge(wb_b_path, wb_c_path, ai_merged_csv):
    answers_b, prefill_b = (read_workbook(wb_b_path) if wb_b_path else ({}, {}))
    answers_c, prefill_c = (read_workbook(wb_c_path) if wb_c_path else ({}, {}))
    ai_codes = load_ai_final_codes(ai_merged_csv)
    all_ids = sorted(set(answers_b) | set(answers_c))

    rows = []
    disagreements = []
    pairs_all = []
    pairs_sample_verify = []

    for rid in all_ids:
        b = answers_b.get(rid)
        c = answers_c.get(rid)
        sheet = (b or c)["sheet"]
        ai_final_code = ai_codes.get(rid, "") if sheet == "sample_verify" else ""

        if b and c:
            pairs_all.append((b["disposition"] or "BLANK", c["disposition"] or "BLANK"))
            if sheet == "sample_verify":
                pairs_sample_verify.append((b["disposition"] or "BLANK", c["disposition"] or "BLANK"))
            disp_disagree = b["disposition"] != c["disposition"]
            code_disagree = b["primary_code"] != c["primary_code"]
            consensus_disposition = b["disposition"] if not disp_disagree else "DISAGREE"
            consensus_primary_code = b["primary_code"] if not code_disagree else "DISAGREE"
            single_rater = ""
            if disp_disagree or code_disagree:
                disagreements.append({
                    "record_id": rid, "sheet": sheet,
                    "B_disposition": b["disposition"], "C_disposition": c["disposition"],
                    "B_primary_code": b["primary_code"], "C_primary_code": c["primary_code"],
                    "B_comment": b["comment"], "C_comment": c["comment"],
                })
        elif b or c:
            single = b or c
            consensus_disposition = single["disposition"]
            consensus_primary_code = single["primary_code"]
            single_rater = "B" if b else "C"
        else:
            continue

        rows.append({
            "record_id": rid, "source_sheet": sheet, "ai_final_code": ai_final_code,
            "B_disposition": b["disposition"] if b else "", "B_primary_code": b["primary_code"] if b else "",
            "B_comment": b["comment"] if b else "",
            "C_disposition": c["disposition"] if c else "", "C_primary_code": c["primary_code"] if c else "",
            "C_comment": c["comment"] if c else "",
            "consensus_disposition": consensus_disposition, "consensus_primary_code": consensus_primary_code,
            "single_rater": single_rater,
        })

    po_all, k_all, n_all = cohen_kappa(pairs_all)
    po_sv, k_sv, n_sv = cohen_kappa(pairs_sample_verify)

    b_overridden, b_n = override_rate(answers_b, prefill_b)
    c_overridden, c_n = override_rate(answers_c, prefill_c)

    # PRE-007 rule-4 trigger table: confirmed-ADVANCE rate per AI exclusion code,
    # sample_verify only, using consensus_disposition where determined.
    by_code = defaultdict(list)
    for r in rows:
        if r["source_sheet"] != "sample_verify" or not r["ai_final_code"]:
            continue
        if r["consensus_disposition"] in ("", "DISAGREE"):
            continue
        by_code[r["ai_final_code"]].append(r)
    trigger_table = {}
    for code, recs in sorted(by_code.items()):
        n = len(recs)
        advanced = sum(1 for r in recs if r["consensus_disposition"] == "ADVANCE")
        lo, hi = wilson_ci(advanced, n)
        trigger_table[code] = {
            "n": n, "confirmed_advance": advanced,
            "confirmed_advance_rate": advanced / n if n else None,
            "wilson_95ci": [lo, hi],
            "rescreen_trigger": advanced > 0,
        }

    summary = {
        "n_records_merged": len(rows),
        "n_with_both_raters": sum(1 for r in rows if not r["single_rater"]),
        "n_single_rater": sum(1 for r in rows if r["single_rater"]),
        "bc_agreement_disposition_all": {"n": n_all, "raw_agreement": po_all, "kappa": k_all},
        "bc_agreement_disposition_sample_verify_only": {"n": n_sv, "raw_agreement": po_sv, "kappa": k_sv},
        "override_rate": {
            "B": {"n_overridden": b_overridden, "n_with_prefill": b_n,
                  "rate": b_overridden / b_n if b_n else None},
            "C": {"n_overridden": c_overridden, "n_with_prefill": c_n,
                  "rate": c_overridden / c_n if c_n else None},
        },
        "n_disagreements": len(disagreements),
        "pre007_rule4_trigger_table_by_ai_code": trigger_table,
        "consensus_disposition_counts": dict(Counter(r["consensus_disposition"] for r in rows)),
    }
    return rows, disagreements, summary


def write_outputs(out_dir, rows, disagreements, summary, suffix=""):
    os.makedirs(out_dir, exist_ok=True)
    merged_csv = os.path.join(out_dir, f"ta_human_verification_merged{suffix}.csv")
    fieldnames = list(rows[0].keys()) if rows else [
        "record_id", "source_sheet", "ai_final_code",
        "B_disposition", "B_primary_code", "B_comment",
        "C_disposition", "C_primary_code", "C_comment",
        "consensus_disposition", "consensus_primary_code", "single_rater",
    ]
    with open(merged_csv, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)

    summary_json = os.path.join(out_dir, f"ta_human_verification_summary{suffix}.json")
    with open(summary_json, "w", encoding="utf-8") as f:
        json.dump({**summary, "disagreements": disagreements}, f, indent=2)
        f.write("\n")

    summary_md = os.path.join(out_dir, f"ta_human_verification_summary{suffix}.md")

    def fmt(v):
        return f"{v:.3f}" if isinstance(v, float) else "n/a"

    lines = [
        "# Stage-2 exclusion-sample verification: human-verification merge summary",
        "",
        f"Records merged: {summary['n_records_merged']} "
        f"(both raters: {summary['n_with_both_raters']}, single rater: {summary['n_single_rater']}).",
        "",
        "## B vs C agreement on your_disposition",
        "",
        f"All records (sample_verify + awaiting): n={summary['bc_agreement_disposition_all']['n']}, "
        f"raw agreement={fmt(summary['bc_agreement_disposition_all']['raw_agreement'])}, "
        f"kappa={fmt(summary['bc_agreement_disposition_all']['kappa'])}.",
        f"sample_verify only: n={summary['bc_agreement_disposition_sample_verify_only']['n']}, "
        f"raw agreement={fmt(summary['bc_agreement_disposition_sample_verify_only']['raw_agreement'])}, "
        f"kappa={fmt(summary['bc_agreement_disposition_sample_verify_only']['kappa'])}.",
        "",
        "## Override rate vs each reviewer's own AI pre-fill",
        "",
        f"B: {summary['override_rate']['B']['n_overridden']} / {summary['override_rate']['B']['n_with_prefill']} "
        f"({fmt(summary['override_rate']['B']['rate'])}).",
        f"C: {summary['override_rate']['C']['n_overridden']} / {summary['override_rate']['C']['n_with_prefill']} "
        f"({fmt(summary['override_rate']['C']['rate'])}).",
        "",
        "## PRE-007 rule-4 trigger table (confirmed-ADVANCE rate per AI exclusion code, sample_verify only)",
        "",
        "| ai_final_code | n | confirmed ADVANCE | rate | Wilson 95% CI | re-screen trigger |",
        "|---|---|---|---|---|---|",
    ]
    for code, a in summary["pre007_rule4_trigger_table_by_ai_code"].items():
        lo, hi = a["wilson_95ci"]
        ci = f"[{lo:.3f}, {hi:.3f}]" if lo is not None else "n/a"
        lines.append(f"| {code} | {a['n']} | {a['confirmed_advance']} | {fmt(a['confirmed_advance_rate'])} | "
                     f"{ci} | {'YES' if a['rescreen_trigger'] else 'no'} |")
    lines += [
        "",
        f"Disagreements for D's adjudication: {summary['n_disagreements']} "
        "(see ta_human_verification_summary.json: disagreements[]).",
        "",
        f"Consensus disposition counts: {summary['consensus_disposition_counts']}",
    ]
    with open(summary_md, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return merged_csv, summary_json, summary_md


# --------------------------------------------------------------------------------------
# Self-test: synthetic random overrides on copies of the real (AI-pre-filled) workbooks
# --------------------------------------------------------------------------------------

def selftest():
    real_dir = "04_screening/formal_2026-10-05_v0.9/ai_stage2"
    src_b = os.path.join(real_dir, "ta_review_B.xlsx")
    src_c = os.path.join(real_dir, "ta_review_C.xlsx")
    if not (os.path.isfile(src_b) and os.path.isfile(src_c)):
        print(f"SELFTEST SKIPPED: need both ta_review_B.xlsx and ta_review_C.xlsx under {real_dir}",
              file=sys.stderr)
        return 0
    tmp_dir = "/tmp/triage/selftest_ta_human"
    if os.path.isdir(tmp_dir):
        shutil.rmtree(tmp_dir)
    os.makedirs(tmp_dir, exist_ok=True)

    rng = random.Random(4242)
    disp_choices = ["ADVANCE", "EXCLUDE_TA", "AWAITING_CLASSIFICATION"]
    code_choices = ["", "FT01", "FT02", "FT03", "FT04", "FT05", "FT06", "FT07", "FT08"]

    filled_paths = {}
    expected_total = None
    for code, src in (("B", src_b), ("C", src_c)):
        dst = os.path.join(tmp_dir, os.path.basename(src))
        shutil.copy(src, dst)
        wb = load_workbook(dst)
        n_rows_this_wb = 0
        for sheet_name in CONTENT_SHEETS:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            n_rows_this_wb += ws.max_row - 1
            header = [c.value for c in ws[1]]
            disp_col = header.index("your_disposition") + 1
            code_col = header.index("your_primary_code") + 1
            cm_col = header.index("your_comment") + 1
            for r in range(2, ws.max_row + 1):
                if rng.random() < 0.3:  # simulate a human override on ~30% of rows
                    disp = rng.choice(disp_choices)
                    ws.cell(row=r, column=disp_col, value=disp)
                    ws.cell(row=r, column=code_col,
                            value=rng.choice(code_choices) if disp == "EXCLUDE_TA" else "")
                    ws.cell(row=r, column=cm_col, value="selftest override")
                # else: leave the AI pre-fill value untouched (simulated accept)
        wb.save(dst)
        filled_paths[code] = dst
        expected_total = n_rows_this_wb
        print(f"selftest: synthetic copy with ~30% overrides -> {dst} ({n_rows_this_wb} rows)",
              file=sys.stderr)

    ai_merged_csv = os.path.join(real_dir, "ta_ai_merged.csv")
    rows, disagreements, summary = merge(filled_paths["B"], filled_paths["C"], ai_merged_csv)
    ok = True

    def check(label, cond):
        nonlocal ok
        status = "PASS" if cond else "FAIL"
        if not cond:
            ok = False
        print(f"[{status}] {label}", file=sys.stderr)

    check(f"merged rows == {expected_total}", summary["n_records_merged"] == expected_total)
    check("every record has both raters in selftest", summary["n_with_both_raters"] == expected_total)
    k = summary["bc_agreement_disposition_all"]["kappa"]
    if k is not None:
        check("kappa in [-1, 1]", -1.0 <= k <= 1.0)
    ra = summary["bc_agreement_disposition_all"]["raw_agreement"]
    if ra is not None:
        check("raw agreement in [0, 1]", 0.0 <= ra <= 1.0)
    for reviewer in ("B", "C"):
        rate = summary["override_rate"][reviewer]["rate"]
        if rate is not None:
            check(f"{reviewer} override rate in [0, 1]", 0.0 <= rate <= 1.0)
        check(f"{reviewer} override rate roughly matches the ~30% synthetic injection rate",
              rate is None or 0.1 <= rate <= 0.5)
    for code, a in summary["pre007_rule4_trigger_table_by_ai_code"].items():
        lo, hi = a["wilson_95ci"]
        if lo is not None:
            check(f"AI code {code}: Wilson CI ordered and in [0,1]",
                  0.0 <= lo <= a["confirmed_advance_rate"] <= hi <= 1.0)
    out_dir = os.path.join(tmp_dir, "out")
    merged_csv, summary_json, summary_md = write_outputs(out_dir, rows, disagreements, summary)
    check("merged CSV written", os.path.isfile(merged_csv))
    with open(merged_csv, encoding="utf-8") as f:
        header = f.readline()
    check("merged CSV has no 'abstract' column", "abstract" not in header.lower())
    check("summary json/md written", os.path.isfile(summary_json) and os.path.isfile(summary_md))
    print("SELFTEST", "PASSED" if ok else "FAILED", file=sys.stderr)
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--wb-b", default=None)
    ap.add_argument("--wb-c", default=None)
    ap.add_argument("--ai-merged-csv", default="04_screening/formal_2026-10-05_v0.9/ai_stage2/ta_ai_merged.csv")
    ap.add_argument("--out-dir", default="04_screening/formal_2026-10-05_v0.9/ai_stage2")
    ap.add_argument("--out-suffix", default="", help="appended to output filenames, e.g. '_provisional'")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()

    if a.selftest:
        sys.exit(selftest())

    if not a.wb_b and not a.wb_c:
        ap.error("provide at least one of --wb-b / --wb-c")

    rows, disagreements, summary = merge(a.wb_b, a.wb_c, a.ai_merged_csv)
    if not rows:
        print("No answered records found in the provided workbook(s); nothing written.", file=sys.stderr)
        sys.exit(1)
    merged_csv, summary_json, summary_md = write_outputs(a.out_dir, rows, disagreements, summary, a.out_suffix)
    print(f"wrote {merged_csv}\nwrote {summary_json}\nwrote {summary_md}")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
