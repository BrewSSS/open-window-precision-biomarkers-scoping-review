"""Guarded full-text workbook writes shared by the three screening stages."""
from __future__ import annotations

import datetime as dt
import csv
import hashlib
import json
import os
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def ids_from_file(path: Path) -> set[str]:
    ids = {line.strip() for line in path.read_text(encoding="utf-8").splitlines()
           if line.strip() and not line.lstrip().startswith("#")}
    if not ids or not all(r.startswith("FS-") for r in ids):
        raise ValueError("ids-file is empty or contains an invalid record ID")
    return ids


def require_in_scope(ids: set[str], scope_csv: Path) -> None:
    with scope_csv.open(newline="", encoding="utf-8") as stream:
        scope = {row["record_id"] for row in csv.DictReader(stream)}
    missing = sorted(ids - scope)
    if missing:
        raise ValueError(f"selected IDs are outside the authoritative full-text scope: {', '.join(missing)}")


def valid_runs(runs_dir: Path, ids: set[str]) -> dict[str, dict]:
    runs = {}
    for rid in sorted(ids):
        path = runs_dir / f"{rid}.json"
        if not path.is_file():
            continue
        try:
            run = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if run.get("record_id") != rid:
            raise ValueError(f"run record ID mismatch: {path}")
        if run.get("valid") and isinstance(run.get("parsed"), dict):
            if run.get("model") != "gpt-6-sol":
                raise ValueError(f"selected run is not the required Sol model: {path}")
            runs[rid] = run
    if not runs:
        raise ValueError(f"no valid selected runs in {runs_dir}")
    return runs


def snapshot(wb):
    # Empty strings round-trip through XLSX as blank cells.
    return {ws.title: {(cell.row, cell.column): None if cell.value == "" else cell.value
                       for row in ws for cell in row} for ws in wb}


def rows_by_id(ws) -> dict[str, int]:
    cols = {cell.value: cell.column for cell in ws[1]}
    out = {}
    for row in range(2, ws.max_row + 1):
        rid = ws.cell(row, cols["record_id"]).value
        if rid:
            if rid in out:
                raise ValueError(f"duplicate workbook record ID: {rid}")
            out[rid] = row
    return out


def latest_audit_values(wb, audit_sheet: str, value_header: str) -> dict[tuple[str, str], object]:
    if audit_sheet not in wb:
        return {}
    ws = wb[audit_sheet]
    cols = {cell.value: cell.column for cell in ws[1]}
    if not {"record_id", "column", value_header} <= cols.keys():
        raise ValueError(f"audit sheet lacks provenance columns: {audit_sheet}")
    result = {}
    for row in range(2, ws.max_row + 1):
        rid = ws.cell(row, cols["record_id"]).value
        col = ws.cell(row, cols["column"]).value
        if rid and col:
            result[(rid, col)] = ws.cell(row, cols[value_header]).value
    return result


def check_existing(value, rid: str, col: str, provenance: dict, stage: str) -> None:
    if value in (None, ""):
        return
    key = (rid, col)
    if key not in provenance or provenance[key] != value:
        raise ValueError(f"{stage}: nonempty {rid}/{col} differs from recorded AI provenance; human value preserved")


def append_audit(wb, name: str, headers: list[str], rows: list[list]) -> None:
    if name in wb:
        ws = wb[name]
        existing = [cell.value for cell in ws[1]]
        if existing != headers:
            raise ValueError(f"audit header mismatch in {name}")
    else:
        ws = wb.create_sheet(name)
        ws.append(headers)
        ws["A1"].font = Font(bold=True)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in rows:
        ws.append(row)
    ws.sheet_state = "hidden"


def guarded_save(path: Path, wb, before: dict, allowed: dict[str, set[tuple[int, int]]],
                 audit_sheet: str, audit_rows: int, dry_run: bool = False) -> dict:
    """Verify every sheet/cell before replacing the file; retain a unique ignored backup."""
    after = snapshot(wb)
    if not set(before) <= set(after) or set(after) - set(before) - {audit_sheet}:
        raise ValueError("unexpected sheet deletion or creation")
    diffs = {}
    for sheet in after:
        old, new = before.get(sheet, {}), after[sheet]
        changed = {key for key in set(old) | set(new) if old.get(key) != new.get(key)}
        if sheet == audit_sheet:
            if sheet in before:
                start = max((r for r, _ in old), default=0) + 1
                permitted = {(r, c) for r in range(start, start + audit_rows)
                             for c in range(1, wb[sheet].max_column + 1)}
                if changed - permitted:
                    raise ValueError(f"historical {audit_sheet} cells changed")
            else:
                permitted = set(new)
            if changed - permitted:
                raise ValueError(f"unexpected {audit_sheet} cell change")
        elif changed - allowed.get(sheet, set()):
            raise ValueError(f"unexpected cell changes in {sheet}: {sorted(changed - allowed.get(sheet, set()))[:8]}")
        diffs[sheet] = len(changed)
    if dry_run:
        return {"dry_run": True, "changed_cells_by_sheet": diffs}

    backup_dir = path.parent / "_backup_batch3"
    backup_dir.mkdir(parents=True, exist_ok=True)
    # Ensure a backup always precedes the replacement, with no overwrite of an earlier backup.
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, backup)
    fd, tmp_name = tempfile.mkstemp(prefix="._ft_write_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        verified = load_workbook(tmp)
        verified_after = snapshot(verified)
        if verified_after != after or verified.sheetnames != wb.sheetnames or any(
                verified[name].sheet_state != wb[name].sheet_state for name in wb.sheetnames):
            raise ValueError("saved workbook differs from verified in-memory workbook")
        verified.close()
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
    root = Path(__file__).resolve().parents[1]
    try:
        backup_label = str(backup.resolve().relative_to(root))
    except ValueError:
        backup_label = backup.name
    return {"backup": backup_label, "changed_cells_by_sheet": diffs, "sha256_after": sha256(path)}


def update_manifest(path: Path, key: str, block: dict, reviewer: str | None,
                    hash_key: str | None, workbook_hash: str) -> None:
    manifest = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    manifest[key] = block
    if reviewer and hash_key:
        manifest.setdefault("workbooks", {}).setdefault(reviewer, {})[hash_key] = workbook_hash
    path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
